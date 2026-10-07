"""
Tests de las jugadas que se agregaron a la notacion:

* libre y toque que terminan el punto: "_P" (punto) y "_U_Y" (usan el bloqueo);
* el toque de bloqueo que deja la pelota viva del lado del que bloquea
  ("_BD_Y"), despues del cual la pelota que se levanta NO es una defensa;
* el saque de potencia ("_P" despues de la zona de destino);
* la calidad del armado (A+ / A0 / A-, y el AX que es la armada mala "_-2");
* el tipo de resolucion del ataque (potente "_PO", colocado "_CO", y el
  block-out que es el resultado U).

Se corren con:
    python -m unittest test_jugadas_nuevas
"""
import unittest

import analisis_voley as av
import generar_informe_volley as gi
import notacion
import sesion_web
import test_carga_visual as t


def punto(equipo_saca, *lineas):
    """Un punto armado a mano, con los equipos puestos como los pone el motor."""
    sesion = sesion_web.SesionPartido()
    for linea in ("Local", "Rival", "", "", equipo_saca) + lineas:
        sesion.enviar(linea)
    return sesion.estado["puntos"][-1]


class TestParser(unittest.TestCase):

    def test_saque_de_potencia(self):
        for texto, resultado in (("5_1_6_P_A", "A"), ("5_1_6_P_E", "E"),
                                 ("5_1_6_P_X/3_3/2_4/4_1_P", "X")):
            bloque = av.parsear_bloque_saque(texto)
            self.assertIsNotNone(bloque, texto)
            self.assertTrue(bloque["saque_potencia"], texto)
            self.assertEqual(bloque["resultado_saque"], resultado)

    def test_el_saque_de_siempre_no_es_de_potencia(self):
        for texto in ("5_1_6_A", "1_6_X/3_3/2_4/4_1_P", "5_1_6_X/3_-1"):
            self.assertIs(av.parsear_bloque_saque(texto)["saque_potencia"], False, texto)

    def test_libre_y_toque_que_terminan_el_punto(self):
        casos = {
            "7_2/1_5/9_F_8_P": ("FP", None), "7_2/1_5/9_T_8_P": ("TP", None),
            "7_2/1_5/9_F_8_U_6": ("FU", 6), "7_2/1_5/9_T_8_U_6": ("TU", 6),
            "7_2/1_5/9_F_8": ("F", None), "7_2/1_5/9_T_0": ("TE", None),
        }
        for texto, (resultado, bloqueo) in casos.items():
            bloque = av.parsear_bloque_defensa(texto)
            self.assertEqual((bloque["resultado"], bloque["jugador_bloqueo"]),
                             (resultado, bloqueo), texto)
            self.assertEqual(bloque["atacante"], 9)

    def test_el_libre_malo_no_lleva_final(self):
        self.assertIsNone(av.parsear_bloque_defensa("7_2/1_5/9_F_0_P"))
        self.assertIsNone(av.parsear_bloque_defensa("7_2/1_5/9_T_0_U_6"))

    def test_toque_de_bloqueo(self):
        bloque = av.parsear_bloque_defensa("7_2/1_5/9_5_BD_6")
        self.assertEqual((bloque["resultado"], bloque["jugador_bloqueo"]), ("BD", 6))


class TestMotor(unittest.TestCase):

    def test_libre_punto_es_del_que_lo_hizo(self):
        p = punto("A", "5_1_6_X/3_3/2_4/4_F_3_P")
        self.assertEqual(p["equipo_gana"], "B")
        self.assertEqual(av.causa_del_punto(p), "FP")

    def test_toque_usando_el_bloqueo_es_del_que_toco(self):
        p = punto("A", "5_1_6_X/3_3/2_4/4_T_3_U_7")
        self.assertEqual(p["equipo_gana"], "B")
        self.assertEqual(av.causa_del_punto(p), "TU")
        self.assertEqual(p["jugadas"][0]["equipo_bloqueo"], "A")

    def test_el_toque_de_bloqueo_pasa_la_pelota_al_que_bloqueo(self):
        # B ataca, A toca en el bloqueo y la sigue jugando A
        p = punto("A", "5_1_6_X/3_3/2_4/4_1_BD_7", "8_2/1_4/9_5_P")
        self.assertEqual(p["equipo_gana"], "A")
        self.assertEqual(p["jugadas"][0]["equipo_bloqueo"], "A")


class TestEstadisticas(unittest.TestCase):

    def test_tras_el_toque_de_bloqueo_no_hay_defensa(self):
        con_bloqueo = punto("A", "5_1_6_X/3_3/2_4/4_1_BD_7", "8_2/1_4/9_5_P")
        sin_bloqueo = punto("A", "5_1_6_X/3_3/2_4/4_1_D", "8_2/1_4/9_5_P")
        self.assertEqual(list(av.defensas_del_punto(con_bloqueo)), [])
        self.assertEqual(len(list(av.defensas_del_punto(sin_bloqueo))), 1)

    def test_para_el_atacante_el_toque_de_bloqueo_es_defendido(self):
        p = punto("A", "5_1_6_X/3_3/2_4/4_1_BD_7", "8_2/1_4/9_5_P")
        ataque = av.calcular_estadisticas_ataque([p])
        self.assertEqual(ataque["B"][4]["defendido"], 1)

    def test_toque_de_bloqueo_y_pelota_no_jugada_es_ataque_efectivo(self):
        p = punto("A", "5_1_6_X/3_3/2_4/4_1_BD_7", "8_0")
        self.assertEqual(p["equipo_gana"], "B")
        self.assertEqual(av.calcular_estadisticas_ataque([p])["B"][4]["efectivo"], 1)
        self.assertEqual(list(av.defensas_del_punto(p)), [])

    def test_toques_de_bloqueo_por_jugador(self):
        p = punto("A", "5_1_6_X/3_3/2_4/4_1_BD_7", "8_2/1_4/9_5_P")
        self.assertEqual(av.calcular_toques_de_bloqueo([p]), {"A": {7: 1}, "B": {}})
        # no es un bloqueo punto
        self.assertEqual(av.calcular_estadisticas_bloqueo([p]), {"A": {}, "B": {}})

    def test_libres_y_toques_no_son_ataques(self):
        p = punto("A", "5_1_6_X/3_3/2_4/4_F_3_P")
        self.assertEqual(av.calcular_estadisticas_ataque([p])["B"], {})
        libres = av.calcular_libres_y_toques([p])["B"][4]
        self.assertEqual(libres["F"]["punto"], 1)
        self.assertEqual(sum(libres["T"].values()), 0)

    def test_saques_normales_y_de_potencia(self):
        puntos = [punto("A", "5_1_6_P_A"), punto("A", "5_1_6_P_E"),
                  punto("A", "5_1_6_A"), punto("A", "5_1_6_X/3_3/2_4/4_1_P")]
        saque = av.calcular_estadisticas_saque(puntos)["A"][5]
        self.assertEqual(saque["potencia"], {"saques": 2, "as": 1, "error": 1})
        self.assertEqual(saque["normal"], {"saques": 2, "as": 1, "error": 0})


class TestVolcadoYExcel(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        sesion = sesion_web.SesionPartido()
        lineas = ("Local", "Rival", "", "", "A",
                  "5_1_6_P_A", "5_1_6_P_E",
                  "3_1_6_X/3_3/2_4/4_1_BD_7", "8_2/1_4/9_5_D", "6_1/2_3/4_T_2_U_9",
                  "8_1_5_P_X/3_2/2_4/4_F_3_P",
                  "4_6_6_X/3_2/2_4/4_F_3", "7_2/1_5/9_F_0")
        for linea in lineas:
            sesion.enviar(linea)
        cls.puntos = sesion.estado["puntos"]
        cls.texto = av.formatear_estadisticas(cls.puntos, {"A": "Local", "B": "Rival"})
        cls.volcado = gi.parse_volcado(cls.texto)

    # Como va el partido: A saca y hace as con potencia (5), A saca y la erra
    # (5, potencia). Saca B (3): A ataca, el 7 de B toca en el bloqueo, B
    # ataca y A defiende, y el toque del 4 de A usa el bloqueo del 9: punto A.
    # Saca A (8, potencia): el libre del 4 de B es punto. Saca B (4): el 4 de
    # A hace un libre que sigue y el libre del 9 de B sale mal: punto A.

    def test_el_parser_lee_las_secciones_nuevas(self):
        local = self.volcado["teams"]["Local"]
        rival = self.volcado["teams"]["Rival"]
        self.assertEqual(local["saques_jugador"]["Jugador 5"]["Potencia"],
                         {"saques": 2, "as": 1, "error": 1})
        self.assertEqual(local["saques_jugador"]["Jugador 8"]["Potencia"]["saques"], 1)
        self.assertEqual(rival["saques_jugador"]["Jugador 3"]["Normal"]["saques"], 1)
        self.assertEqual(rival["toques_bloqueo_jugador"], {"Jugador 7": 1})
        self.assertEqual(local["toques_bloqueo_jugador"], {})
        self.assertEqual(local["libres_jugador"]["Jugador 4"]["Libre"],
                         {"total": 1, "punto": 0, "usado": 0, "sigue": 1, "malo": 0})
        self.assertEqual(local["libres_jugador"]["Jugador 4"]["Toque"]["usado"], 1)
        self.assertEqual(rival["libres_jugador"]["Jugador 4"]["Libre"]["punto"], 1)
        self.assertEqual(rival["libres_jugador"]["Jugador 9"]["Libre"]["malo"], 1)
        self.assertEqual(rival["causas"]["hechos"]["Libre punto"], 1)
        self.assertEqual(local["causas"]["hechos"]["Toque usando el bloqueo"], 1)

    def libro(self, equipo, rival):
        try:
            import openpyxl                      # noqa: F401
        except ImportError:
            self.skipTest("necesita openpyxl")
        import valores_excel
        libro, _ = gi.build_workbook(equipo, rival, self.volcado)
        valores_excel.convertir_a_valores(libro)
        return libro

    @staticmethod
    def filas(hoja, columnas):
        return [[hoja.cell(row=f, column=c).value for c in range(1, columnas + 1)]
                for f in range(1, hoja.max_row + 1)]

    def test_el_excel_trae_los_saques(self):
        hoja = self.libro("Local", "Rival")["Saque"]
        por_tipo = {fila[0]: fila[1:] for fila in self.filas(hoja, 8)}
        # saques, % de los saques, as, errores, en juego
        self.assertEqual(por_tipo["Potencia"][:5], [3, 1, 1, 1, 1])
        self.assertEqual(por_tipo["Normal"][0], 0)
        self.assertIn(["Jugador 5", 2, 1, 1, 0, 0.5, 0.5, 2, 1], self.filas(hoja, 9))

    def test_el_excel_trae_libres_y_toques(self):
        filas = self.filas(self.libro("Local", "Rival")["Ataque jugador"], 8)
        self.assertIn(["Jugador 4", "Toque", 1, 0, 1, 0, 0, 1], filas)
        self.assertIn(["Jugador 4", "Libre", 1, 0, 0, 1, 0, 0], filas)

    def test_el_excel_trae_los_toques_de_bloqueo(self):
        filas = self.filas(self.libro("Rival", "Local")["Ataque jugador"], 3)
        inicio = filas.index(["Jugador", "Toques de bloqueo", "% del total del equipo"])
        self.assertEqual(filas[inicio + 1], ["Jugador 7", 1, 1])

    def test_el_excel_trae_las_causas_nuevas(self):
        hoja = self.libro("Local", "Rival")["Fases y armador"]
        causas = {fila[0]: fila[2] for fila in self.filas(hoja, 3)}
        self.assertEqual(causas["Toque usando el bloqueo"], 1)
        self.assertEqual(causas["Libre punto"], 0)

    def test_un_volcado_viejo_sigue_armando_el_excel(self):
        """Sin las secciones nuevas las tablas salen vacias, no con error."""
        corte = self.texto.index("Toques de bloqueo por jugador:")
        resto = self.texto[corte:]
        siguiente_equipo = resto.find("\n\n--- ")
        viejo = self.texto[:corte] + (resto[siguiente_equipo:] if siguiente_equipo >= 0 else "")
        volcado = gi.parse_volcado(viejo)
        libro, _ = gi.build_workbook("Local", "Rival", volcado)
        self.assertIn("Saque", libro.sheetnames)


class TestCargaTocando(unittest.TestCase):
    """Lo nuevo se tiene que poder armar tocando, con el mismo texto."""

    def armar(self, espera, objetivo, equipo_saca="A"):
        armador = notacion.Armador(espera, equipo_saca, planteles={},
                                   sacador_conocido=False)
        candidatos = [int(n) for n in dict.fromkeys(__import__("re").findall(r"\d+", objetivo))]
        return t.buscar_toques(armador, objetivo, candidatos)

    def test_se_arman_tocando(self):
        for espera, linea in (
                ("saque", "5_1_6_P_A"),
                ("saque", "5_1_6_P_X/3_3/2_4/4_1_BD_7"),
                ("continuacion", "7_2/1_5/9_F_8_P"),
                ("continuacion", "7_2/1_5/9_F_8"),
                ("continuacion", "7_2/1_5/9_F_0")):
            self.assertIsNotNone(self.armar(espera, linea), linea)

    def test_la_potencia_se_marca_una_sola_vez(self):
        armador = notacion.Armador("saque", "A", sacador_conocido=False)
        armador.tocar("otro", 5)
        armador.tocar("z1")
        armador.tocar("z6")
        armador.tocar("potencia")
        self.assertNotIn("potencia", [o["id"] for o in armador.opciones()])
        armador.tocar("as")
        armador.tocar("as_sin_receptor")
        self.assertEqual(armador.linea, "5_1_6_P_A")
        self.assertTrue(armador.cerrada)



class TestArmadoYTipoDeAtaque(unittest.TestCase):

    def test_parser(self):
        casos = {
            "5_1_6_X/3_3/2_4_A+/4_1_PO_P": ("+", True, "PO"),
            "5_1_6_X/3_3/2_4_X_A-/4_1_CO_D": ("-", False, "CO"),
            "5_1_6_X/3_3/2_4_A0_X/4_1_U_7": ("0", False, None),
            "5_1_6_X/3_3/2_4/4_1_P": (None, True, None),
        }
        for texto, esperado in casos.items():
            b = av.parsear_bloque_saque(texto)
            self.assertEqual((b["calidad_armado"], b["armado_valido"], b["tipo_ataque"]),
                             esperado, texto)
        self.assertEqual(av.parsear_bloque_saque("5_1_6_X/3_3/2_-2")["calidad_armado"], "X")
        self.assertEqual(av.parsear_bloque_defensa("9_A_1_PO_P")["tipo_ataque"], "PO")

    def test_no_se_aceptan_dos_calidades(self):
        self.assertIsNone(av.parsear_bloque_saque("5_1_6_X/3_3/2_4_A+_A0/4_1_P"))
        self.assertIsNone(av.parsear_bloque_saque("5_1_6_X/3_3/2_4/4_1_PO_CO_P"))

    def test_la_direccion_sigue_siendo_obligatoria(self):
        self.assertIsNone(av.parsear_bloque_saque("5_1_6_X/3_3/2_4/4_PO_P"))
        self.assertIsNone(av.parsear_bloque_saque("5_1_6_X/3_3/2_4/4_3_PO_P"))

    def test_calidad_por_armador_incluye_el_ax(self):
        puntos = [punto("A", "5_1_6_X/3_3/2_4_A+/4_1_P"), punto("A", "5_1_6_X/3_3/2_-2"),
                  punto("A", "5_1_6_X/3_3/2_4/4_1_P"),
                  punto("A", "5_1_6_X/3_3/2_4_X_A-/4_1_P")]   # sin armado: no cuenta
        calidad = av.calcular_calidad_armado(puntos)
        self.assertEqual(calidad["B"][2], {"+": 1, "0": 0, "-": 0, "X": 1, "sin": 1})
        self.assertEqual(calidad["A"], {})

    def test_ataque_segun_calidad_del_armado(self):
        p = punto("A", "5_1_6_X/3_3/2_4_A-/4_1_O")
        datos = av.calcular_ataque_por_calidad_armado([p])["B"][4]
        self.assertEqual(datos["-"], {"efectivo": 0, "defendido": 0, "fuera": 1})
        # el ataque de primera no viene de un armado
        p = punto("A", "5_1_6_X/3_3/2_4/4_1_D", "9_A_1_PO_P")
        self.assertEqual(av.calcular_ataque_por_calidad_armado([p])["A"], {})

    def test_tipo_y_direccion(self):
        p = punto("A", "5_1_6_X/3_3/2_4/4_1_PO_D", "7_2/1_4/9_6_U_4")
        tipos = av.calcular_ataque_por_tipo([p])
        self.assertEqual(tipos["B"][4], {"Potente": {1: {"efectivo": 0, "defendido": 1, "fuera": 0}}})
        # el block-out se reconoce por el resultado, sin marca
        self.assertEqual(tipos["A"][9], {"Block out": {6: {"efectivo": 1, "defendido": 0, "fuera": 0}}})


class TestArmadoYTipoEnElExcel(unittest.TestCase):
    """A saca. B arma A+ y el 4 de B hace punto potente a la 1. Saca B: el 8
    de A arma A-, el 9 coloca a la 5 y lo defienden; B arma A0 y el 4 hace
    block-out hacia la 6. Saca B: el 8 de A arma mal (AX). Saca B: el 8 arma
    sin calificar y el 9 hace punto a la 1 sin tipo."""

    @classmethod
    def setUpClass(cls):
        sesion = sesion_web.SesionPartido()
        for linea in ("Local", "Rival", "", "", "A",
                      "5_1_6_X/3_3/2_4_A+/4_1_PO_P",
                      "7_1_5_X/1_2/8_3_A-/9_5_CO_D", "4_2/2_4_A0/4_6_U_7",
                      "7_1_5_X/1_2/8_-2",
                      "7_1_5_X/1_2/8_3/9_1_P"):
            sesion.enviar(linea)
        texto = av.formatear_estadisticas(sesion.estado["puntos"], {"A": "Local", "B": "Rival"})
        cls.volcado = gi.parse_volcado(texto)

    def libro(self, equipo, rival):
        import valores_excel
        libro, _ = gi.build_workbook(equipo, rival, self.volcado)
        valores_excel.convertir_a_valores(libro)
        return libro

    @staticmethod
    def filas(hoja, columnas):
        return [[hoja.cell(row=f, column=c).value for c in range(1, columnas + 1)]
                for f in range(1, hoja.max_row + 1)]

    def test_el_parser(self):
        local = self.volcado["teams"]["Local"]
        self.assertEqual(local["calidad_armado_jugador"]["Jugador 8"],
                         {"A+": 0, "A0": 0, "A-": 1, "AX": 1, "Sin calificar": 1})
        self.assertEqual(local["ataque_calidad"]["Jugador 9"]["A-"], (0, 1, 0))
        self.assertIn(("Jugador 9", "Colocado", "5", 0, 1, 0), local["ataque_tipo"])
        rival = self.volcado["teams"]["Rival"]
        self.assertIn(("Jugador 4", "Block out", "6", 1, 0, 0), rival["ataque_tipo"])

    def test_hoja_armado(self):
        filas = self.filas(self.libro("Local", "Rival")["Armado"], 11)
        self.assertIn(["Jugador 8", 0, 0, 1, 1, 2, 0, 0, 0.5, 0.5, 1], filas)

    def test_hoja_ataque_por_calidad(self):
        filas = self.filas(self.libro("Local", "Rival")["Ataque jugador"], 9)
        self.assertIn(["Jugador 9", "A-", 1, 0, 1, 0, 0, 0, 0], filas)
        self.assertIn(["Jugador 9", "Sin calificar", 1, 1, 0, 0, 1, 0, 1], filas)

    def test_hoja_tipo_y_direccion(self):
        filas = self.filas(self.libro("Rival", "Local")["Zona y dirección"], 10)
        self.assertIn(["Block out", 1, 1, 0, 0, 1, 0, 1, None, None], filas)
        # zonas 1, 5, 6; total; % de cada una; % de la mas usada
        self.assertIn(["Jugador 4", "Potente", 1, 0, 0, 1, 1, 0, 0, 1], filas)
        self.assertIn(["Jugador 4", "Block out", 0, 0, 1, 1, 0, 0, 1, 1], filas)


class TestArmadoYTipoTocando(unittest.TestCase):

    def test_se_arman_tocando(self):
        for espera, linea in (("saque", "5_1_6_X/3_3/2_4_A+/4_1_PO_P"),
                              ("continuacion", "7_2/1_5_A-_X/9_5_CO_D"),
                              ("continuacion", "7_2/1_5_X_A0/9_5_U_4"),
                              ("continuacion", "9_A_1_PO_P")):
            armador = notacion.Armador(espera, "A", sacador_conocido=False)
            candidatos = [int(n) for n in dict.fromkeys(__import__("re").findall(r"\d+", linea))]
            self.assertIsNotNone(t.buscar_toques(armador, linea, candidatos), linea)

    def test_una_sola_calidad_y_un_solo_tipo(self):
        armador = notacion.Armador("continuacion", "A", equipo_con_la_pelota="A")
        for toque in (("otro", 7), ("c2",), ("otro", 1), ("z5",), ("armado_mas",)):
            armador.tocar(*toque)
        ids = [o["id"] for o in armador.opciones()]
        self.assertNotIn("armado_cero", ids)
        self.assertNotIn("armado_menos", ids)
        self.assertIn("sin_armado", ids)
        for toque in (("otro", 9), ("ataque",), ("z5",)):
            armador.tocar(*toque)
        # despues de la direccion se pregunta como ataco, y recien despues
        # que paso
        self.assertEqual(armador.estado, "ATACA_FORMA")
        self.assertEqual([o["id"] for o in armador.opciones()],
                         ["potente", "colocado", "sin_forma"])
        armador.tocar("colocado")
        self.assertEqual(armador.estado, "ATACA_RESULTADO")
        ids = [o["id"] for o in armador.opciones()]
        self.assertNotIn("potente", ids)
        self.assertIn("malla", ids)
        armador.tocar("defendido")
        self.assertEqual(armador.linea, "7_2/1_5_A+/9_5_CO_D")



# Un partido corto con rotacion, tiempos y cambios:
#   Local (A) arranca 3 88 15 13 16 21, armador 3, libero 9 por 15 y 21.
#   1) saca el 3 de 1 a 5: as sobre el 4 del rival         -> 1-0
#   T_B (tiempo del rival con 1-0)
#   2) saca el 3 de 1 a 5, recibe el 4 con 2, ataca el rival -> 1-1
#   C_28_16 (Local: sale el 16, entra el 28, con 1-1)
#   3) saca el rival; Local recibe y hace punto              -> 2-1
#   T_A (tiempo de Local con 2-1)
#   4) saca el 88 de 1 a 1 y la erra                         -> 2-2
LINEAS_PARTIDO = ("Local", "Rival", "3_S 88 15 13 16 21  9_L_15_21", "1 2_S 3 4 5 6", "A",
                  "1_5_A_4", "T_B", "1_5_X/4_2/2_3/3_1_P", "C_28_16",
                  "6_1_X/88_3/3_4/13_1_P", "T_A", "1_1_E")


def sesion_del_partido():
    sesion = sesion_web.SesionPartido()
    for linea in LINEAS_PARTIDO:
        resultado = sesion.enviar(linea)
        assert resultado["ok"], (linea, resultado["mensaje"])
    return sesion


class TestTiemposYCambios(unittest.TestCase):

    def test_el_tiempo_guarda_el_marcador_exacto(self):
        tiempos = sesion_del_partido().estado["tiempos"]
        self.assertEqual([(t["equipo"], t["marcador"]) for t in tiempos],
                         [("B", {"A": 1, "B": 0}), ("A", {"A": 2, "B": 1})])

    def test_el_cambio_guarda_el_marcador(self):
        cambio = sesion_del_partido().estado["cambios"][0]
        self.assertEqual((cambio["sale"], cambio["entra"], cambio["marcador"]),
                         (16, 28, {"A": 1, "B": 1}))

    def test_la_x_saca_el_tiempo(self):
        sesion = sesion_web.SesionPartido()
        for linea in LINEAS_PARTIDO[:7]:
            sesion.enviar(linea)
        self.assertEqual(sesion.instantanea()["deshacer"], "tiempo")
        sesion.enviar("x")
        self.assertEqual(sesion.estado["tiempos"], [])
        self.assertEqual(sesion.estado["marcador"], {"A": 1, "B": 0})   # el punto sigue

    def test_no_se_pide_tiempo_con_la_pelota_en_juego(self):
        sesion = sesion_web.SesionPartido()
        for linea in ("Local", "Rival", "", "", "A", "5_1_6_X/3_3/2_4/4_1_D"):
            sesion.enviar(linea)
        self.assertFalse(sesion.enviar("T_A")["ok"])

    def test_rallies_antes_y_despues(self):
        estado = sesion_del_partido().estado
        lista = av.intervenciones(estado["cambios"], estado["tiempos"], estado["nombres"])
        self.assertEqual([i["texto"] for i in lista],
                         ["Tiempo Rival", "Cambio Local: sale #16 / entra #28", "Tiempo Local"])
        antes, despues = av.rallies_alrededor(estado["puntos"], lista[0])
        self.assertEqual((antes, despues), ({"A": 1, "B": 0}, {"A": 1, "B": 2}))

    def test_el_volcado_trae_la_cronologia(self):
        estado = sesion_del_partido().estado
        texto = av.formatear_estadisticas(estado["puntos"], estado["nombres"],
                                          estado["armadores"], estado["rotaciones_por_set"],
                                          estado["cambios"], estado["tiempos"])
        self.assertIn("  Set 1 | 1-1 | Cambio Local: sale #16 / entra #28 | antes 1-1 | despues 1-1",
                      texto)
        # del lado del rival el marcador va al reves
        self.assertIn("  Set 1 | 0-1 | Tiempo Rival | antes 0-1 | despues 2-1", texto)


class TestRotaciones(unittest.TestCase):

    def test_cada_punto_guarda_la_formacion(self):
        puntos = sesion_del_partido().estado["puntos"]
        self.assertEqual(puntos[0]["formacion"]["A"], [3, 88, 15, 13, 16, 21])
        self.assertEqual(puntos[2]["formacion"]["A"], [3, 88, 15, 13, 28, 21])  # ya entro el 28

    def test_rotacion_y_formacion(self):
        datos = av.calcular_rotacion_y_formacion(sesion_del_partido().estado["puntos"])["A"]
        self.assertEqual(datos[(1, "3 88 15 13 16 21")]["saque_jugados"], 2)
        self.assertEqual(datos[(1, "3 88 15 13 28 21")]["rec_ganados"], 1)
        self.assertEqual(datos[(6, "88 15 13 28 21 3")]["ganados"], 0)


class TestEstrategiaDeSaque(unittest.TestCase):

    def test_as_con_receptor(self):
        bloque = av.parsear_bloque_saque("5_1_6_A_7")
        self.assertEqual((bloque["resultado_saque"], bloque["receptor_as"]), ("A", 7))
        self.assertIsNone(bloque["receptor"])   # no es una recepcion
        self.assertIsNone(av.parsear_bloque_saque("5_1_6_E_7"))

    def test_por_estrategia_y_por_receptor(self):
        puntos = sesion_del_partido().estado["puntos"]
        estrategia = av.calcular_saque_por_estrategia(puntos)["A"]
        self.assertEqual(estrategia[(3, "1 a 5")]["saques"], 2)
        self.assertEqual(estrategia[(3, "1 a 5")]["as"], 1)
        self.assertEqual(estrategia[(3, "1 a 5")]["rec2"], 1)
        self.assertEqual(estrategia[(88, "1 a 1")]["error"], 1)
        receptores = av.calcular_saque_por_receptor(puntos)["A"]
        self.assertEqual(receptores[4]["saques"], 2)   # el as y la recepcion
        self.assertEqual(receptores[4]["ganados"], 1)

    def test_el_as_pide_el_receptor_tocando(self):
        armador = notacion.Armador("saque", "A", planteles={"B": [4, 5]})
        for toque in ("z1", "z5", "as"):
            armador.tocar(toque)
        self.assertEqual(armador.estado, "AS_RECEPTOR")
        self.assertIn("j4", [o["id"] for o in armador.opciones()])   # los del rival
        armador.tocar("j4")
        self.assertEqual(armador.linea, "1_5_A_4")
        self.assertTrue(armador.cerrada)


class TestSaqueRotacionesYTiemposEnElExcel(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        estado = sesion_del_partido().estado
        texto = av.formatear_estadisticas(estado["puntos"], estado["nombres"],
                                          estado["armadores"], estado["rotaciones_por_set"],
                                          estado["cambios"], estado["tiempos"])
        import valores_excel
        libro, _ = gi.build_workbook("Local", "Rival", gi.parse_volcado(texto))
        valores_excel.convertir_a_valores(libro)
        cls.libro = libro

    def filas(self, hoja, columnas):
        ws = self.libro[hoja]
        return [[ws.cell(row=f, column=c).value for c in range(1, columnas + 1)]
                for f in range(1, ws.max_row + 1)]

    def test_estrategia_de_saque(self):
        filas = self.filas("Saque", 12)
        self.assertIn(["De 1 a 5", 2, 1, 0, 0, 1, 0, 0, 0.5, 0, 0, 0.5], filas)
        self.assertIn(["Receptor 4", 2, 1, 0, 1, 0, 0, 0.5, 0, 0.5, None, None], filas)

    def test_rotaciones(self):
        filas = self.filas("Rotaciones", 11)
        self.assertIn(["Set 1", 3, 88, 15, 13, 16, 21, 3, "9 por 15 y 21", None, None], filas)
        self.assertIn(["Armador en Z1", 3, 2, 2 / 3, 2, 1, 0.5, 1, 1, 1, None], filas)
        self.assertIn(["Armador en Z1", "3 88 15 13 28 21", 1, 1, 1, 0, 0, "", 1, 1, 1], filas)

    def test_tiempos_y_cambios(self):
        filas = self.filas("Tiempos y cambios", 9)
        self.assertIn(["Set 1", "2-1", "Tiempo Local", 2, 1, 0, 1, 2 / 3, 0], filas)



class TestCalidadDelArmadoPorZona(unittest.TestCase):

    def test_por_zona(self):
        puntos = [punto("A", "5_1_6_X/3_3/2_4_A+/4_1_P"), punto("A", "5_1_6_X/3_3/2_4_A0/4_1_P"),
                  punto("A", "5_1_6_X/3_3/2_2_A-/4_1_P"), punto("A", "5_1_6_X/3_3/2_6/4_1_P")]
        datos = av.calcular_calidad_armado_por_zona(puntos)["B"][2]
        self.assertEqual(datos["4"], {"+": 1, "0": 1, "-": 0, "sin": 0})
        self.assertEqual(datos["2"], {"+": 0, "0": 0, "-": 1, "sin": 0})
        self.assertEqual(datos["6-5"]["sin"], 1)

    def test_llega_al_excel(self):
        sesion = sesion_web.SesionPartido()
        # ataques afuera: gana A y sigue sacando A, asi arma siempre el 2 de B
        for linea in ("Local", "Rival", "", "", "A", "5_1_6_X/3_3/2_4_A+/4_1_O",
                      "5_1_6_X/3_3/2_4_A-/4_1_O", "5_1_6_X/3_3/2_2_A+/4_1_O"):
            sesion.enviar(linea)
        import valores_excel
        volcado = gi.parse_volcado(av.formatear_estadisticas(
            sesion.estado["puntos"], {"A": "Local", "B": "Rival"}))
        libro, _ = gi.build_workbook("Rival", "Local", volcado)
        valores_excel.convertir_a_valores(libro)
        hoja = libro["Armado"]
        filas = [[hoja.cell(row=f, column=c).value for c in range(1, 11)]
                 for f in range(1, hoja.max_row + 1)]
        self.assertIn(["Zona 4", 1, 0, 1, 2, 0.5, 0, 0.5, 0, None], filas)
        self.assertIn(["Jugador 2", "Zona 2", 1, 0, 0, 1, 1, 0, 0, 0], filas)


class TestFichaDelJugador(unittest.TestCase):
    """La pestana Jugadores lee los volcados guardados: se guarda uno de
    verdad en una carpeta aparte y se arma la ficha como lo hace el servidor."""

    @classmethod
    def setUpClass(cls):
        import tempfile
        from pathlib import Path
        from unittest import mock
        import estadisticas_jugadores as ej
        cls.carpeta = Path(tempfile.mkdtemp())
        sesion = sesion_web.SesionPartido()
        for linea in ("Palestino", "Rival", "3_S 88 15 13 16 21  9_L_15_21", "1 2_S 3 4 5 6", "A",
                      "1_5_P_A_4", "T_B", "1_5_X/4_2/2_3/3_1_P", "C_28_16",
                      "6_1_X/9_3/3_4_A+/13_1_PO_P", "T_A", "1_1_E",
                      "1_6_X/9_2/3_2_A-/15_5_CO_BD_4", "5_1/2_4/4_1_D", "21_2/3_4_A0/13_6_U_5",
                      "1_5_X/4_3/2_4/4_1_D", "9_3/3_4_A+/13_F_8_P"):
            assert sesion.enviar(linea)["ok"], linea
        with mock.patch.object(av, "nombre_de_volcado_libre",
                               return_value=cls.carpeta / "partido_20261005_120000.txt"):
            sesion.guardar()
        agregado = ej.agregar(carpeta=cls.carpeta)
        cls.ficha = {d: ej.ficha("Palestino", d, agregado) for d in ("3", "13", "28", "9")}

    def test_saque(self):
        saque = self.ficha["3"]["saque"]
        self.assertEqual(saque["total"], {"saques": 2, "as": 1, "error": 0, "en_juego": 1})
        self.assertEqual([t["saques"] for t in saque["por_tipo"]], [1, 1])   # normal, potencia
        self.assertEqual(saque["estrategias"][0]["par"], "1 a 5")
        self.assertEqual(self.ficha["3"]["indicadores"]["aces"], 1)

    def test_calidad_del_armado_por_zona(self):
        calidad = self.ficha["3"]["armado_calidad"]
        self.assertEqual(calidad["total"]["A+"], 2)
        por_zona = {z["zona"]: z for z in calidad["por_zona"]}
        self.assertEqual((por_zona["4"]["A+"], por_zona["4"]["A0"]), (2, 1))
        self.assertEqual(por_zona["2"]["A-"], 1)

    def test_ataque_por_calidad_y_tipo(self):
        ataque = self.ficha["13"]["ataque"]
        por_calidad = {f["calidad"]: f for f in ataque["por_calidad_armado"]}
        self.assertEqual(por_calidad["A+"]["punto"], 1)
        por_tipo = {t["tipo"]: t["valores"] for t in ataque["por_tipo"]}
        self.assertEqual(por_tipo, {"Potente": [1, 0, 0], "Block out": [0, 0, 1]})  # 1, 5, 6

    def test_libres(self):
        self.assertEqual(self.ficha["13"]["libres"][0]["punto"], 1)

    def test_rotacion_y_sexteto(self):
        self.assertEqual(self.ficha["13"]["rotacion"]["titular"][0]["zona"], 4)
        self.assertTrue(self.ficha["9"]["rotacion"]["titular"][0]["libero"])
        cancha = self.ficha["28"]["rotacion"]["en_cancha"]
        self.assertEqual((cancha["jugados"], cancha["ganados"]), (4, 3))

    def test_cambios(self):
        cambio = self.ficha["28"]["cambios"][0]
        self.assertEqual((cambio["que"], cambio["otro"], cambio["marcador"]), ("Entra", "16", "1-1"))

    def test_resumen_por_partido(self):
        partido = self.ficha["3"]["por_partido"][0]
        self.assertEqual((partido["saques"], partido["aces"], partido["potencia"]), (2, 1, 1))
        self.assertEqual((partido["armado_mas"], partido["armado_calificados"]), (2, 4))
        self.assertEqual(partido["en_cancha"], 6)


if __name__ == "__main__":
    unittest.main()
