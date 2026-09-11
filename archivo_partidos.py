"""
Lectura de los partidos ya archivados en disco (Datos/*.txt e Informes/*.xlsx).

Este modulo es de solo lectura: no conoce la sesion en curso ni la modifica.
Se ocupa de tres cosas que la interfaz web necesita para la pestana Partidos:

1) Armar el listado cruzando las dos carpetas. Un volcado y su informe son el
   mismo partido, asi que se emparejan por (equipos, fecha) y salen en una
   sola fila. Para el listado alcanza con la cabecera del .txt (hasta
   "=== Estadisticas por equipo ==="), que es la parte barata de leer; el
   resultado se cachea por (ruta, mtime) porque el buscador de la web pide
   la lista seguido y el disco no tiene por que enterarse.

2) Devolver un partido entero ya parseado. El parser es el de
   generar_informe_volley (parse_volcado), mas las secciones de rotaciones y
   cambios que ese parser no mira porque el Excel no las usa.

3) Devolver un informe .xlsx como datos planos. openpyxl guarda la formula
   pero no su resultado, asi que si el archivo nunca paso por Excel las
   celdas vienen vacias; para eso se reusa valores_excel.convertir_a_valores,
   que es el evaluador que ya tiene el proyecto.

Todo nombre de archivo que llega de la red se resuelve con ruta_segura(): el
servidor escucha en toda la WiFi y no hay ninguna razon para abrir algo que
no este en Datos/ o en Informes/.
"""
import importlib.util
import re
import sys
import types
from pathlib import Path

import analisis_voley as av

CARPETA_DATOS = av.CARPETA_DATOS
CARPETA_INFORMES = av.CARPETA_INFORMES

# Se mira una sola vez y antes de cualquier import con truco (ver
# _modulo_parser): despues sys.modules puede tener un openpyxl de juguete y
# find_spec ya no serviria para saber si el de verdad esta instalado.
HAY_OPENPYXL = importlib.util.find_spec("openpyxl") is not None

MENSAJE_SIN_OPENPYXL = "Falta openpyxl: pip install openpyxl"
MENSAJE_ARCHIVO_ABIERTO = "Cerra el .xlsx en Excel y volve a intentar"


class RutaInvalida(ValueError):
    """El nombre pedido no cae dentro de las carpetas del proyecto."""


class FaltaOpenpyxl(RuntimeError):
    """Se pidio algo del .xlsx y openpyxl no esta instalado."""


# ======================================================================
# 1) RUTAS
# ======================================================================

def ruta_segura(nombre: str, carpetas) -> Path:
    """Resuelve `nombre` dentro de alguna de esas carpetas.

    Se compara despues de resolve() porque es la unica forma de que
    "../../algo", un enlace simbolico o una ruta absoluta no se escapen; los
    nombres se concatenan con el operador de Path, nunca a mano."""
    if not nombre or not str(nombre).strip():
        raise RutaInvalida("Falta el nombre del archivo.")

    candidato = Path(str(nombre).strip())
    for carpeta in carpetas:
        base = Path(carpeta).resolve()
        # un nombre absoluto se acepta solo si ya cae adentro; asi el cliente
        # puede mandar de vuelta la ruta que le devolvio /api/excel
        destino = (candidato if candidato.is_absolute() else base / candidato).resolve()
        if destino == base or base in destino.parents:
            return destino
    raise RutaInvalida(f"El archivo {nombre!r} no esta en Datos/ ni en Informes/.")


# Cada tipo de archivo vive en una sola carpeta, asi que el tipo alcanza para
# saber donde buscarlo y con que Content-Type contestar.
CARPETA_POR_TIPO = {"txt": lambda: CARPETA_DATOS, "xlsx": lambda: CARPETA_INFORMES}
MIME_POR_TIPO = {
    "txt": "text/plain; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


def ruta_de_tipo(nombre: str, tipo: str | None = None) -> tuple[Path, str]:
    """Resuelve un nombre pedido desde la web y devuelve (ruta, tipo).

    Si no viene el tipo se deduce de la extension. Se valida antes de tocar
    el disco: un tipo que no conocemos o una ruta que se sale de las carpetas
    del proyecto no llegan a abrir nada."""
    tipo = (tipo or Path(str(nombre)).suffix.lstrip(".")).lower()
    if tipo not in CARPETA_POR_TIPO:
        raise RutaInvalida(f"Tipo de archivo no valido: {tipo!r} (txt o xlsx).")

    ruta = ruta_segura(nombre, [CARPETA_POR_TIPO[tipo]()])
    if ruta.suffix.lower() != f".{tipo}":
        raise RutaInvalida(f"El archivo {ruta.name!r} no es un .{tipo}.")
    if es_temporal(ruta.name):
        raise RutaInvalida("Ese es un temporal de Excel, no un archivo del proyecto.")
    return ruta, tipo


def es_temporal(nombre: str) -> bool:
    """Los "~$Informe_....xlsx" que deja Excel mientras tiene el archivo
    abierto no son informes: son un candado de dos KB."""
    return Path(nombre).name.startswith("~$")


# ======================================================================
# 2) EL PARSER DEL VOLCADO
# ======================================================================

def _modulo_parser():
    """generar_informe_volley, aunque no haya openpyxl.

    parse_volcado es todo biblioteca estandar, pero vive en un modulo que
    importa openpyxl arriba para construir los estilos del Excel. Sin
    openpyxl instalado ese import falla y no se podria ni mirar un .txt, asi
    que en ese caso se le da un openpyxl de juguete: las constantes de estilo
    se construyen con objetos que aceptan cualquier cosa y el parser queda
    disponible. Es preferible a tener un segundo parser del volcado, que es
    justo lo que no queremos mantener por duplicado."""
    if not HAY_OPENPYXL:
        _instalar_openpyxl_de_juguete()
    import generar_informe_volley

    return generar_informe_volley


def _instalar_openpyxl_de_juguete():
    if "openpyxl" in sys.modules:
        return

    class Cualquiera:
        """Acepta cualquier llamada: solo tiene que dejar pasar los estilos."""

        def __init__(self, *args, **kwargs):
            pass

        def __call__(self, *args, **kwargs):
            return self

    openpyxl = types.ModuleType("openpyxl")
    estilos = types.ModuleType("openpyxl.styles")
    utiles = types.ModuleType("openpyxl.utils")
    for nombre in ("Font", "PatternFill", "Alignment", "Border", "Side"):
        setattr(estilos, nombre, Cualquiera)
    utiles.get_column_letter = Cualquiera()
    utiles.column_index_from_string = Cualquiera()
    openpyxl.styles = estilos
    openpyxl.utils = utiles
    sys.modules.update({"openpyxl": openpyxl, "openpyxl.styles": estilos,
                        "openpyxl.utils": utiles})


# ======================================================================
# 3) CABECERA Y RESUMEN DE UN VOLCADO
# ======================================================================

CORTE_CABECERA = "=== Estadisticas por equipo ==="

RE_NOMBRE_VOLCADO = re.compile(
    r"^partido_(?P<anio>\d{4})(?P<mes>\d{2})(?P<dia>\d{2})"
    r"_(?P<hora>\d{2})(?P<minuto>\d{2})(?P<segundo>\d{2})\.txt$", re.IGNORECASE)

RE_NOMBRE_INFORME = re.compile(
    r"^Informe_(?P<equipo>.+?)_vs_(?P<rival>.+?)"
    r"_(?P<fecha>\d{4}-\d{2}-\d{2})(?P<extra>_.+)?\.xlsx$", re.IGNORECASE)

RE_SETS_GANADOS = re.compile(r"^Sets:\s*(.+?)\s+(\d+)\s*-\s*(\d+)\s+(.+)$")
RE_PARCIAL = re.compile(r"^Set\s+(\d+):\s*(.+?)\s+(\d+)\s*-\s*(\d+)\s+(.+)$")
RE_MARCADOR_FINAL = re.compile(r"^Marcador final:\s*(.+?)\s+(\d+)\s*-\s*(\d+)\s+(.+)$")
RE_TOTAL_PUNTOS = re.compile(r"^Total de puntos cargados:\s*(\d+)$")

RE_ROTACION_SET = re.compile(r"^Set\s+(\d+):$")
RE_ROTACION_EQUIPO = re.compile(r"^(.+?):\s*(.+/.+)$")
RE_CAMBIO = re.compile(
    r"^Set\s+(\d+)\s+-\s+(.+?):\s*entra\s+(\S+),\s*sale\s+(\S+)"
    r"\s*\(zona\s+([\w-]+)\)\s*(.*)$")


def cabecera_volcado(ruta) -> str:
    """El .txt hasta donde arrancan las estadisticas por equipo.

    Se corta ahi porque el listado solo necesita nombres, sets y puntos, y
    esa parte son cuatro lineas mientras que el bloque de estadisticas son
    ochocientas."""
    lineas = []
    with open(ruta, encoding="utf-8") as archivo:
        for linea in archivo:
            if linea.strip() == CORTE_CABECERA:
                break
            lineas.append(linea.rstrip("\n"))
    return "\n".join(lineas)


def resumen_volcado(texto: str) -> dict:
    """Nombres, sets, parciales y puntos cargados de la cabecera de un .txt."""
    equipo = rival = None
    sets_ganados = None
    parciales = []
    puntos = None

    for linea in texto.splitlines():
        s = linea.strip()
        m = RE_SETS_GANADOS.match(s)
        if m:
            equipo, rival = m.group(1).strip(), m.group(4).strip()
            sets_ganados = (int(m.group(2)), int(m.group(3)))
            continue
        m = RE_PARCIAL.match(s)
        if m:
            equipo = equipo or m.group(2).strip()
            rival = rival or m.group(5).strip()
            parciales.append((int(m.group(1)), int(m.group(3)), int(m.group(4))))
            continue
        m = RE_MARCADOR_FINAL.match(s)
        if m:
            # partido de un solo set: no hay linea "Sets:" ni parciales
            equipo, rival = m.group(1).strip(), m.group(4).strip()
            parciales.append((1, int(m.group(2)), int(m.group(3))))
            continue
        m = RE_TOTAL_PUNTOS.match(s)
        if m:
            puntos = int(m.group(1))

    # El set que todavia no se jugo se guarda como 0-0 (el marcador del set en
    # curso al momento de guardar). No dice nada, y en la lista solo ensucia.
    jugados = [p for p in parciales if p[1] or p[2]]
    if sets_ganados is None and jugados:
        sets_ganados = (sum(1 for _, a, b in jugados if a > b),
                        sum(1 for _, a, b in jugados if b > a))

    return {
        "equipo": equipo,
        "rival": rival,
        "sets": f"{sets_ganados[0]}-{sets_ganados[1]}" if sets_ganados else None,
        "sets_ganados": list(sets_ganados) if sets_ganados else None,
        "parciales": [f"{a}-{b}" for _, a, b in jugados],
        "puntos": puntos,
    }


def secciones_de_cancha(texto: str) -> dict:
    """Rotaciones por set y cambios registrados.

    parse_volcado no las mira porque el informe de Excel no las usa, pero en
    la pantalla del partido son justo lo que explica por que roto quien."""
    rotaciones, cambios = [], []
    seccion = None
    set_actual = None

    for linea in texto.splitlines():
        s = linea.strip()
        if s.startswith("=== Rotaciones"):
            seccion, set_actual = "rotaciones", None
            continue
        if s == "=== Cambios ===":
            seccion = "cambios"
            continue
        if s.startswith("==="):
            seccion = None
            continue
        if not s or seccion is None:
            continue

        if seccion == "rotaciones":
            m = RE_ROTACION_SET.match(s)
            if m:
                set_actual = {"set": int(m.group(1)), "equipos": []}
                rotaciones.append(set_actual)
                continue
            m = RE_ROTACION_EQUIPO.match(s)
            if m and set_actual is not None:
                jugadores = [j.strip() for j in m.group(2).split("/")]
                armador = next((j[:-2] for j in jugadores if j.endswith("-S")), None)
                set_actual["equipos"].append({
                    "equipo": m.group(1).strip(),
                    "jugadores": [j[:-2] if j.endswith("-S") else j for j in jugadores],
                    "armador": armador,
                })
            continue

        m = RE_CAMBIO.match(s)
        if m:
            cambios.append({
                "set": int(m.group(1)), "equipo": m.group(2).strip(),
                "entra": m.group(3), "sale": m.group(4), "zona": m.group(5),
                "detalle": m.group(6).strip(" ()"),
            })

    return {"rotaciones": rotaciones, "cambios": cambios}


# ======================================================================
# 4) LISTADO
# ======================================================================

# (ruta -> (mtime, resumen)). El buscador de la web filtra en el cliente, pero
# la lista se vuelve a pedir cada vez que se entra a la pestana y releer
# veinte volcados en cada visita no tiene sentido si no cambio ninguno.
_CACHE_CABECERA: dict[str, tuple[float, dict]] = {}


def _resumen_cacheado(ruta: Path) -> dict:
    clave = str(ruta)
    mtime = ruta.stat().st_mtime
    guardado = _CACHE_CABECERA.get(clave)
    if guardado is not None and guardado[0] == mtime:
        return guardado[1]
    resumen = resumen_volcado(cabecera_volcado(ruta))
    _CACHE_CABECERA[clave] = (mtime, resumen)
    return resumen


def _fecha_y_hora(nombre: str) -> tuple[str | None, str | None]:
    m = RE_NOMBRE_VOLCADO.match(Path(nombre).name)
    if not m:
        return None, None
    p = m.groupdict()
    return f"{p['anio']}-{p['mes']}-{p['dia']}", f"{p['hora']}:{p['minuto']}"


def _clave_partido(equipo, rival, fecha) -> tuple:
    """Los dos equipos ordenados, porque el informe puede estar hecho desde
    cualquiera de los dos lados (Informe_A_vs_B o Informe_B_vs_A)."""
    nombres = sorted(str(n or "").strip().lower() for n in (equipo, rival))
    return (nombres[0], nombres[1], fecha or "")


def listar_partidos(carpeta_datos=None, carpeta_informes=None) -> list[dict]:
    """Una fila por partido, la mas nueva arriba.

    Un partido puede tener volcado, informe o los dos. Los informes que no
    corresponden a ningun volcado (porque el .txt se borro) igual aparecen:
    el archivo existe y se puede abrir."""
    carpeta_datos = Path(carpeta_datos or CARPETA_DATOS)
    carpeta_informes = Path(carpeta_informes or CARPETA_INFORMES)

    filas, por_clave = [], {}

    for ruta in sorted(carpeta_datos.glob("*.txt")) if carpeta_datos.is_dir() else []:
        if es_temporal(ruta.name):
            continue
        try:
            resumen = _resumen_cacheado(ruta)
        except OSError:
            continue
        fecha, hora = _fecha_y_hora(ruta.name)
        fila = {
            "id": ruta.name,
            "fecha": fecha or "",
            "hora": hora or "",
            "equipo": resumen["equipo"] or "?",
            "rival": resumen["rival"] or "?",
            "sets": resumen["sets"] or "",
            "parciales": resumen["parciales"],
            "puntos": resumen["puntos"] or 0,
            "volcado": ruta.name,
            "informe": None,
        }
        filas.append(fila)
        por_clave.setdefault(_clave_partido(fila["equipo"], fila["rival"], fila["fecha"]),
                             []).append(fila)

    for ruta in sorted(carpeta_informes.glob("*.xlsx")) if carpeta_informes.is_dir() else []:
        if es_temporal(ruta.name):
            continue
        m = RE_NOMBRE_INFORME.match(ruta.name)
        if not m:
            continue
        equipo, rival, fecha = m.group("equipo"), m.group("rival"), m.group("fecha")
        candidatos = por_clave.get(_clave_partido(equipo, rival, fecha), [])
        pegado = False
        for fila in candidatos:
            # el que ya tiene informe no se pisa: si hay dos volcados del mismo
            # dia y un solo informe, el informe cuelga del primero
            if fila["informe"] is None:
                fila["informe"] = ruta.name
                pegado = True
                break
        if pegado:
            continue
        # llegar aca con un "_sin_formulas" (o cualquier otro sufijo) y con
        # volcados de ese partido significa que el informe bueno ya quedo
        # colgado de su fila: esta copia no es un partido aparte
        if m.group("extra") and candidatos:
            continue
        filas.append({
            "id": ruta.name,
            "fecha": fecha,
            "hora": "",
            "equipo": equipo,
            "rival": rival,
            "sets": "",
            "parciales": [],
            "puntos": 0,
            "volcado": None,
            "informe": ruta.name,
        })

    filas.sort(key=lambda f: (f["fecha"], f["hora"], f["id"]), reverse=True)
    return filas


# ======================================================================
# 5) UN PARTIDO
# ======================================================================

def leer_partido(ruta) -> dict:
    """El volcado entero, ya parseado, listo para mandar como JSON."""
    ruta = Path(ruta)
    texto = ruta.read_text(encoding="utf-8")
    volcado = _modulo_parser().parse_volcado(texto)

    resumen = resumen_volcado(texto)
    fecha, hora = _fecha_y_hora(ruta.name)
    partido = {
        "archivo": ruta.name,
        "fecha": fecha or "",
        "hora": hora or "",
        "equipo": resumen["equipo"] or volcado["team1_name"] or "?",
        "rival": resumen["rival"] or volcado["team2_name"] or "?",
        "sets": resumen["sets"] or "",
        "parciales": resumen["parciales"],
        "puntos": resumen["puntos"] or volcado["total_puntos_cargados"] or 0,
        # el orden A/B del volcado, para que el selector de equipo respete
        # cual es el local
        "orden_equipos": [n for n in (volcado["team1_name"], volcado["team2_name"])
                          if n in volcado["teams"]],
        "equipos": volcado["teams"],
    }
    partido.update(secciones_de_cancha(texto))
    # si el volcado nombra equipos que no tienen bloque de estadisticas, igual
    # se listan todos los que si lo tienen
    for nombre in volcado["teams"]:
        if nombre not in partido["orden_equipos"]:
            partido["orden_equipos"].append(nombre)
    return partido


# ======================================================================
# 6) UN INFORME
# ======================================================================

# Los colores con los que generar_informe_volley pinta cada tipo de fila. Es
# la forma mas fiel de distinguirlas: el que las escribio ya decidio cual es
# titulo y cual es total, no hace falta adivinarlo por el texto.
RELLENO_TITULO = "1F3864"
RELLENO_ENCABEZADO = "2E5C9A"
RELLENO_TOTAL = "F2F2F2"


def _relleno(celda) -> str:
    """El color de fondo en RRGGBB, sin el canal alfa.

    openpyxl devuelve ocho digitos y el alfa cambia segun quien escribio el
    archivo (00 el que genera el proyecto, FF si paso por Excel), asi que se
    compara solo el color."""
    try:
        if celda.fill.patternType != "solid":
            return ""
        rgb = celda.fill.fgColor.rgb
        return rgb[-6:].upper() if isinstance(rgb, str) else ""
    except AttributeError:
        return ""


def _texto_de_celda(celda) -> str:
    """El valor como se muestra, respetando el formato de la celda."""
    valor = celda.value
    if valor is None:
        return ""
    if isinstance(valor, str):
        return valor          # incluye la formula sin evaluar, a proposito
    if isinstance(valor, bool):
        return "si" if valor else "no"
    if isinstance(valor, (int, float)):
        formato = celda.number_format or ""
        if "%" in formato:
            return f"{valor * 100:.1f}%"
        if isinstance(valor, float):
            return str(int(valor)) if valor == int(valor) else f"{valor:.2f}"
        return str(valor)
    return str(valor)


def _tipo_de_fila(celdas, textos) -> str:
    relleno = _relleno(celdas[0]) if celdas else ""
    if not any(textos):
        return "vacia"
    if relleno == RELLENO_TITULO:
        return "titulo"
    if relleno == RELLENO_ENCABEZADO:
        # los subtitulos van combinados (una sola celda con texto), los
        # encabezados de tabla tienen un nombre por columna
        return "subtitulo" if not any(textos[1:]) else "encabezado"
    if relleno == RELLENO_TOTAL:
        return "total"
    try:
        if celdas and celdas[0].font is not None and celdas[0].font.italic:
            return "nota"
    except AttributeError:
        pass
    return "dato"


def _libro_con_valores(ruta):
    """Abre el .xlsx de forma que las celdas tengan numeros y no formulas.

    Primero se prueba con los valores que cacheo Excel; si el archivo nunca
    se abrio en Excel esas celdas vienen en None, y entonces se evalua con
    valores_excel, que es el evaluador que ya usa el proyecto para guardar el
    informe sin formulas. Devuelve (libro, formulas_que_no_se_pudieron)."""
    import openpyxl

    import valores_excel

    libro = openpyxl.load_workbook(ruta, data_only=False)
    formulas = [(hoja.title, celda.coordinate)
                for hoja in libro
                for fila in hoja.iter_rows()
                for celda in fila
                if isinstance(celda.value, str) and celda.value.startswith("=")]
    if not formulas:
        return libro, []

    cacheado = openpyxl.load_workbook(ruta, data_only=True)
    if all(cacheado[hoja][coord].value is not None for hoja, coord in formulas):
        return cacheado, []

    _, fallidas = valores_excel.convertir_a_valores(libro)
    return libro, fallidas


def leer_informe(ruta) -> dict:
    """El .xlsx como hojas, filas y celdas de texto.

    Cada fila viene con su tipo (titulo, subtitulo, encabezado, dato, total,
    nota) para que la web la pinte igual que el Excel sin volver a decidir
    nada."""
    if not HAY_OPENPYXL:
        raise FaltaOpenpyxl(MENSAJE_SIN_OPENPYXL)

    libro, fallidas = _libro_con_valores(Path(ruta))

    hojas = []
    for hoja in libro:
        filas = []
        ancho = 0
        for celdas in hoja.iter_rows():
            textos = [_texto_de_celda(c) for c in celdas]
            while textos and textos[-1] == "":
                textos.pop()
            tipo = _tipo_de_fila(celdas[:len(textos) or 1], textos)
            ancho = max(ancho, len(textos))
            filas.append({"tipo": tipo, "celdas": textos})
        while filas and filas[-1]["tipo"] == "vacia":
            filas.pop()
        hojas.append({
            "nombre": hoja.title,
            "oculta": hoja.sheet_state != "visible",
            "columnas": ancho,
            "filas": filas,
        })

    return {
        "archivo": Path(ruta).name,
        "hojas": hojas,
        # si alguna formula no se pudo evaluar se avisa: en la celda quedo la
        # formula a la vista, que es mejor que un cero inventado
        "avisos": [f"{len(fallidas)} formula(s) quedaron sin evaluar: {fallidas[0]}"]
        if fallidas else [],
    }
