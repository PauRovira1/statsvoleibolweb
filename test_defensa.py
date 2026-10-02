"""
Tests de las estadisticas de defensa.

Se corren con:
    python -m unittest test_defensa

La defensa ya estaba en los datos desde siempre -- el motor la parsea y la
describe en vivo -- pero no la contaba nadie. Lo que estos tests cuidan es lo
que puede salir mal callado al empezar a contarla:

* que se le atribuya al equipo equivocado, porque el bloque no guarda de quien
  es la pelota y hay que deducirlo repasando el punto desde el saque;
* que se mezcle con la recepcion, que se anota igual pero no es lo mismo;
* que el .txt y el Excel se desincronicen, o sea que el volcado diga una cosa
  y el informe otra.
"""
import textwrap
import unittest
from pathlib import Path

import analisis_voley as av
import generar_informe_volley as gi
import sesion_web
import test_carga_visual as t

VOLCADOS = sorted(Path(__file__).resolve().parent.joinpath("Datos").glob("*.txt"))


def partido(ruta):
    """Rehace un partido de verdad y devuelve su estado."""
    sesion = sesion_web.SesionPartido()
    for linea in t.lineas_del_volcado(ruta):
        sesion.enviar(linea)
    return sesion


class TestDeQuienEsLaDefensa(unittest.TestCase):
    """El bloque no guarda de quien es la pelota: se deduce repasando el punto
    desde el saque. Equivocarse ahi le da las defensas de un equipo al otro y
    no hay nada que lo delate."""

    def punto(self, equipo_saca, resultados):
        """Un punto armado a mano: el primer bloque es el saque."""
        jugadas = [{"resultado": resultados[0]}]
        for i, r in enumerate(resultados[1:], start=1):
            jugadas.append({"defensor": 10 + i, "calidad_defensa": 2, "resultado": r})
        return {"equipo_saca": equipo_saca, "jugadas": jugadas}

    def test_la_primera_defensa_es_del_que_saco(self):
        # A saca, B recibe y ataca, A defiende
        p = self.punto("A", ["D", "P"])
        self.assertEqual([e for _, e in av.defensas_del_punto(p)], ["A"])

    def test_los_equipos_se_alternan_mientras_la_pelota_sigue(self):
        p = self.punto("A", ["D", "D", "D", "P"])
        self.assertEqual([e for _, e in av.defensas_del_punto(p)], ["A", "B", "A"])

    def test_si_saca_el_otro_se_invierte(self):
        p = self.punto("B", ["D", "D", "P"])
        self.assertEqual([e for _, e in av.defensas_del_punto(p)], ["B", "A"])

    def test_un_punto_sin_continuacion_no_tiene_defensas(self):
        self.assertEqual(list(av.defensas_del_punto(self.punto("A", ["P"]))), [])

    def test_un_punto_sin_jugadas_no_rompe(self):
        vacio = {"equipo_saca": "A", "jugadas": []}
        self.assertEqual(list(av.defensas_del_punto(vacio)), [])


class TestElConteo(unittest.TestCase):

    @classmethod
    def setUpClass(cls):
        if not VOLCADOS:
            raise unittest.SkipTest("no hay volcados")
        cls.estado = partido(VOLCADOS[0]).estado
        cls.puntos = cls.estado["puntos"]

    def test_cuenta_todas_las_defensas_del_partido(self):
        # No es "todo bloque con defensor": lo que viene de un libre se anota
        # igual pero no es una defensa (ver TestQueCuentaComoDefensa). Lo que
        # se cuida aca es que no se pierda ni se invente ninguna de las que si
        # cuentan, recorriendo los puntos igual que el motor.
        crudas = sum(1 for p in self.puntos for _ in av.defensas_del_punto(p))
        calculadas = sum(x["total"]
                         for equipo in av.calcular_estadisticas_defensa(self.puntos).values()
                         for x in equipo.values())
        self.assertEqual(calculadas, crudas)
        self.assertGreater(crudas, 0)

    def test_no_se_mezcla_con_la_recepcion(self):
        # se anotan igual en la notacion pero son dos habilidades distintas
        defensa = av.calcular_estadisticas_defensa(self.puntos)
        recepcion = av.calcular_estadisticas_recepcion(self.puntos)
        total_def = sum(x["total"] for e in defensa.values() for x in e.values())
        total_rec = sum(x["total"] for e in recepcion.values() for x in e.values())
        self.assertGreater(total_def, 0)
        self.assertGreater(total_rec, 0)
        self.assertNotEqual(total_def, total_rec)

    def test_cada_jugador_reparte_su_total_entre_las_calidades(self):
        for equipo in av.calcular_estadisticas_defensa(self.puntos).values():
            for datos in equipo.values():
                suma = sum(c for c, _ in datos["calidades"].values())
                self.assertEqual(suma, datos["total"])

    def test_tiene_las_mismas_calidades_que_la_recepcion(self):
        # la pelota que no se llega a jugar es un 0 como cualquier otro: no
        # tiene una calidad propia, cierra el punto y eso se ve en la causa
        self.assertEqual(av.CALIDADES_DEFENSA, (-1, 0, 1, 2, 3))
        self.assertNotIn(-2, av.CALIDADES_DEFENSA)

    def test_la_que_no_se_pudo_jugar_cuenta_como_calidad_0(self):
        punto = {"equipo_saca": "A", "jugadas": [
            av.parsear_bloque_saque("5_1_6_X/3_3/2_4/4_1_D"),
            av.parsear_bloque_defensa("7_0"),
        ]}
        stats = av.calcular_estadisticas_defensa([punto])
        cantidad, _ = stats["A"][7]["calidades"][0]
        self.assertEqual(cantidad, 1)
        self.assertEqual(stats["A"][7]["total"], 1)

    def test_la_suma_de_los_sets_da_el_partido(self):
        por_set = av.calcular_estadisticas_defensa_por_set(self.puntos)
        entero = av.calcular_estadisticas_defensa(self.puntos)
        for equipo in ("A", "B"):
            del_partido = sum(x["total"] for x in entero[equipo].values())
            de_los_sets = sum(x["total"] for datos in por_set.values()
                              for x in datos[equipo].values())
            self.assertEqual(de_los_sets, del_partido, equipo)


class TestElVolcadoYElExcel(unittest.TestCase):
    """El .txt es la fuente y el Excel se arma leyendolo: si el parser no
    entiende la seccion, el informe sale en cero sin ningun error."""

    @classmethod
    def setUpClass(cls):
        if not VOLCADOS:
            raise unittest.SkipTest("no hay volcados")
        estado = partido(VOLCADOS[0]).estado
        cls.texto = av.formatear_estadisticas(
            estado["puntos"], estado["nombres"],
            {k: set(v) for k, v in estado["armadores"].items()})
        cls.puntos = estado["puntos"]

    def test_el_volcado_trae_la_seccion(self):
        self.assertIn("Defensas por jugador:", self.texto)

    def test_el_parser_la_entiende(self):
        volcado = gi.parse_volcado(self.texto)
        equipo = next(iter(volcado["teams"]))
        self.assertTrue(volcado["teams"][equipo]["defensas"])

    def test_lo_parseado_es_lo_calculado(self):
        volcado = gi.parse_volcado(self.texto)
        equipo = next(iter(volcado["teams"]))
        del_texto = sum(sum(v.values())
                        for v in volcado["teams"][equipo]["defensas"].values())
        calculado = sum(x["total"] for x in
                        av.calcular_estadisticas_defensa(self.puntos)["A"].values())
        self.assertEqual(del_texto, calculado)

    def test_las_calidades_son_las_mismas_que_en_recepcion(self):
        volcado = gi.parse_volcado(self.texto)
        equipo = next(iter(volcado["teams"]))
        for v in volcado["teams"][equipo]["defensas"].values():
            self.assertEqual(set(v), {"cal3", "cal2", "cal1", "cal0", "pase"})

    def test_un_volcado_viejo_suma_la_perdida_a_la_calidad_0(self):
        """Los .txt guardados antes traen esas pelotas en una fila propia.

        Si el parser la descarta, el informe de un partido viejo sale con
        menos defensas de las que se cargaron y nadie se entera."""
        viejo = textwrap.dedent("""            Defensas por jugador:
              Jugador 7: 4 defensas
                  Defensa perdida: 1 (25.0%)
                  Pase al otro lado: 1 (25.0%)
                  Calidad 0: 1 (25.0%)
                  Calidad 1: 1 (25.0%)
                  Calidad 2: 0 (0.0%)
                  Calidad 3: 0 (0.0%)
            """)
        defensas = gi.parse_volcado(self.texto.replace(
            self.texto[self.texto.index("Defensas por jugador:"):], viejo))
        jugador = defensas["teams"][next(iter(defensas["teams"]))]["defensas"]["Jugador 7"]
        self.assertEqual(jugador["cal0"], 2)          # la perdida vieja y el 0
        self.assertEqual(sum(jugador.values()), 4)    # el total no cambia

    def test_el_excel_tiene_su_hoja_con_datos(self):
        try:
            import openpyxl                      # noqa: F401
        except ImportError:
            self.skipTest("necesita openpyxl")
        import valores_excel
        volcado = gi.parse_volcado(self.texto)
        equipo, rival = list(volcado["teams"])[:2]
        libro, _ = gi.build_workbook(equipo, rival, volcado)
        self.assertIn("Defensa", libro.sheetnames)

        valores_excel.convertir_a_valores(libro)
        hoja = libro["Defensa"]
        totales = [hoja.cell(row=f, column=2).value for f in range(1, hoja.max_row + 1)
                   if hoja.cell(row=f, column=1).value == "TOTAL"]
        self.assertTrue(totales)
        self.assertGreater(totales[0], 0)

    def test_el_total_del_excel_es_el_del_volcado(self):
        try:
            import openpyxl                      # noqa: F401
        except ImportError:
            self.skipTest("necesita openpyxl")
        import valores_excel
        volcado = gi.parse_volcado(self.texto)
        equipo, rival = list(volcado["teams"])[:2]
        del_texto = sum(sum(v.values())
                        for v in volcado["teams"][equipo]["defensas"].values())
        libro, _ = gi.build_workbook(equipo, rival, volcado)
        valores_excel.convertir_a_valores(libro)
        hoja = libro["Defensa"]
        primer_total = next(hoja.cell(row=f, column=2).value
                            for f in range(1, hoja.max_row + 1)
                            if hoja.cell(row=f, column=1).value == "TOTAL")
        self.assertEqual(primer_total, del_texto)


class TestQueCuentaComoDefensa(unittest.TestCase):
    """Defender es levantar un ATAQUE.

    Un libre no es un ataque: el rival decidio no atacar y mando una pelota
    facil. Se anota igual que una defensa (Y_C) y por eso se contaba, pero
    contarlo infla la defensa de todos por igual y hace parecer mejor a quien
    mas pelotas faciles recibio."""

    def punto(self, resultados, equipo_saca="A"):
        """El primer bloque es el saque; los demas llevan defensor."""
        jugadas = [{"resultado": resultados[0]}]
        for i, r in enumerate(resultados[1:], start=1):
            jugadas.append({"defensor": 10 + i, "calidad_defensa": 2, "resultado": r})
        return {"equipo_saca": equipo_saca, "jugadas": jugadas}

    def cuantas(self, punto):
        return len(list(av.defensas_del_punto(punto)))

    def test_lo_que_viene_de_un_ataque_defendido_cuenta(self):
        self.assertEqual(self.cuantas(self.punto(["D", "P"])), 1)

    def test_lo_que_viene_de_un_libre_no_cuenta(self):
        self.assertEqual(self.cuantas(self.punto(["F", "P"])), 0)

    def test_lo_que_viene_de_un_bloqueo_rejugable_no_cuenta(self):
        # la pelota rebota en el bloqueo y vuelve al lado del que ataco: el
        # que la levanta esta recuperando la suya, no defendiendo nada
        self.assertEqual(self.cuantas(self.punto(["R", "P"])), 0)

    def test_el_rejugable_deja_la_pelota_del_mismo_lado(self):
        """A diferencia del libre, el rejugable NO cambia de lado.

        A saca, B ataca, el bloqueo de A la devuelve al lado de B (R), B la
        recupera -- eso no se cuenta -- y vuelve a atacar; esa segunda la
        defiende A. Si el rejugable hiciera cambiar de lado, esa defensa
        quedaria anotada a nombre de B."""
        equipos = [e for _, e in av.defensas_del_punto(self.punto(["R", "D", "P"]))]
        self.assertEqual(equipos, ["A"])

    def test_en_un_punto_largo_se_saltean_solo_los_del_libre(self):
        # ataque defendido, X, ataque defendido: dos defensas, no tres
        for medio in ("F", "R", "T", "V"):
            self.assertEqual(self.cuantas(self.punto(["D", medio, "D", "P"])), 2, medio)

    def test_el_libre_no_rompe_de_quien_es_la_pelota(self):
        # aunque no se cuente, la pelota igual cambio de lado: si no, las
        # defensas siguientes se le adjudicarian al equipo equivocado
        equipos = [e for _, e in av.defensas_del_punto(self.punto(["D", "F", "D", "P"]))]
        self.assertEqual(equipos, ["A", "A"])

    def test_lo_que_viene_de_un_toque_no_cuenta(self):
        # la notacion dice literalmente "toca la pelota sin atacar"
        self.assertEqual(self.cuantas(self.punto(["T", "P"])), 0)

    def test_lo_que_viene_de_un_overpass_no_cuenta(self):
        # la recepcion se fue derecho al otro lado: nadie ataco
        self.assertEqual(self.cuantas(self.punto(["V", "P"])), 0)

    def test_solo_cuenta_lo_que_sigue_a_un_ataque_defendido(self):
        # la regla entera en una linea. Si alguien agrega un resultado nuevo,
        # no entra como defensa sin decidirlo.
        self.assertEqual(av.TRAS_UN_ATAQUE, ("D",))

    def test_el_toque_y_el_overpass_igual_cambian_de_lado(self):
        # no se cuentan, pero la pelota cruzo: si se dejara de alternar, la
        # defensa siguiente quedaria a nombre del equipo equivocado
        for medio in ("T", "V", "F"):
            equipos = [e for _, e in av.defensas_del_punto(self.punto(["D", medio, "D", "P"]))]
            self.assertEqual(equipos, ["A", "A"], medio)

    def test_en_los_volcados_de_verdad_bajan_las_defensas(self):
        if not VOLCADOS:
            self.skipTest("no hay volcados")
        puntos = partido(VOLCADOS[0]).estado["puntos"]
        contadas = sum(x["total"]
                       for e in av.calcular_estadisticas_defensa(puntos).values()
                       for x in e.values())
        con_libre = sum(1 for p in puntos for i, b in enumerate(p["jugadas"][1:], start=1)
                        if b.get("defensor") is not None
                        and b.get("calidad_defensa") is not None)
        self.assertLess(contadas, con_libre)


class TestTodosLosVolcadosGuardados(unittest.TestCase):

    def test_todos_traen_la_seccion(self):
        if not VOLCADOS:
            self.skipTest("no hay volcados")
        sin = [r.name for r in VOLCADOS
               if "Defensas por jugador:" not in r.read_text(encoding="utf-8")]
        self.assertEqual(sin, [], "estos volcados hay que regenerarlos")

    def test_el_parser_los_lee_a_todos(self):
        if not VOLCADOS:
            self.skipTest("no hay volcados")
        for ruta in VOLCADOS:
            volcado = gi.parse_volcado(ruta.read_text(encoding="utf-8"))
            equipo = next(iter(volcado["teams"]))
            self.assertIn("defensas", volcado["teams"][equipo], ruta.name)


if __name__ == "__main__":
    unittest.main()
