"""
Tests para analisis_voley.py.

Se corren con:
    python -m unittest test_analisis_voley.py -v
"""

import io
import os
import re

from pathlib import Path

import analisis_voley as av
import sesion_web
import unittest
from unittest.mock import patch

from analisis_voley import (
    CONTRASENA_CARGA,
    parsear_bloque_saque,
    parsear_bloque_defensa,
    calcular_estadisticas_armado,
    calcular_estadisticas_recepcion,
    calcular_estadisticas_recepcion_por_set,
    calcular_estadisticas_ataque,
    calcular_armado_por_calidad_recepcion,
    calcular_recepcion_por_tipo_saque,
    calcular_estadisticas_bloqueo,
    cargar_jugadas,
    calcular_estadisticas_armado_por_set,
    calcular_armado_por_armador,
    calcular_armado_por_armador_por_set,
    calcular_armado_por_calidad_recepcion_por_armador,
    parsear_rotacion,
    jugador_que_saca,
    rotar,
    rotacion_en_cancha,
    veces_que_roto,
    fase_del_punto,
    calcular_puntos_por_fase,
    FASES_RALLY,
    causa_del_punto,
    calcular_puntos_por_causa,
    zona_del_armador,
    calcular_puntos_por_zona_armador,
    CAUSAS_GANADAS,
    CAUSAS_ERROR,
    otro_equipo,
    aplicar_cambio,
    copiar_rotaciones,
    jugar_punto,
    guardar_reporte_txt,
    preguntar_equipo,
    formatear_estadisticas,
    preguntar_si_no,
    generar_informe_excel,
)

import almacenamiento as alm

_token_de_verdad = alm.token_blob


def setUpModule():
    """Ningun test de este archivo habla con el Blob.

    Varios guardan un reporte de verdad (guardar_reporte_txt), y guardar
    ahora publica el archivo. Con BLOB_READ_WRITE_TOKEN en el entorno -- que
    es lo normal en la maquina del que ademas despliega esto -- correr las
    pruebas subiria partidos de prueba al store de produccion."""
    alm.token_blob = lambda: ""


def tearDownModule():
    alm.token_blob = _token_de_verdad


class TestParsearBloqueSaque(unittest.TestCase):

    def test_as(self):
        bloque = parsear_bloque_saque("5_1_6_A")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["sacador"], 5)
        self.assertEqual(bloque["zona_saque"], "1")
        self.assertEqual(bloque["zona_destino_saque"], "6")
        self.assertEqual(bloque["resultado_saque"], "A")
        self.assertIsNone(bloque["receptor"])

    def test_error_de_saque(self):
        bloque = parsear_bloque_saque("5_1_6_E")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["resultado_saque"], "E")

    def test_jugada_completa_punto(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_3/2_4/4_1_P")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["resultado_saque"], "X")
        self.assertEqual(bloque["receptor"], 3)
        self.assertEqual(bloque["calidad_recepcion"], 3)
        self.assertEqual(bloque["colocador"], 2)
        self.assertEqual(bloque["zona_colocacion"], "4")
        self.assertEqual(bloque["atacante"], 4)
        self.assertEqual(bloque["zona_ataque"], "1")
        self.assertEqual(bloque["resultado"], "P")

    def test_calidad_recepcion_es_entero(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_0/2_4/4_1_D")
        self.assertEqual(bloque["calidad_recepcion"], 0)
        self.assertIsInstance(bloque["calidad_recepcion"], int)

    def test_jugador_con_varios_digitos(self):
        bloque = parsear_bloque_saque("12_1_6_X/10_2/11_4/13_5_D")
        self.assertEqual(bloque["sacador"], 12)
        self.assertEqual(bloque["receptor"], 10)
        self.assertEqual(bloque["colocador"], 11)
        self.assertEqual(bloque["atacante"], 13)

    def test_formato_invalido(self):
        self.assertIsNone(parsear_bloque_saque("esto no es una jugada"))
        self.assertIsNone(parsear_bloque_saque("5_1_6_Z"))
        self.assertIsNone(parsear_bloque_saque("5_1_6_X/3_9/2_4/4_1_P"))  # calidad fuera de rango
        self.assertIsNone(parsear_bloque_saque("5_2_6_A"))  # zona de saque invalida (2 no permitido)


class TestParsearBloqueDefensa(unittest.TestCase):

    def test_defensa_calidad_0_no_termina_el_punto(self):
        # antes "7_0" solo (sin armado/ataque) significaba defensa perdida;
        # ahora la calidad 0 es solo informativa y hace falta la cadena completa.
        bloque = parsear_bloque_defensa("7_0/1_5/9_5_D")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["defensor"], 7)
        self.assertEqual(bloque["calidad_defensa"], 0)
        self.assertIsInstance(bloque["calidad_defensa"], int)
        self.assertEqual(bloque["colocador"], 1)
        self.assertEqual(bloque["atacante"], 9)
        self.assertEqual(bloque["resultado"], "D")

    def test_defensa_completa(self):
        bloque = parsear_bloque_defensa("7_2/1_5/9_5_D")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["defensor"], 7)
        self.assertEqual(bloque["calidad_defensa"], 2)
        self.assertEqual(bloque["colocador"], 1)
        self.assertEqual(bloque["zona_colocacion"], "5")
        self.assertEqual(bloque["atacante"], 9)
        self.assertEqual(bloque["zona_ataque"], "5")
        self.assertEqual(bloque["resultado"], "D")

    def test_formato_invalido(self):
        self.assertIsNone(parsear_bloque_defensa("7_0"))  # ya no existe el atajo de "defensa perdida"
        self.assertIsNone(parsear_bloque_defensa("7_4/1_5/9_5_D"))  # calidad fuera de rango (0-3)
        self.assertIsNone(parsear_bloque_defensa("7_2/1_5/9_5"))  # falta resultado


class TestResultadoAtaqueConBloqueo(unittest.TestCase):

    def test_bloqueo_punto(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_3/2_4/4_1_B_6_P")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["resultado"], "B")
        self.assertEqual(bloque["jugador_bloqueo"], 6)

    def test_toque_de_bloqueo_usado(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_3/2_4/4_1_U_6")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["resultado"], "U")
        self.assertEqual(bloque["jugador_bloqueo"], 6)

    def test_bloqueo_rejugable(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_3/2_4/4_1_R_6")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["resultado"], "R")
        self.assertEqual(bloque["jugador_bloqueo"], 6)

    def test_jugador_de_bloqueo_multidigito(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_3/2_4/4_1_B_12_P")
        self.assertEqual(bloque["jugador_bloqueo"], 12)

    def test_defensa_con_bloqueo(self):
        bloque = parsear_bloque_defensa("7_2/1_5/9_5_R_10")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["resultado"], "R")
        self.assertEqual(bloque["jugador_bloqueo"], 10)

    def test_resultado_bloqueo_invalido(self):
        self.assertIsNone(parsear_bloque_saque("5_1_6_X/3_3/2_4/4_1_B_6"))  # falta el _P final
        self.assertIsNone(parsear_bloque_saque("5_1_6_X/3_3/2_4/4_1_Z_6"))  # letra invalida
        self.assertIsNone(parsear_bloque_saque("5_1_6_X/3_3/2_4/4_1_U_"))   # falta numero de jugador

    def test_resultado_sin_jugador_de_bloqueo(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_3/2_4/4_1_P")
        self.assertIsNone(bloque["jugador_bloqueo"])
        bloque = parsear_bloque_saque("5_1_6_X/3_3/2_4/4_1_D")
        self.assertIsNone(bloque["jugador_bloqueo"])


class TestJugarPuntoConBloqueo(unittest.TestCase):

    def test_bloqueo_punto_para_el_equipo_que_defiende(self):
        with patch("builtins.input", side_effect=["5_1_6_X/3_3/2_4/4_1_B_7_P"]):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "A")  # A saca, B ataca, A bloquea -> punto para A
        self.assertEqual(len(secuencia), 1)
        self.assertEqual(secuencia[0]["jugador_bloqueo"], 7)

    def test_toque_de_bloqueo_punto_para_el_que_ataco(self):
        with patch("builtins.input", side_effect=["5_1_6_X/3_3/2_4/4_1_U_7"]):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "B")  # B ataco y el toque de bloqueo no cambia el punto

    def test_bloqueo_rejugable_mismo_equipo_sigue_atacando(self):
        entradas = [
            "5_1_6_X/3_3/2_4/4_1_R_7",  # B ataca, bloqueo rejugable
            "8_2/1_5/9_5_P",             # B (sin cambiar de equipo) recupera y hace punto
        ]
        with patch("builtins.input", side_effect=entradas):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "B")
        self.assertEqual(len(secuencia), 2)
        # el bloqueo rejugable no cambia el equipo que arma/ataca (sigue B)
        self.assertEqual(secuencia[1]["equipo_set"], "B")

    def test_defendido_normal_si_cambia_el_equipo(self):
        entradas = [
            "5_1_6_X/3_3/2_4/4_1_D",    # B ataca, A defiende
            "7_2/1_5/9_5_P",             # ahora A arma y ataca (cambio de equipo)
        ]
        with patch("builtins.input", side_effect=entradas):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "A")
        self.assertEqual(secuencia[1]["equipo_set"], "A")


class TestBolaLibre(unittest.TestCase):

    def test_parseo_libre_en_bloque_inicial(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_3/2_4/9_F_8")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["resultado"], "F")
        self.assertEqual(bloque["atacante"], 9)
        self.assertEqual(bloque["zona_ataque"], "8")
        self.assertIsNone(bloque["jugador_bloqueo"])

    def test_parseo_libre_en_bloque_defensa(self):
        bloque = parsear_bloque_defensa("7_2/1_5/9_F_3")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["resultado"], "F")
        self.assertEqual(bloque["atacante"], 9)
        self.assertEqual(bloque["zona_ataque"], "3")

    def test_libre_permite_las_9_zonas(self):
        for zona in "123456789":
            bloque = parsear_bloque_saque(f"5_1_6_X/3_3/2_4/9_F_{zona}")
            self.assertIsNotNone(bloque, f"deberia aceptar la zona {zona}")
            self.assertEqual(bloque["zona_ataque"], zona)

    def test_ataque_normal_no_confunde_con_libre(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_3/2_4/9_1_P")
        self.assertEqual(bloque["resultado"], "P")
        self.assertEqual(bloque["atacante"], 9)
        self.assertEqual(bloque["zona_ataque"], "1")

    def test_libre_invalido(self):
        self.assertIsNone(parsear_bloque_saque("5_1_6_X/3_3/2_4/9_F_0"))   # zona 0 no existe
        self.assertIsNone(parsear_bloque_saque("5_1_6_X/3_3/2_4/9_F"))     # falta la zona

    def test_libre_siempre_cambia_de_equipo_y_sigue_el_punto(self):
        entradas = [
            "5_1_6_X/3_3/2_4/9_F_8",   # B no puede atacar, hace libre
            "7_2/1_5/9_1_P",            # A recibe el libre y hace punto
        ]
        with patch("builtins.input", side_effect=entradas):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "A")
        self.assertEqual(len(secuencia), 2)
        self.assertEqual(secuencia[0]["resultado"], "F")
        # el libre pasa el control al otro equipo (A), igual que una defensa
        self.assertEqual(secuencia[1]["equipo_set"], "A")

    def test_libre_encadenado_con_otro_libre(self):
        entradas = [
            "5_1_6_X/3_3/2_4/9_F_8",   # B hace libre
            "7_2/1_5/9_F_4",            # A tampoco puede atacar, otro libre
            "6_1/2_3/1_5_P",            # B recibe y hace punto
        ]
        with patch("builtins.input", side_effect=entradas):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "B")
        self.assertEqual(len(secuencia), 3)


class TestToque(unittest.TestCase):

    def test_parseo_en_bloque_inicial(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_3/2_4/9_T_8")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["resultado"], "T")
        self.assertEqual(bloque["atacante"], 9)
        self.assertEqual(bloque["zona_ataque"], "8")
        self.assertIsNone(bloque["jugador_bloqueo"])
        self.assertFalse(bloque["es_segunda"])

    def test_parseo_en_bloque_defensa(self):
        bloque = parsear_bloque_defensa("7_2/1_5/9_T_3")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["resultado"], "T")
        self.assertEqual(bloque["atacante"], 9)
        self.assertEqual(bloque["zona_ataque"], "3")

    def test_admite_las_9_zonas(self):
        for zona in "123456789":
            bloque = parsear_bloque_saque(f"5_1_6_X/3_3/2_4/9_T_{zona}")
            self.assertIsNotNone(bloque, f"deberia aceptar la zona {zona}")
            self.assertEqual(bloque["zona_ataque"], zona)

    def test_no_se_confunde_con_libre_ni_ataque_normal(self):
        toque = parsear_bloque_saque("5_1_6_X/3_3/2_4/9_T_8")
        libre = parsear_bloque_saque("5_1_6_X/3_3/2_4/9_F_8")
        ataque = parsear_bloque_saque("5_1_6_X/3_3/2_4/9_1_P")
        self.assertEqual(toque["resultado"], "T")
        self.assertEqual(libre["resultado"], "F")
        self.assertEqual(ataque["resultado"], "P")

    def test_formato_invalido(self):
        self.assertIsNone(parsear_bloque_saque("5_1_6_X/3_3/2_4/9_T_0"))  # zona 0 no existe
        self.assertIsNone(parsear_bloque_saque("5_1_6_X/3_3/2_4/9_T"))    # falta la zona

    def test_toque_cambia_de_equipo_y_sigue_el_punto(self):
        entradas = [
            "5_1_6_X/3_3/2_4/9_T_8",   # B toca, no ataca
            "7_2/1_5/9_1_P",            # A recibe y hace punto
        ]
        with patch("builtins.input", side_effect=entradas):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "A")
        self.assertEqual(len(secuencia), 2)
        self.assertEqual(secuencia[0]["resultado"], "T")
        self.assertEqual(secuencia[1]["equipo_set"], "A")

    def test_no_cuenta_como_ataque(self):
        puntos = [
            _punto("B", "A", [
                {**parsear_bloque_saque("3_5_4_X/1_3/2_4/9_T_8"), "equipo_atacante": "A"},
            ]),
        ]
        stats = calcular_estadisticas_ataque(puntos)
        self.assertEqual(stats["A"], {})


class TestAtaqueAfuera(unittest.TestCase):

    def test_parseo_afuera_en_bloque_inicial(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_3/2_4/4_1_O")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["resultado"], "O")
        self.assertIsNone(bloque["jugador_bloqueo"])

    def test_parseo_afuera_en_bloque_defensa(self):
        bloque = parsear_bloque_defensa("7_2/1_5/9_5_O")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["resultado"], "O")

    def test_afuera_punto_para_la_defensa(self):
        with patch("builtins.input", side_effect=["5_1_6_X/3_3/2_4/4_1_O"]):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "A")  # A saca, B ataca y se va afuera -> punto para A
        self.assertEqual(len(secuencia), 1)

    def test_afuera_tras_un_intercambio(self):
        entradas = [
            "5_1_6_X/3_3/2_4/4_1_D",   # B ataca, A defiende
            "7_2/1_5/9_5_O",            # A contraataca y se va afuera -> punto para B
        ]
        with patch("builtins.input", side_effect=entradas):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "B")


class TestAtaqueMalla(unittest.TestCase):

    def test_parseo_malla_en_bloque_inicial(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_3/2_4/4_1_M")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["resultado"], "M")
        self.assertIsNone(bloque["jugador_bloqueo"])

    def test_parseo_malla_en_bloque_defensa(self):
        bloque = parsear_bloque_defensa("7_2/1_5/9_5_M")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["resultado"], "M")

    def test_malla_punto_para_la_defensa(self):
        with patch("builtins.input", side_effect=["5_1_6_X/3_3/2_4/4_1_M"]):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "A")  # A saca, B ataca y va a la malla -> punto para A

    def test_malla_cuenta_como_fuera_en_estadisticas_de_ataque(self):
        puntos = [
            _punto("B", "A", [
                {**parsear_bloque_saque("3_5_4_X/1_3/2_4/9_1_M"), "equipo_atacante": "A"},
            ]),
        ]
        stats = calcular_estadisticas_ataque(puntos)
        self.assertEqual(stats["A"][9]["fuera"], 1)


class TestPasadaDeSegunda(unittest.TestCase):
    """Formato nuevo: la segunda NO lleva armado por separado (Y_C/A_S_Z_R2),
    porque ella misma ES el segundo toque."""

    def test_parseo_en_bloque_inicial(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_3/9_S_6_D")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["receptor"], 3)
        self.assertEqual(bloque["atacante"], 9)
        self.assertEqual(bloque["zona_ataque"], "6")
        self.assertEqual(bloque["resultado"], "D")
        self.assertTrue(bloque["es_segunda"])
        self.assertIsNone(bloque["colocador"])
        self.assertIsNone(bloque["zona_colocacion"])

    def test_parseo_en_bloque_defensa(self):
        bloque = parsear_bloque_defensa("7_2/9_S_4_P")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["defensor"], 7)
        self.assertEqual(bloque["atacante"], 9)
        self.assertEqual(bloque["zona_ataque"], "4")
        self.assertEqual(bloque["resultado"], "P")
        self.assertTrue(bloque["es_segunda"])
        self.assertIsNone(bloque["colocador"])

    def test_ya_no_admite_armado_intermedio(self):
        # el formato viejo (con W_Z3 antes de la segunda) ya no es valido
        self.assertIsNone(parsear_bloque_saque("5_1_6_X/3_3/2_4/9_S_6_D"))
        self.assertIsNone(parsear_bloque_defensa("7_2/1_5/9_S_4_P"))

    def test_admite_las_9_zonas(self):
        for zona in "123456789":
            bloque = parsear_bloque_saque(f"5_1_6_X/3_3/9_S_{zona}_D")
            self.assertIsNotNone(bloque, f"deberia aceptar la zona {zona}")

    def test_solo_admite_d_p_o_como_resultado(self):
        self.assertIsNone(parsear_bloque_saque("5_1_6_X/3_3/9_S_4_B"))
        self.assertIsNone(parsear_bloque_saque("5_1_6_X/3_3/9_S_4_M"))

    def test_ataque_normal_no_tiene_es_segunda(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_3/2_4/9_1_P")
        self.assertFalse(bloque["es_segunda"])

    def test_punto_directo_de_segunda(self):
        with patch("builtins.input", side_effect=["5_1_6_X/3_3/9_S_6_P"]):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "B")  # B recibe y hace la segunda -> punto para B

    def test_segunda_como_continuacion_cambia_de_equipo(self):
        entradas = [
            "5_1_6_X/3_3/2_4/4_1_D",  # B ataca, A defiende
            "2_2/10_S_6_D",              # A hace la segunda directo (sin armado)
            "7_1/1_4/9_5_P",             # B recibe y hace punto
        ]
        with patch("builtins.input", side_effect=entradas):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "B")
        self.assertEqual(len(secuencia), 3)
        self.assertTrue(secuencia[1]["es_segunda"])
        self.assertIsNone(secuencia[1]["equipo_set"])  # no hubo armado, no cuenta

    def test_no_cuenta_en_estadisticas_de_ataque(self):
        puntos = [
            _punto("B", "A", [
                {**parsear_bloque_saque("3_5_4_X/1_3/9_S_6_P"), "equipo_atacante": "A"},
            ]),
        ]
        stats = calcular_estadisticas_ataque(puntos)
        self.assertEqual(stats["A"], {})


class TestErrorEnJuego(unittest.TestCase):

    def test_f_en_el_saque_da_el_punto_al_receptor(self):
        with patch("builtins.input", side_effect=["f"]):
            ganador, secuencia, entradas = jugar_punto("A")
        self.assertEqual(ganador, "B")
        self.assertEqual(secuencia, [])
        self.assertEqual(entradas, ["f"])

    def test_f_mayuscula_tambien_funciona(self):
        with patch("builtins.input", side_effect=["F"]):
            ganador, secuencia, entradas = jugar_punto("A")
        self.assertEqual(ganador, "B")

    def test_f_a_mitad_del_punto_da_el_punto_al_otro_equipo(self):
        entradas = [
            "5_1_6_X/3_3/2_4/4_1_D",  # B ataca, A defiende
            "f",                         # A comete una falta mientras jugaba
        ]
        with patch("builtins.input", side_effect=entradas):
            ganador, secuencia, entradas_crudas = jugar_punto("A")
        self.assertEqual(ganador, "B")  # el punto es para el equipo contrario de quien fallo (A)
        self.assertEqual(entradas_crudas, ["5_1_6_X/3_3/2_4/4_1_D", "f"])


class TestArmadoOpcional(unittest.TestCase):

    def test_sin_marcador_armado_valido(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_3/2_4/4_1_P")
        self.assertTrue(bloque["armado_valido"])

    def test_marcador_x_en_bloque_inicial(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_3/2_4_X/4_1_P")
        self.assertIsNotNone(bloque)
        self.assertFalse(bloque["armado_valido"])
        self.assertEqual(bloque["colocador"], 2)
        self.assertEqual(bloque["zona_colocacion"], "4")
        self.assertEqual(bloque["resultado"], "P")  # el resto de la jugada no se ve afectado

    def test_marcador_x_en_bloque_defensa(self):
        bloque = parsear_bloque_defensa("7_2/1_5_X/9_5_P")
        self.assertIsNotNone(bloque)
        self.assertFalse(bloque["armado_valido"])
        self.assertEqual(bloque["colocador"], 1)
        self.assertEqual(bloque["zona_colocacion"], "5")

    def test_marcador_x_con_libre(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_3/2_4_X/9_F_8")
        self.assertFalse(bloque["armado_valido"])
        self.assertEqual(bloque["resultado"], "F")

    def test_armado_no_valido_no_cuenta_en_estadisticas(self):
        puntos = [
            _punto("B", "A", [
                {**parsear_bloque_saque("3_5_4_X/1_3/2_4_X/9_1_P"), "equipo_set": "A"},
            ]),
        ]
        stats = calcular_estadisticas_armado(puntos)
        total_a = sum(cantidad for cantidad, _ in stats["A"].values())
        self.assertEqual(total_a, 0)

    def test_armado_valido_si_cuenta_junto_a_uno_no_valido(self):
        puntos = [
            _punto("B", "A", [
                {**parsear_bloque_saque("3_5_4_X/1_3/2_4/9_1_D"), "equipo_set": "A"},
                {**parsear_bloque_defensa("7_2/1_5_X/9_5_P"), "equipo_set": "B"},
            ]),
        ]
        stats = calcular_estadisticas_armado(puntos)
        total_a = sum(cantidad for cantidad, _ in stats["A"].values())
        total_b = sum(cantidad for cantidad, _ in stats["B"].values())
        self.assertEqual(total_a, 1)  # el armado normal si cuenta
        self.assertEqual(total_b, 0)  # el marcado con _X no cuenta


class TestOverpassEnRecepcion(unittest.TestCase):

    def test_parseo_overpass(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_-1")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["receptor"], 3)
        self.assertEqual(bloque["calidad_recepcion"], -1)
        self.assertEqual(bloque["resultado"], "V")
        self.assertIsNone(bloque["colocador"])
        self.assertIsNone(bloque["atacante"])

    def test_overpass_no_admite_continuacion(self):
        # si hubo overpass la jugada termina ahi, no puede seguir con armado/ataque
        self.assertIsNone(parsear_bloque_saque("5_1_6_X/3_-1/2_4/4_1_P"))

    def test_overpass_cambia_de_equipo_y_sigue_el_punto(self):
        entradas = [
            "5_1_6_X/3_-1",   # B recibe pero se va directo al otro lado
            "7_2/1_5/9_1_P",   # A recibe ese overpass y hace punto
        ]
        with patch("builtins.input", side_effect=entradas):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "A")
        self.assertEqual(len(secuencia), 2)
        self.assertEqual(secuencia[1]["equipo_set"], "A")

    def test_overpass_cuenta_en_estadisticas_de_recepcion(self):
        puntos = [
            _punto("A", "A", [{**parsear_bloque_saque("5_1_6_X/3_-1"), "equipo_receptor": "B"}]),
        ]
        stats = calcular_estadisticas_recepcion(puntos)
        self.assertEqual(stats["B"][3]["total"], 1)
        cantidad, porcentaje = stats["B"][3]["calidades"][-1]
        self.assertEqual(cantidad, 1)
        self.assertAlmostEqual(porcentaje, 100.0)


class TestOverpassEnDefensa(unittest.TestCase):

    def test_parseo_overpass_en_defensa(self):
        bloque = parsear_bloque_defensa("7_-1")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["defensor"], 7)
        self.assertEqual(bloque["calidad_defensa"], -1)
        self.assertEqual(bloque["resultado"], "V")
        self.assertIsNone(bloque["colocador"])
        self.assertIsNone(bloque["atacante"])

    def test_overpass_en_defensa_no_admite_continuacion(self):
        self.assertIsNone(parsear_bloque_defensa("7_-1/1_5/9_5_P"))

    def test_overpass_en_defensa_cambia_de_equipo_y_sigue_el_punto(self):
        entradas = [
            "5_1_6_X/3_3/2_4/4_1_D",  # B ataca, A defiende
            "7_-1",                     # la defensa de A se va directo al otro lado
            "6_1/2_3/1_5_P",            # B recibe ese overpass y hace punto
        ]
        with patch("builtins.input", side_effect=entradas):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "B")
        self.assertEqual(len(secuencia), 3)
        self.assertEqual(secuencia[1]["resultado"], "V")
        self.assertEqual(secuencia[2]["equipo_set"], "B")


class TestEntradasCrudasYReporte(unittest.TestCase):

    def test_jugar_punto_devuelve_las_entradas_tal_cual_se_tipearon(self):
        with patch("builtins.input", side_effect=["5_1_6_A"]):
            ganador, secuencia, entradas = jugar_punto("A")
        self.assertEqual(entradas, ["5_1_6_A"])

    def test_entradas_invalidas_no_se_incluyen(self):
        entradas_tipeadas = [
            "esto no es valido",
            "5_1_6_X/3_3/2_4/4_1_D",
            "formato mal",
            "7_2/1_5/9_5_P",
        ]
        with patch("builtins.input", side_effect=entradas_tipeadas):
            ganador, secuencia, entradas = jugar_punto("A")
        self.assertEqual(entradas, ["5_1_6_X/3_3/2_4/4_1_D", "7_2/1_5/9_5_P"])

    def test_guardar_reporte_txt_incluye_inputs_y_stats(self):
        puntos = [_punto("A", "A", [parsear_bloque_saque("5_1_6_A")])]
        marcador = {"A": 1, "B": 0}
        entradas = ["A", "5_1_6_A"]

        nombre_archivo = guardar_reporte_txt(entradas, puntos, marcador)
        try:
            self.assertTrue(os.path.exists(nombre_archivo))
            with open(nombre_archivo, encoding="utf-8") as archivo:
                contenido = archivo.read()
            self.assertIn("=== Jugadas cargadas ===", contenido)
            self.assertIn("A\n5_1_6_A", contenido)
            self.assertIn("Marcador final: A 1 - 0 B", contenido)
            self.assertIn("Total de puntos cargados: 1", contenido)
            self.assertIn("=== Estadisticas por equipo ===", contenido)
            self.assertIn("--- A ---", contenido)
            self.assertIn("--- B ---", contenido)
            # todo lo del equipo A tiene que aparecer antes que el bloque de equipo B
            self.assertLess(contenido.index("--- A ---"), contenido.index("--- B ---"))
        finally:
            os.remove(nombre_archivo)


class TestDefensaPerdida(unittest.TestCase):

    def test_parseo_defensa_perdida(self):
        bloque = parsear_bloque_defensa("7_-2")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["defensor"], 7)
        self.assertEqual(bloque["calidad_defensa"], -2)
        self.assertEqual(bloque["resultado"], "L")
        self.assertIsNone(bloque["colocador"])
        self.assertIsNone(bloque["atacante"])

    def test_defensa_perdida_no_admite_continuacion(self):
        self.assertIsNone(parsear_bloque_defensa("7_-2/1_5/9_5_P"))

    def test_defensa_perdida_da_el_punto_al_equipo_contrario(self):
        entradas = [
            "5_1_6_X/3_3/2_4/4_1_D",  # B ataca, A defiende
            "7_-2",                     # la defensa de A se pierde por completo
        ]
        with patch("builtins.input", side_effect=entradas):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "B")  # el punto es para quien ataco (B)
        self.assertEqual(len(secuencia), 2)
        self.assertEqual(secuencia[1]["resultado"], "L")


class TestArmadoOverpass(unittest.TestCase):

    def test_parseo_en_bloque_inicial(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_3/1_-1")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["receptor"], 3)
        self.assertEqual(bloque["colocador"], 1)
        self.assertIsNone(bloque["zona_colocacion"])
        self.assertEqual(bloque["resultado"], "K")
        self.assertIsNone(bloque["atacante"])

    def test_parseo_en_bloque_defensa(self):
        bloque = parsear_bloque_defensa("7_2/1_-1")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["defensor"], 7)
        self.assertEqual(bloque["colocador"], 1)
        self.assertIsNone(bloque["zona_colocacion"])
        self.assertEqual(bloque["resultado"], "K")

    def test_no_admite_continuacion(self):
        self.assertIsNone(parsear_bloque_saque("5_1_6_X/3_3/1_-1/4_1_P"))
        self.assertIsNone(parsear_bloque_defensa("7_2/1_-1/4_1_P"))

    def test_no_se_confunde_con_armado_malo(self):
        overpass = parsear_bloque_defensa("7_2/1_-1")
        malo = parsear_bloque_defensa("7_2/1_-2")
        self.assertEqual(overpass["resultado"], "K")
        self.assertEqual(malo["resultado"], "N")

    def test_sigue_el_punto_y_cambia_de_equipo(self):
        entradas = [
            "5_1_6_X/3_3/2_4/4_1_D",  # B ataca, A defiende
            "7_2/1_-1",                  # A defiende bien pero la armada se pasa sola
            "9_1/2_4/4_1_P",             # B recibe ese pase y hace punto
        ]
        with patch("builtins.input", side_effect=entradas):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "B")
        self.assertEqual(len(secuencia), 3)
        self.assertEqual(secuencia[1]["resultado"], "K")
        self.assertEqual(secuencia[2]["equipo_set"], "B")

    def test_no_cuenta_en_estadisticas_de_armado(self):
        puntos = [
            _punto("B", "A", [
                {**parsear_bloque_saque("5_1_6_X/3_3/1_-1"), "equipo_set": "A"},
            ]),
        ]
        stats = calcular_estadisticas_armado(puntos)
        total_a = sum(cantidad for cantidad, _ in stats["A"].values())
        self.assertEqual(total_a, 0)


class TestArmadoMalo(unittest.TestCase):

    def test_parseo_en_bloque_inicial(self):
        bloque = parsear_bloque_saque("5_1_6_X/3_3/1_-2")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["receptor"], 3)
        self.assertEqual(bloque["colocador"], 1)
        self.assertIsNone(bloque["zona_colocacion"])
        self.assertEqual(bloque["resultado"], "N")
        self.assertIsNone(bloque["atacante"])

    def test_parseo_en_bloque_defensa(self):
        bloque = parsear_bloque_defensa("7_2/1_-2")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["defensor"], 7)
        self.assertEqual(bloque["colocador"], 1)
        self.assertIsNone(bloque["zona_colocacion"])
        self.assertEqual(bloque["resultado"], "N")

    def test_no_admite_continuacion(self):
        self.assertIsNone(parsear_bloque_saque("5_1_6_X/3_3/1_-2/4_1_P"))
        self.assertIsNone(parsear_bloque_defensa("7_2/1_-2/4_1_P"))

    def test_armado_malo_en_el_saque_da_el_punto_al_que_saco(self):
        with patch("builtins.input", side_effect=["5_1_6_X/3_3/1_-2"]):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "A")  # B recibio y armo mal -> punto para A
        self.assertEqual(len(secuencia), 1)

    def test_armado_malo_en_continuacion_da_el_punto_al_equipo_contrario(self):
        entradas = [
            "5_1_6_X/3_3/2_4/4_1_D",  # B ataca, A defiende
            "7_2/1_-2",                  # A defiende bien pero arma mal
        ]
        with patch("builtins.input", side_effect=entradas):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "B")  # el punto es para quien ataco (B)
        self.assertEqual(len(secuencia), 2)

    def test_no_cuenta_en_estadisticas_de_armado(self):
        puntos = [
            _punto("B", "A", [
                {**parsear_bloque_saque("5_1_6_X/3_3/1_-2"), "equipo_set": "A"},
            ]),
        ]
        stats = calcular_estadisticas_armado(puntos)
        total_a = sum(cantidad for cantidad, _ in stats["A"].values())
        self.assertEqual(total_a, 0)


class TestAtaqueDePrimera(unittest.TestCase):

    def test_parseo_punto(self):
        bloque = parsear_bloque_defensa("9_A_1_P")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["atacante"], 9)
        self.assertEqual(bloque["zona_ataque"], "1")
        self.assertEqual(bloque["resultado"], "P")
        self.assertIsNone(bloque["defensor"])
        self.assertIsNone(bloque["colocador"])

    def test_parseo_defendido(self):
        bloque = parsear_bloque_defensa("9_A_6_D")
        self.assertIsNotNone(bloque)
        self.assertEqual(bloque["zona_ataque"], "6")
        self.assertEqual(bloque["resultado"], "D")

    def test_formato_invalido(self):
        self.assertIsNone(parsear_bloque_defensa("9_A_2_P"))   # zona 2 no permitida
        self.assertIsNone(parsear_bloque_defensa("9_A_1_X"))   # resultado invalido
        self.assertIsNone(parsear_bloque_defensa("_A_1_P"))    # falta jugador

    def test_flujo_completo_da_el_punto_al_equipo_correcto(self):
        entradas = [
            "5_1_6_X/3_3/2_4/4_1_D",  # B ataca, A defiende
            "9_A_1_P",                  # A ataca de primera y hace punto
        ]
        with patch("builtins.input", side_effect=entradas):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(ganador, "A")
        self.assertIsNone(secuencia[1]["equipo_set"])  # no hubo armado, no cuenta

    def test_ataque_de_primera_no_cuenta_como_armado(self):
        puntos = [
            _punto("A", "A", [
                {**parsear_bloque_saque("5_1_6_X/3_3/2_4/4_1_D"), "equipo_set": "B"},
                {**parsear_bloque_defensa("9_A_1_P"), "equipo_set": None},
            ]),
        ]
        stats = calcular_estadisticas_armado(puntos)
        total_a = sum(cantidad for cantidad, _ in stats["A"].values())
        self.assertEqual(total_a, 0)


class TestEstadisticasAtaque(unittest.TestCase):

    def test_clasificacion_efectivo_defendido_fuera(self):
        puntos = [
            _punto("B", "A", [
                {**parsear_bloque_saque("3_5_4_X/1_3/2_4/9_1_P"), "equipo_atacante": "A"},
            ]),
            _punto("A", "B", [
                {**parsear_bloque_saque("3_5_4_X/1_3/2_4/9_1_D"), "equipo_atacante": "B"},
                {**parsear_bloque_defensa("7_2/1_5/9_5_O"), "equipo_atacante": "A"},
            ]),
        ]
        stats = calcular_estadisticas_ataque(puntos)
        self.assertEqual(stats["A"][9]["efectivo"], 1)
        self.assertEqual(stats["A"][9]["fuera"], 1)
        self.assertEqual(stats["B"][9]["defendido"], 1)

    def test_bloqueo_y_toque_de_bloqueo(self):
        puntos = [
            _punto("B", "A", [
                {**parsear_bloque_saque("3_5_4_X/1_3/2_4/9_1_B_6_P"), "equipo_atacante": "A"},
            ]),
            _punto("B", "A", [
                {**parsear_bloque_saque("3_5_4_X/1_3/2_4/9_1_U_6"), "equipo_atacante": "A"},
            ]),
        ]
        stats = calcular_estadisticas_ataque(puntos)
        self.assertEqual(stats["A"][9]["fuera"], 1)     # el bloqueo cuenta como error del atacante
        self.assertEqual(stats["A"][9]["efectivo"], 1)  # el toque de bloqueo sigue siendo punto

    def test_libre_no_cuenta_como_ataque(self):
        puntos = [
            _punto("B", "A", [
                {**parsear_bloque_saque("3_5_4_X/1_3/2_4/9_F_8"), "equipo_atacante": "A"},
            ]),
        ]
        stats = calcular_estadisticas_ataque(puntos)
        self.assertEqual(stats["A"], {})

    def test_ataque_de_primera_cuenta_en_total_pero_no_tiene_zona_de_origen(self):
        puntos = [
            _punto("A", "A", [
                {**parsear_bloque_saque("5_1_6_X/3_3/2_4/4_1_D"), "equipo_atacante": "B"},
                {**parsear_bloque_defensa("9_A_1_P"), "equipo_atacante": "A"},
            ]),
        ]
        stats = calcular_estadisticas_ataque(puntos)
        registro = stats["A"][9]
        self.assertEqual(registro["efectivo"], 1)
        total_en_zonas = sum(
            datos_zona["efectivo"] + datos_zona["defendido"] + datos_zona["fuera"]
            for datos_zona in registro["zonas"].values()
        )
        self.assertEqual(total_en_zonas, 0)

    def test_desglose_por_zona_de_origen_y_destino(self):
        puntos = [
            _punto("B", "A", [
                {**parsear_bloque_saque("3_5_4_X/1_3/2_4/9_1_P"), "equipo_atacante": "A"},
            ]),
        ]
        stats = calcular_estadisticas_ataque(puntos)
        zona_4 = stats["A"][9]["zonas"]["4"]
        self.assertEqual(zona_4["efectivo"], 1)
        self.assertEqual(zona_4["destinos"][1]["efectivo"], 1)
        self.assertEqual(zona_4["destinos"][5]["efectivo"], 0)
        self.assertEqual(zona_4["destinos"][6]["efectivo"], 0)

    def test_ataques_separan_zona_1_y_zona_2(self):
        puntos = [
            _punto("B", "A", [
                {**parsear_bloque_saque("3_5_4_X/1_3/2_1/9_1_P"), "equipo_atacante": "A"},
            ]),
            _punto("B", "A", [
                {**parsear_bloque_saque("3_5_4_X/1_3/2_2/9_5_P"), "equipo_atacante": "A"},
            ]),
        ]
        stats = calcular_estadisticas_ataque(puntos)
        zonas = stats["A"][9]["zonas"]
        self.assertEqual(zonas["1"]["efectivo"], 1)
        self.assertEqual(zonas["2"]["efectivo"], 1)
        self.assertNotIn("1-2", zonas)

    def test_armado_tambien_separa_zona_1_y_zona_2(self):
        puntos = [
            _punto("B", "A", [
                {**parsear_bloque_saque("3_5_4_X/1_3/2_1/9_1_P"), "equipo_set": "A"},
            ]),
            _punto("B", "A", [
                {**parsear_bloque_saque("3_5_4_X/1_3/2_2/9_5_P"), "equipo_set": "A"},
            ]),
        ]
        stats = calcular_estadisticas_armado(puntos)
        cantidad_1, _ = stats["A"]["1"]
        cantidad_2, _ = stats["A"]["2"]
        self.assertEqual(cantidad_1, 1)
        self.assertEqual(cantidad_2, 1)
        self.assertNotIn("1-2", stats["A"])

    def test_jugar_punto_asigna_equipo_atacante(self):
        with patch("builtins.input", side_effect=["5_1_6_X/3_3/2_4/4_1_P"]):
            ganador, secuencia, _ = jugar_punto("A")
        self.assertEqual(secuencia[0]["equipo_atacante"], "B")  # B recibe y ataca


class TestNombresDeEquipos(unittest.TestCase):

    def test_preguntar_equipo_acepta_letra(self):
        with patch("builtins.input", side_effect=["A"]):
            self.assertEqual(preguntar_equipo("¿Quien saca? ", {"A": "Tigres", "B": "Leones"}), "A")

    def test_preguntar_equipo_acepta_nombre_elegido(self):
        with patch("builtins.input", side_effect=["Tigres"]):
            self.assertEqual(preguntar_equipo("¿Quien saca? ", {"A": "Tigres", "B": "Leones"}), "A")

    def test_preguntar_equipo_nombre_case_insensitive(self):
        with patch("builtins.input", side_effect=["leones"]):
            self.assertEqual(preguntar_equipo("¿Quien saca? ", {"A": "Tigres", "B": "Leones"}), "B")

    def test_preguntar_equipo_reintenta_si_es_invalido(self):
        with patch("builtins.input", side_effect=["Panteras", "B"]):
            self.assertEqual(preguntar_equipo("¿Quien saca? ", {"A": "Tigres", "B": "Leones"}), "B")

    def test_jugar_punto_usa_el_nombre_en_los_prompts(self):
        nombres = {"A": "Tigres", "B": "Leones"}
        with patch("builtins.input", side_effect=["5_1_6_A"]) as mock_input:
            jugar_punto("A", nombres)
        primer_prompt = mock_input.call_args_list[0].args[0]
        self.assertIn("Tigres", primer_prompt)

    def test_reporte_usa_los_nombres_elegidos(self):
        puntos = [_punto("A", "A", [parsear_bloque_saque("5_1_6_A")])]
        texto = formatear_estadisticas(puntos, {"A": "Tigres", "B": "Leones"})
        self.assertIn("--- Tigres ---", texto)
        self.assertIn("--- Leones ---", texto)


def _punto(equipo_saca, equipo_gana, jugadas):
    return {"equipo_saca": equipo_saca, "equipo_gana": equipo_gana, "jugadas": jugadas}


class TestEstadisticasArmado(unittest.TestCase):

    def test_distribucion_por_zona_y_equipo(self):
        puntos = [
            _punto("B", "A", [
                {**parsear_bloque_saque("3_5_4_X/1_3/2_4/9_1_P"), "equipo_set": "A"},
            ]),
            _punto("A", "B", [
                {**parsear_bloque_saque("3_5_4_X/1_3/2_4/9_1_D"), "equipo_set": "B"},
                {**parsear_bloque_defensa("7_2/1_5/9_5_D"), "equipo_set": "B"},
                {**parsear_bloque_defensa("4_1/2_6/3_1_P"), "equipo_set": "A"},
            ]),
        ]

        stats = calcular_estadisticas_armado(puntos)

        cantidad_a_4, pct_a_4 = stats["A"]["4"]
        self.assertEqual(cantidad_a_4, 1)
        cantidad_a_65, pct_a_65 = stats["A"]["6-5"]
        self.assertEqual(cantidad_a_65, 1)
        self.assertAlmostEqual(pct_a_4, 50.0)
        self.assertAlmostEqual(pct_a_65, 50.0)

        cantidad_b_4, pct_b_4 = stats["B"]["4"]
        self.assertEqual(cantidad_b_4, 1)
        cantidad_b_65, pct_b_65 = stats["B"]["6-5"]
        self.assertEqual(cantidad_b_65, 1)
        self.assertAlmostEqual(pct_b_4, 50.0)
        self.assertAlmostEqual(pct_b_65, 50.0)

    def test_sin_datos_no_rompe(self):
        stats = calcular_estadisticas_armado([])
        for equipo in ("A", "B"):
            for grupo in ("1", "2", "6-5", "3", "4"):
                cantidad, porcentaje = stats[equipo][grupo]
                self.assertEqual(cantidad, 0)
                self.assertEqual(porcentaje, 0.0)

    def test_defensa_calidad_0_cuenta_como_armado(self):
        # al sacarse la regla de "defensa perdida", una defensa de calidad 0
        # sigue teniendo armado y ataque, y por lo tanto debe contar.
        puntos = [
            _punto("A", "B", [
                {**parsear_bloque_saque("3_5_4_X/1_3/2_4/9_1_D"), "equipo_set": "B"},
                {**parsear_bloque_defensa("7_0/1_5/9_5_P"), "equipo_set": "A"},
            ]),
        ]
        stats = calcular_estadisticas_armado(puntos)
        total_a = sum(cantidad for cantidad, _ in stats["A"].values())
        total_b = sum(cantidad for cantidad, _ in stats["B"].values())
        self.assertEqual(total_a, 1)
        self.assertEqual(total_b, 1)


class TestEstadisticasRecepcion(unittest.TestCase):

    def test_cantidad_y_porcentaje_por_jugador(self):
        puntos = [
            _punto("B", "A", [{**parsear_bloque_saque("3_5_4_X/7_3/2_4/9_1_P"), "equipo_receptor": "A"}]),
            _punto("A", "B", [{**parsear_bloque_saque("3_5_4_X/7_1/2_4/9_1_D"), "equipo_receptor": "B"}]),
            _punto("B", "A", [{**parsear_bloque_saque("3_5_4_X/7_3/2_4/9_1_P"), "equipo_receptor": "A"}]),
        ]

        stats = calcular_estadisticas_recepcion(puntos)

        self.assertEqual(stats["A"][7]["total"], 2)
        cantidad_c3, pct_c3 = stats["A"][7]["calidades"][3]
        self.assertEqual(cantidad_c3, 2)
        self.assertAlmostEqual(pct_c3, 100.0)

        self.assertEqual(stats["B"][7]["total"], 1)
        cantidad_c1, pct_c1 = stats["B"][7]["calidades"][1]
        self.assertEqual(cantidad_c1, 1)
        self.assertAlmostEqual(pct_c1, 100.0)

    def test_saque_directo_no_genera_recepcion(self):
        puntos = [
            _punto("A", "A", [{**parsear_bloque_saque("5_1_6_A"), "equipo_receptor": None}]),
            _punto("A", "B", [{**parsear_bloque_saque("5_1_6_E"), "equipo_receptor": None}]),
        ]
        stats = calcular_estadisticas_recepcion(puntos)
        self.assertEqual(stats["A"], {})
        self.assertEqual(stats["B"], {})

    def test_jugadores_del_mismo_numero_se_separan_por_equipo(self):
        puntos = [
            _punto("B", "A", [{**parsear_bloque_saque("3_5_4_X/7_2/2_4/9_1_P"), "equipo_receptor": "A"}]),
            _punto("A", "B", [{**parsear_bloque_saque("3_5_4_X/7_0/2_4/9_1_P"), "equipo_receptor": "B"}]),
        ]
        stats = calcular_estadisticas_recepcion(puntos)
        self.assertEqual(set(stats["A"].keys()), {7})
        self.assertEqual(set(stats["B"].keys()), {7})
        cantidad_a, pct_a = stats["A"][7]["calidades"][2]
        self.assertEqual(cantidad_a, 1)
        self.assertAlmostEqual(pct_a, 100.0)
        cantidad_b, pct_b = stats["B"][7]["calidades"][0]
        self.assertEqual(cantidad_b, 1)
        self.assertAlmostEqual(pct_b, 100.0)


class TestArmadoPorCalidadRecepcion(unittest.TestCase):
    """Es una estadistica GLOBAL del equipo (no desglosada por jugador)."""

    def test_es_global_y_suma_jugadores_distintos_juntos(self):
        puntos = [
            _punto("B", "A", [{**parsear_bloque_saque("3_5_4_X/7_3/2_4/9_1_P"), "equipo_receptor": "A"}]),
            _punto("B", "A", [{**parsear_bloque_saque("3_5_4_X/2_3/2_1/9_1_P"), "equipo_receptor": "A"}]),
        ]
        stats = calcular_armado_por_calidad_recepcion(puntos)
        # jugadores distintos (7 y 2) con la misma calidad (3) se suman juntos
        self.assertEqual(stats["A"][3], {"4": 1, "1": 1})

    def test_muestra_hacia_donde_fue_la_armada_segun_la_calidad(self):
        puntos = [
            _punto("B", "A", [{**parsear_bloque_saque("3_5_4_X/7_3/2_4/9_1_P"), "equipo_receptor": "A"}]),
            _punto("B", "A", [{**parsear_bloque_saque("3_5_4_X/7_3/2_1/9_1_P"), "equipo_receptor": "A"}]),
            _punto("B", "A", [{**parsear_bloque_saque("3_5_4_X/7_1/2_2/9_5_D"), "equipo_receptor": "A"}]),
        ]
        stats = calcular_armado_por_calidad_recepcion(puntos)
        self.assertEqual(stats["A"][3], {"4": 1, "1": 1})
        self.assertEqual(stats["A"][1], {"2": 1})

    def test_overpass_no_tiene_armada_asociada(self):
        puntos = [
            _punto("B", "A", [{**parsear_bloque_saque("5_1_6_X/3_-1"), "equipo_receptor": "A"}]),
        ]
        stats = calcular_armado_por_calidad_recepcion(puntos)
        self.assertEqual(stats["A"][-1], {})

    def test_armado_marcado_sin_armado_real_no_cuenta(self):
        puntos = [
            _punto("B", "A", [{**parsear_bloque_saque("3_5_4_X/7_3/2_4_X/9_1_P"), "equipo_receptor": "A"}]),
        ]
        stats = calcular_armado_por_calidad_recepcion(puntos)
        self.assertEqual(stats["A"][3], {})


def _total_tipo_saque(stats_equipo, tipo):
    return sum(sum(cantidades.values()) for cantidades in stats_equipo[tipo].values())


class TestRecepcionPorTipoSaque(unittest.TestCase):
    """Es una estadistica GLOBAL del equipo (no desglosada por jugador), con el
    desglose anidado: tipo de saque -> par concreto de zonas -> calidad."""

    def test_saque_paralelo_1_a_5(self):
        # la zona 1 de un lado queda enfrentada a la zona 5 del otro
        puntos = [
            _punto("B", "A", [{**parsear_bloque_saque("5_1_5_X/3_3/2_4/4_1_P"), "equipo_receptor": "A"}]),
        ]
        stats = calcular_recepcion_por_tipo_saque(puntos)
        self.assertEqual(stats["A"]["paralelo"]["1 a 5"][3], 1)
        self.assertEqual(_total_tipo_saque(stats["A"], "cruzado"), 0)

    def test_cada_par_paralelo_se_cuenta_por_separado(self):
        puntos = [
            _punto("B", "A", [{**parsear_bloque_saque("5_5_1_X/3_3/2_4/4_1_P"), "equipo_receptor": "A"}]),
            _punto("B", "A", [{**parsear_bloque_saque("5_6_6_X/3_2/2_4/4_1_P"), "equipo_receptor": "A"}]),
        ]
        stats = calcular_recepcion_por_tipo_saque(puntos)
        self.assertEqual(stats["A"]["paralelo"]["5 a 1"][3], 1)
        self.assertEqual(stats["A"]["paralelo"]["6 a 6"][2], 1)
        self.assertEqual(sum(stats["A"]["paralelo"]["1 a 5"].values()), 0)
        self.assertEqual(_total_tipo_saque(stats["A"], "paralelo"), 2)

    def test_cada_par_cruzado_se_cuenta_por_separado(self):
        puntos = [
            _punto("B", "A", [{**parsear_bloque_saque("5_1_1_X/3_3/2_4/4_1_P"), "equipo_receptor": "A"}]),
            _punto("B", "A", [{**parsear_bloque_saque("5_5_5_X/3_2/2_4/4_1_P"), "equipo_receptor": "A"}]),
        ]
        stats = calcular_recepcion_por_tipo_saque(puntos)
        self.assertEqual(stats["A"]["cruzado"]["1 a 1"][3], 1)
        self.assertEqual(stats["A"]["cruzado"]["5 a 5"][2], 1)
        self.assertEqual(_total_tipo_saque(stats["A"], "paralelo"), 0)

    def test_saque_cruzado_desde_zona_6(self):
        puntos = [
            _punto("B", "A", [{**parsear_bloque_saque("5_6_1_X/3_3/2_4/4_1_P"), "equipo_receptor": "A"}]),
            _punto("B", "A", [{**parsear_bloque_saque("5_6_5_X/3_2/2_4/4_1_P"), "equipo_receptor": "A"}]),
        ]
        stats = calcular_recepcion_por_tipo_saque(puntos)
        self.assertEqual(stats["A"]["cruzado"]["6 a 1"][3], 1)
        self.assertEqual(stats["A"]["cruzado"]["6 a 5"][2], 1)

    def test_saque_a_otras_zonas_no_se_cuenta(self):
        puntos = [
            _punto("B", "A", [{**parsear_bloque_saque("5_1_2_X/3_3/2_4/4_1_P"), "equipo_receptor": "A"}]),
            _punto("B", "A", [{**parsear_bloque_saque("5_1_9_X/3_3/2_4/4_1_P"), "equipo_receptor": "A"}]),
        ]
        stats = calcular_recepcion_por_tipo_saque(puntos)
        self.assertEqual(_total_tipo_saque(stats["A"], "paralelo"), 0)
        self.assertEqual(_total_tipo_saque(stats["A"], "cruzado"), 0)

    def test_overpass_cuenta_en_calidad_menos_1(self):
        puntos = [
            _punto("B", "A", [{**parsear_bloque_saque("5_1_5_X/3_-1"), "equipo_receptor": "A"}]),
        ]
        stats = calcular_recepcion_por_tipo_saque(puntos)
        self.assertEqual(stats["A"]["paralelo"]["1 a 5"][-1], 1)

    def test_es_global_y_suma_jugadores_distintos_juntos(self):
        puntos = [
            _punto("B", "A", [{**parsear_bloque_saque("5_1_5_X/3_3/2_4/4_1_P"), "equipo_receptor": "A"}]),
            _punto("B", "A", [{**parsear_bloque_saque("5_1_5_X/9_3/2_4/4_1_P"), "equipo_receptor": "A"}]),
        ]
        stats = calcular_recepcion_por_tipo_saque(puntos)
        self.assertEqual(stats["A"]["paralelo"]["1 a 5"][3], 2)

    def test_el_reporte_muestra_el_par_de_zonas(self):
        puntos = [
            _punto("B", "A", [{**parsear_bloque_saque("5_6_6_X/3_3/2_4/4_1_P"), "equipo_receptor": "A"}]),
        ]
        texto = formatear_estadisticas(puntos)
        self.assertIn("Paralelo: 1 recibidos", texto)
        self.assertIn("De 6 a 6: 1 recibidos", texto)


class TestDeshacer(unittest.TestCase):

    def test_x_en_el_primer_prompt_pide_deshacer_el_punto_anterior(self):
        with patch("builtins.input", side_effect=["x"]):
            resultado = jugar_punto("A")
        self.assertEqual(resultado, ("DESHACER", "x"))

    def test_x_mayuscula_tambien_funciona(self):
        with patch("builtins.input", side_effect=["X"]):
            resultado = jugar_punto("A")
        self.assertEqual(resultado, ("DESHACER", "X"))

    def test_x_deshace_la_ultima_continuacion_y_pide_de_nuevo(self):
        entradas = [
            "5_1_6_X/3_3/2_4/4_1_D",  # B ataca, A defiende
            "7_2/1_5/9_5_D",            # continuacion con error, la vamos a deshacer
            "x",                          # deshace esa continuacion
            "7_2/1_5/9_5_P",            # continuacion correcta -> punto
        ]
        with patch("builtins.input", side_effect=entradas):
            ganador, secuencia, entradas_crudas = jugar_punto("A")
        self.assertEqual(len(secuencia), 2)
        self.assertEqual(secuencia[1]["resultado"], "P")
        self.assertEqual(entradas_crudas, ["5_1_6_X/3_3/2_4/4_1_D", "7_2/1_5/9_5_P"])
        self.assertEqual(ganador, "A")

    def test_x_sin_continuaciones_cancela_el_punto_y_reinicia_el_saque(self):
        entradas = [
            "5_1_6_X/3_3/2_4/4_1_D",  # primer intento, pide continuacion
            "x",                          # no hay continuaciones -> cancela todo el punto
            "5_1_6_A",                   # se vuelve a pedir el saque: esta vez as
        ]
        with patch("builtins.input", side_effect=entradas):
            ganador, secuencia, entradas_crudas = jugar_punto("A")
        self.assertEqual(ganador, "A")
        self.assertEqual(len(secuencia), 1)
        self.assertEqual(secuencia[0]["resultado_saque"], "A")
        self.assertEqual(entradas_crudas, ["5_1_6_A"])  # el intento cancelado no queda


class TestRecepcionPorSet(unittest.TestCase):

    def _punto_set(self, numero, jugada, equipo_receptor="A"):
        punto = _punto("B", "A", [{**parsear_bloque_saque(jugada), "equipo_receptor": equipo_receptor}])
        punto["set"] = numero
        return punto

    def test_separa_las_recepciones_por_set(self):
        puntos = [
            self._punto_set(1, "5_1_5_X/7_3/2_4/4_1_P"),
            self._punto_set(1, "5_1_5_X/7_2/2_4/4_1_P"),
            self._punto_set(2, "5_1_5_X/7_0/2_4/4_1_P"),
        ]
        por_set = calcular_estadisticas_recepcion_por_set(puntos)
        self.assertEqual(sorted(por_set), [1, 2])
        self.assertEqual(por_set[1]["A"][7]["total"], 2)
        self.assertEqual(por_set[2]["A"][7]["total"], 1)
        self.assertEqual(por_set[2]["A"][7]["calidades"][0][0], 1)

    def test_la_suma_de_los_sets_coincide_con_el_total(self):
        puntos = [
            self._punto_set(1, "5_1_5_X/7_3/2_4/4_1_P"),
            self._punto_set(2, "5_1_5_X/7_0/2_4/4_1_P"),
        ]
        total = calcular_estadisticas_recepcion(puntos)["A"][7]["total"]
        por_set = calcular_estadisticas_recepcion_por_set(puntos)
        self.assertEqual(sum(por_set[n]["A"][7]["total"] for n in por_set), total)

    def test_puntos_sin_set_se_tratan_como_set_1(self):
        puntos = [
            _punto("B", "A", [{**parsear_bloque_saque("5_1_5_X/7_3/2_4/4_1_P"), "equipo_receptor": "A"}]),
        ]
        por_set = calcular_estadisticas_recepcion_por_set(puntos)
        self.assertEqual(list(por_set), [1])

    def test_el_reporte_muestra_las_secciones_por_set(self):
        puntos = [
            self._punto_set(1, "5_1_5_X/7_3/2_4/4_1_P"),
            self._punto_set(2, "5_1_5_X/7_0/2_4/4_1_P"),
        ]
        texto = formatear_estadisticas(puntos)
        self.assertIn("Recepciones por jugador:", texto)
        self.assertIn("Recepciones por jugador (set 1):", texto)
        self.assertIn("Recepciones por jugador (set 2):", texto)

    def test_con_un_solo_set_no_repite_el_desglose(self):
        puntos = [self._punto_set(1, "5_1_5_X/7_3/2_4/4_1_P")]
        texto = formatear_estadisticas(puntos)
        self.assertIn("Recepciones por jugador:", texto)
        self.assertNotIn("(set 1)", texto)


class TestInformeExcelAlSalir(unittest.TestCase):

    def test_preguntar_si_no_acepta_variantes(self):
        for respuesta in ("s", "S", "si", "Si", "sí", "y", "yes"):
            with patch("builtins.input", side_effect=[respuesta]):
                self.assertTrue(preguntar_si_no("? "), f"deberia aceptar {respuesta!r}")
        for respuesta in ("n", "N", "no", "No"):
            with patch("builtins.input", side_effect=[respuesta]):
                self.assertFalse(preguntar_si_no("? "), f"deberia rechazar {respuesta!r}")

    def test_preguntar_si_no_reintenta_si_es_invalido(self):
        with patch("builtins.input", side_effect=["quizas", "s"]):
            self.assertTrue(preguntar_si_no("? "))

    def test_si_dice_que_no_no_genera_nada(self):
        nombres = {"A": "Local", "B": "Rival"}
        with patch("builtins.input", side_effect=["n"]):
            self.assertIsNone(generar_informe_excel("no_deberia_leerse.txt", nombres))

    def test_avisa_si_el_equipo_no_esta_en_el_volcado(self):
        puntos = [_punto("A", "A", [parsear_bloque_saque("5_1_6_A")])]
        marcador = {"A": 1, "B": 0}
        nombres = {"A": "Local", "B": "Rival"}
        # el volcado se guarda con nombres A/B, asi que "Local" no va a aparecer
        nombre_txt = guardar_reporte_txt(["A", "5_1_6_A"], puntos, marcador)
        try:
            with patch("builtins.input", side_effect=["s", "A"]):
                self.assertIsNone(generar_informe_excel(nombre_txt, nombres))
        finally:
            os.remove(nombre_txt)


class TestCambioDeSet(unittest.TestCase):

    def test_w_activa_cambio_de_set(self):
        with patch("builtins.input", side_effect=["w"]):
            resultado = jugar_punto("A")
        self.assertEqual(resultado, ("CAMBIO_SET", "w"))

    def test_w_mayuscula_tambien_funciona(self):
        with patch("builtins.input", side_effect=["W"]):
            resultado = jugar_punto("A")
        self.assertEqual(resultado, ("CAMBIO_SET", "W"))

    def test_reporte_sin_cambio_de_set_mantiene_el_texto_original(self):
        puntos = [_punto("A", "A", [parsear_bloque_saque("5_1_6_A")])]
        marcador = {"A": 1, "B": 0}
        nombre_archivo = guardar_reporte_txt(["A", "5_1_6_A"], puntos, marcador)
        try:
            with open(nombre_archivo, encoding="utf-8") as archivo:
                contenido = archivo.read()
            self.assertIn("Marcador final: A 1 - 0 B", contenido)
            self.assertNotIn("Sets:", contenido)
        finally:
            os.remove(nombre_archivo)

    def test_reporte_con_cambio_de_set_muestra_el_historial(self):
        puntos = [_punto("A", "A", [parsear_bloque_saque("5_1_6_A")])]
        marcador = {"A": 3, "B": 5}
        historial_sets = [{"A": 25, "B": 20}, {"A": 3, "B": 5}]
        sets_ganados = {"A": 1, "B": 0}
        nombre_archivo = guardar_reporte_txt(
            ["A", "5_1_6_A"], puntos, marcador,
            nombres=None, historial_sets=historial_sets, sets_ganados=sets_ganados,
        )
        try:
            with open(nombre_archivo, encoding="utf-8") as archivo:
                contenido = archivo.read()
            self.assertIn("Sets: A 1 - 0 B", contenido)
            self.assertIn("Set 1: A 25 - 20 B", contenido)
            self.assertIn("Set 2: A 3 - 5 B", contenido)
            self.assertIn("Marcador del set actual: A 3 - 5 B", contenido)
        finally:
            os.remove(nombre_archivo)


class TestEstadisticasBloqueo(unittest.TestCase):

    def _punto_jugado(self, entradas, equipo_saca="A"):
        with patch("builtins.input", side_effect=entradas), \
                patch("sys.stdout", new_callable=io.StringIO):
            ganador, secuencia, _ = jugar_punto(equipo_saca)
        return {"set": 1, "equipo_saca": equipo_saca, "equipo_gana": ganador, "jugadas": secuencia}

    def test_el_bloqueo_en_el_saque_es_del_equipo_que_saco(self):
        # A saca, B recibe y ataca: el que bloquea juega para A
        punto = self._punto_jugado(["5_1_6_X/3_3/2_4/4_1_B_7_P"])
        self.assertEqual(calcular_estadisticas_bloqueo([punto]), {"A": {7: 1}, "B": {}})

    def test_el_bloqueo_en_una_continuacion_es_del_equipo_que_defiende(self):
        # A saca, B ataca y le defienden: ahora ataca A y bloquea el 9 de B
        punto = self._punto_jugado(["5_1_6_X/3_3/2_4/4_1_D", "8_2/1_4/9_5_B_9_P"])
        self.assertEqual(calcular_estadisticas_bloqueo([punto]), {"A": {}, "B": {9: 1}})

    def test_el_bloqueo_punto_siempre_es_del_equipo_que_gana_el_punto(self):
        for entradas in (["5_1_6_X/3_3/2_4/4_1_B_7_P"],
                         ["5_1_6_X/3_3/2_4/4_1_D", "8_2/1_4/9_5_B_9_P"]):
            punto = self._punto_jugado(entradas)
            bloqueos = calcular_estadisticas_bloqueo([punto])
            equipo_que_bloqueo = next(e for e, jugadores in bloqueos.items() if jugadores)
            self.assertEqual(equipo_que_bloqueo, punto["equipo_gana"], entradas)

    def test_usar_el_bloqueo_y_el_bloqueo_rejugable_no_son_bloqueo_punto(self):
        # U es punto del atacante y R deja la pelota en juego: ninguno suma al bloqueo
        puntos = [
            self._punto_jugado(["5_1_6_X/3_3/2_4/4_1_U_7"]),
            self._punto_jugado(["5_1_6_X/3_3/2_4/4_1_R_7", "8_2/1_4/9_5_P"]),
        ]
        self.assertEqual(calcular_estadisticas_bloqueo(puntos), {"A": {}, "B": {}})

    def test_usar_el_bloqueo_cuenta_como_ataque_efectivo_del_atacante(self):
        punto = self._punto_jugado(["5_1_6_X/3_3/2_4/4_1_U_7"])
        registro = calcular_estadisticas_ataque([punto])["B"][4]
        self.assertEqual(registro["efectivo"], 1)
        self.assertEqual(registro["defendido"], 0)
        self.assertEqual(registro["fuera"], 0)
        self.assertEqual(punto["equipo_gana"], "B")

    def test_acumula_varios_bloqueos_del_mismo_jugador(self):
        puntos = [self._punto_jugado(["5_1_6_X/3_3/2_4/4_1_B_7_P"]) for _ in range(3)]
        puntos.append(self._punto_jugado(["5_1_6_X/3_3/2_4/4_1_B_9_P"]))
        self.assertEqual(calcular_estadisticas_bloqueo(puntos)["A"], {7: 3, 9: 1})

    def test_el_reporte_muestra_la_seccion_dentro_de_los_ataques(self):
        puntos = [self._punto_jugado(["5_1_6_X/3_3/2_4/4_1_B_7_P"]) for _ in range(2)]
        texto = formatear_estadisticas(puntos, {"A": "Local", "B": "Rival"})
        self.assertIn("Bloqueos punto por jugador:", texto)
        self.assertIn("Jugador 7: 2 bloqueos punto (100.0%)", texto)
        self.assertIn("Total: 2 bloqueos punto", texto)
        # el equipo que no bloqueo igual muestra la seccion, vacia
        self.assertIn("(sin bloqueos punto)", texto)
        # va despues de la tabla de ataques, no mezclada con ella
        seccion_rival = texto.split("--- Rival ---")[1]
        self.assertLess(
            seccion_rival.index("Ataques por jugador:"),
            seccion_rival.index("Bloqueos punto por jugador:"),
        )


class TestParsearRotacion(unittest.TestCase):

    def test_rotacion_valida(self):
        jugadores, armador = parsear_rotacion("28_S 5 13 88 3 40")
        self.assertEqual(jugadores, [28, 5, 13, 88, 3, 40])
        self.assertEqual(armador, 28)

    def test_el_armador_puede_estar_en_cualquier_zona(self):
        jugadores, armador = parsear_rotacion("1 9_S 6 15 10 2")
        self.assertEqual(jugadores, [1, 9, 6, 15, 10, 2])
        self.assertEqual(armador, 9)

    def test_acepta_comas_y_barras_como_separador(self):
        for entrada in ("28_S,5,13,88,3,40", "28_S/5/13/88/3/40", "  28_S   5 13 88 3 40  "):
            self.assertEqual(parsear_rotacion(entrada)[0], [28, 5, 13, 88, 3, 40], entrada)

    def test_rechaza_cantidad_distinta_de_seis(self):
        for entrada in ("28_S 5 13", "28_S 5 13 88 3 40 7"):
            with self.assertRaises(ValueError) as contexto:
                parsear_rotacion(entrada)
            self.assertIn("6 jugadores", str(contexto.exception))

    def test_rechaza_jugador_repetido(self):
        with self.assertRaises(ValueError) as contexto:
            parsear_rotacion("28_S 5 13 5 3 40")
        self.assertIn("repetido", str(contexto.exception))

    def test_rechaza_numero_invalido(self):
        with self.assertRaises(ValueError):
            parsear_rotacion("28_S cinco 13 88 3 40")

    def test_exige_exactamente_un_armador(self):
        for entrada in ("28 5 13 88 3 40", "28_S 5_S 13 88 3 40"):
            with self.assertRaises(ValueError) as contexto:
                parsear_rotacion(entrada)
            self.assertIn("armador", str(contexto.exception))


class TestJugadorQueSaca(unittest.TestCase):

    ROTACIONES = {
        "A": {"jugadores": [28, 5, 13, 88, 3, 40], "armador": 28},
        "B": {"jugadores": [1, 9, 6, 15, 10, 2], "armador": 9},
    }

    def _punto(self, equipo_saca, equipo_gana, numero_set=1):
        return {"set": numero_set, "equipo_saca": equipo_saca,
                "equipo_gana": equipo_gana, "jugadas": []}

    def test_el_que_saca_primero_arranca_en_la_zona_1(self):
        self.assertEqual(jugador_que_saca(self.ROTACIONES, [], 1, "A"), 28)

    def test_el_que_recibe_primero_saca_con_el_de_la_zona_2(self):
        # B recupera el saque, asi que rota antes de sacar: entra el de zona 2
        puntos = [self._punto("A", "B")]
        self.assertEqual(jugador_que_saca(self.ROTACIONES, puntos, 1, "B"), 9)

    def test_ganar_sacando_no_rota(self):
        puntos = [self._punto("A", "A"), self._punto("A", "A")]
        self.assertEqual(jugador_que_saca(self.ROTACIONES, puntos, 1, "A"), 28)

    def test_cada_side_out_avanza_una_posicion(self):
        puntos, esperados = [], [5, 13, 88, 3, 40, 28]
        for esperado in esperados:
            puntos.append(self._punto("B", "A"))   # side-out de A
            self.assertEqual(jugador_que_saca(self.ROTACIONES, puntos, 1, "A"), esperado)
            puntos.append(self._punto("A", "B"))   # side-out de B, para devolver el saque

    def test_la_rotacion_se_reinicia_en_cada_set(self):
        puntos = [self._punto("B", "A", numero_set=1) for _ in range(3)]
        self.assertEqual(jugador_que_saca(self.ROTACIONES, puntos, 1, "A"), 88)
        self.assertEqual(jugador_que_saca(self.ROTACIONES, puntos, 2, "A"), 28)

    def test_sin_rotacion_cargada_devuelve_none(self):
        self.assertIsNone(jugador_que_saca({}, [], 1, "A"))
        self.assertIsNone(jugador_que_saca({"A": self.ROTACIONES["A"]}, [], 1, "B"))


class TestSacadorOpcional(unittest.TestCase):

    def _jugar(self, entradas, jugador_saca=None):
        with patch("builtins.input", side_effect=entradas), \
                patch("sys.stdout", new_callable=io.StringIO) as salida:
            resultado = jugar_punto("A", {"A": "Local", "B": "Rival"}, jugador_saca)
        return resultado, salida.getvalue()

    def test_las_dos_formas_significan_lo_mismo(self):
        largo, _ = self._jugar(["28_1_5_X/3_3/2_4/4_1_P"], jugador_saca=28)
        corto, _ = self._jugar(["1_5_X/3_3/2_4/4_1_P"], jugador_saca=28)
        self.assertEqual(largo[1][0], corto[1][0])
        self.assertEqual(corto[1][0]["sacador"], 28)

    def test_guarda_el_sacador_aunque_no_se_escriba(self):
        (_, secuencia, crudas), _ = self._jugar(["1_5_A"], jugador_saca=28)
        self.assertEqual(secuencia[0]["sacador"], 28)
        # la entrada cruda queda tal cual se tipeo, para poder repetir la carga
        self.assertEqual(crudas, ["1_5_A"])

    def test_rechaza_un_sacador_distinto_al_que_le_toca(self):
        (_, secuencia, _), salida = self._jugar(["99_1_5_A", "1_5_A"], jugador_saca=28)
        self.assertIn("Saca el jugador 28, no el 99", salida)
        self.assertEqual(secuencia[0]["sacador"], 28)

    def test_sin_rotacion_el_numero_sigue_siendo_obligatorio(self):
        (_, secuencia, _), salida = self._jugar(["1_5_A", "28_1_5_A"])
        self.assertIn("Falta el numero del sacador", salida)
        self.assertEqual(secuencia[0]["sacador"], 28)

    def test_sin_rotacion_no_valida_quien_saca(self):
        (_, secuencia, _), salida = self._jugar(["99_1_5_A"])
        self.assertNotIn("Saca el jugador", salida)
        self.assertEqual(secuencia[0]["sacador"], 99)


class TestCambios(unittest.TestCase):

    NOMBRES = {"A": "Local", "B": "Rival"}

    def _rotaciones(self):
        return {
            "A": {"jugadores": [28, 5, 13, 88, 3, 40], "armador": 28},
            "B": {"jugadores": [1, 9, 6, 15, 10, 2], "armador": 9},
        }

    def _aplicar(self, rotaciones, entrada, entradas_extra=()):
        with patch("builtins.input", side_effect=list(entradas_extra)), \
                patch("sys.stdout", new_callable=io.StringIO) as salida:
            cambio = aplicar_cambio(rotaciones, entrada, self.NOMBRES)
        return cambio, salida.getvalue()

    def test_el_que_entra_ocupa_la_zona_del_que_sale(self):
        rotaciones = self._rotaciones()
        cambio, _ = self._aplicar(rotaciones, "C_7_13")
        self.assertEqual(rotaciones["A"]["jugadores"], [28, 5, 7, 88, 3, 40])
        self.assertEqual(cambio, {"equipo": "A", "entra": 7, "sale": 13, "zona": 3,
                                  "armador": False, "armador_desplazado": None})

    def test_deduce_el_equipo_del_jugador_que_sale(self):
        rotaciones = self._rotaciones()
        cambio, _ = self._aplicar(rotaciones, "C_77_15")
        self.assertEqual(cambio["equipo"], "B")
        self.assertEqual(rotaciones["B"]["jugadores"], [1, 9, 6, 77, 10, 2])
        self.assertEqual(rotaciones["A"]["jugadores"], [28, 5, 13, 88, 3, 40])

    def test_el_que_reemplaza_al_armador_queda_de_armador(self):
        rotaciones = self._rotaciones()
        cambio, _ = self._aplicar(rotaciones, "C_7_28")
        self.assertEqual(rotaciones["A"]["armador"], 7)
        self.assertTrue(cambio["armador"])
        self.assertIsNone(cambio["armador_desplazado"])

    def test_la_marca_s_hace_armador_al_que_entra(self):
        # cambio de armador: entra por un jugador que no era el armador
        rotaciones = self._rotaciones()
        cambio, salida = self._aplicar(rotaciones, "C_7_S_3")
        self.assertEqual(rotaciones["A"]["armador"], 7)
        self.assertEqual(rotaciones["A"]["jugadores"], [28, 5, 13, 88, 7, 40])
        self.assertTrue(cambio["armador"])
        self.assertEqual(cambio["armador_desplazado"], 28)
        self.assertIn("deja de serlo el 28", salida)

    def test_el_armador_anterior_deja_de_serlo(self):
        rotaciones = self._rotaciones()
        self._aplicar(rotaciones, "C_7_S_3")
        # el 28 sigue en cancha pero ya no es el armador
        self.assertIn(28, rotaciones["A"]["jugadores"])
        self.assertNotEqual(rotaciones["A"]["armador"], 28)

    def test_la_marca_s_sobre_el_armador_que_sale_no_cambia_nada_mas(self):
        rotaciones = self._rotaciones()
        cambio, _ = self._aplicar(rotaciones, "C_7_S_28")
        self.assertEqual(rotaciones["A"]["armador"], 7)
        self.assertIsNone(cambio["armador_desplazado"])

    def test_la_marca_s_acepta_minusculas(self):
        rotaciones = self._rotaciones()
        cambio, _ = self._aplicar(rotaciones, "c_7_s_3")
        self.assertTrue(cambio["armador"])
        self.assertEqual(rotaciones["A"]["armador"], 7)

    def test_siempre_queda_exactamente_un_armador(self):
        rotaciones = self._rotaciones()
        for entrada in ("C_7_S_3", "C_50_5", "C_60_S_88", "C_70_7"):
            self._aplicar(rotaciones, entrada)
            self.assertIn(rotaciones["A"]["armador"], rotaciones["A"]["jugadores"], entrada)

    def test_un_cambio_normal_no_toca_al_armador(self):
        rotaciones = self._rotaciones()
        self._aplicar(rotaciones, "C_7_13")
        self.assertEqual(rotaciones["A"]["armador"], 28)

    def test_el_cambio_no_altera_el_orden_de_saque(self):
        rotaciones = self._rotaciones()
        puntos = [{"set": 1, "equipo_saca": "B", "equipo_gana": "A", "jugadas": []}]
        self.assertEqual(jugador_que_saca(rotaciones, puntos, 1, "A"), 5)
        self._aplicar(rotaciones, "C_7_5")
        self.assertEqual(jugador_que_saca(rotaciones, puntos, 1, "A"), 7)

    def test_rechaza_al_que_ya_esta_en_cancha(self):
        rotaciones = self._rotaciones()
        cambio, salida = self._aplicar(rotaciones, "C_5_13")
        self.assertIsNone(cambio)
        self.assertIn("ya esta en cancha", salida)
        self.assertEqual(rotaciones["A"]["jugadores"], [28, 5, 13, 88, 3, 40])

    def test_rechaza_si_el_que_sale_no_esta_en_cancha(self):
        rotaciones = self._rotaciones()
        cambio, salida = self._aplicar(rotaciones, "C_7_77")
        self.assertIsNone(cambio)
        self.assertIn("no esta en cancha", salida)

    def test_rechaza_entrar_y_salir_el_mismo(self):
        cambio, salida = self._aplicar(self._rotaciones(), "C_13_13")
        self.assertIsNone(cambio)
        self.assertIn("no puede entrar y salir", salida)

    def test_rechaza_si_no_hay_rotacion_cargada(self):
        cambio, salida = self._aplicar({}, "C_7_13")
        self.assertIsNone(cambio)
        self.assertIn("No hay rotacion cargada", salida)

    def test_si_el_numero_esta_en_los_dos_equipos_pregunta(self):
        rotaciones = self._rotaciones()
        rotaciones["B"]["jugadores"][0] = 13   # ahora el 13 esta en los dos
        prompts = []

        def responder(prompt=""):
            prompts.append(prompt)
            return "B"

        with patch("builtins.input", responder), patch("sys.stdout", new_callable=io.StringIO):
            cambio = aplicar_cambio(rotaciones, "C_7_13", self.NOMBRES)

        self.assertIn("Los dos equipos tienen al jugador 13", "".join(prompts))
        self.assertEqual(cambio["equipo"], "B")
        self.assertEqual(rotaciones["B"]["jugadores"], [7, 9, 6, 15, 10, 2])
        self.assertEqual(rotaciones["A"]["jugadores"], [28, 5, 13, 88, 3, 40])

    def test_copiar_rotaciones_no_comparte_las_listas(self):
        # asi la formacion inicial del set no se modifica con los cambios
        rotaciones = self._rotaciones()
        copia = copiar_rotaciones(rotaciones)
        self._aplicar(rotaciones, "C_7_28")
        self.assertEqual(copia["A"]["jugadores"], [28, 5, 13, 88, 3, 40])
        self.assertEqual(copia["A"]["armador"], 28)

    def test_jugar_punto_devuelve_el_cambio_en_el_prompt_del_saque(self):
        with patch("builtins.input", side_effect=["C_7_28"]), \
                patch("sys.stdout", new_callable=io.StringIO):
            resultado = jugar_punto("A", self.NOMBRES, 28)
        self.assertEqual(resultado, ("CAMBIO", "C_7_28"))

    def test_no_deja_cambiar_en_medio_del_punto(self):
        entradas = ["1_5_X/3_3/2_4/4_1_D", "C_7_28", "8_2/1_4/9_5_P"]
        with patch("builtins.input", side_effect=entradas), \
                patch("sys.stdout", new_callable=io.StringIO) as salida:
            ganador, secuencia, _ = jugar_punto("A", self.NOMBRES, 28)
        self.assertIn("Los cambios se hacen entre puntos", salida.getvalue())
        self.assertEqual(ganador, "A")
        self.assertEqual(len(secuencia), 2)


class TestArmadoPorSetYPorArmador(unittest.TestCase):

    def _punto_armado(self, numero_set, jugada, equipo_set="A"):
        punto = _punto("B", "A", [{**parsear_bloque_saque(jugada), "equipo_set": equipo_set}])
        punto["set"] = numero_set
        return punto

    def _puntos(self):
        return [
            self._punto_armado(1, "3_5_4_X/1_3/99_4/9_1_P"),
            self._punto_armado(1, "3_5_4_X/1_3/99_2/9_1_P"),
            self._punto_armado(1, "3_5_4_X/1_3/7_4/9_1_P"),    # armado de emergencia
            self._punto_armado(2, "3_5_4_X/1_3/2_4/9_1_P", equipo_set="B"),
            self._punto_armado(2, "3_5_4_X/1_3/99_3/9_1_P"),
        ]

    def test_separa_los_armados_por_set(self):
        por_set = calcular_estadisticas_armado_por_set(self._puntos())
        self.assertEqual(sorted(por_set), [1, 2])
        self.assertEqual(por_set[1]["A"]["4"][0], 2)
        self.assertEqual(por_set[1]["A"]["2"][0], 1)
        self.assertEqual(por_set[2]["A"]["3"][0], 1)
        self.assertEqual(por_set[2]["B"]["4"][0], 1)

    def test_la_suma_de_los_sets_coincide_con_el_total(self):
        puntos = self._puntos()
        total = calcular_estadisticas_armado(puntos)
        por_set = calcular_estadisticas_armado_por_set(puntos)
        for equipo in ("A", "B"):
            for grupo in ("1", "2", "6-5", "3", "4"):
                self.assertEqual(
                    total[equipo][grupo][0],
                    sum(por_set[n][equipo][grupo][0] for n in por_set),
                    f"{equipo} zona {grupo}",
                )

    def test_agrupa_los_armados_por_jugador_y_zona(self):
        datos = calcular_armado_por_armador(self._puntos())["A"]
        self.assertEqual(datos[99]["total"], 3)
        self.assertEqual(datos[99]["zonas"]["4"][0], 1)
        self.assertEqual(datos[99]["zonas"]["2"][0], 1)
        self.assertEqual(datos[99]["zonas"]["3"][0], 1)
        self.assertAlmostEqual(datos[99]["zonas"]["4"][1], 100 / 3)
        self.assertEqual(datos[7]["total"], 1)

    def test_no_mezcla_los_armadores_de_los_dos_equipos(self):
        datos = calcular_armado_por_armador(self._puntos())
        self.assertNotIn(2, datos["A"])
        self.assertEqual(datos["B"][2]["total"], 1)

    def test_el_armado_marcado_con_x_no_cuenta(self):
        # el _X dice que en realidad no hubo armado
        puntos = [self._punto_armado(1, "3_5_4_X/1_3/99_4_X/9_1_P")]
        self.assertEqual(calcular_armado_por_armador(puntos)["A"], {})

    def test_por_armador_separado_por_set(self):
        por_set = calcular_armado_por_armador_por_set(self._puntos())
        self.assertEqual(por_set[1]["A"][99]["total"], 2)
        self.assertEqual(por_set[2]["A"][99]["total"], 1)
        self.assertNotIn(7, por_set[2]["A"])

    def test_el_reporte_solo_lista_a_los_marcados_con_s(self):
        texto = formatear_estadisticas(
            self._puntos(), {"A": "Local", "B": "Rival"}, {"A": {99}},
        )
        seccion = texto[texto.index("Armado por armador:"):texto.index("Recepciones por jugador:")]
        self.assertIn("Jugador 99: 3 armados", seccion)
        self.assertNotIn("Jugador 7:", seccion)
        # el armado de emergencia del 7 se ve igual, para que cuadre con el total
        self.assertIn("Otros jugadores: 1 armados", seccion)

    def test_sin_armadores_conocidos_lista_a_todos(self):
        texto = formatear_estadisticas(self._puntos(), {"A": "Local", "B": "Rival"})
        seccion = texto[texto.index("Armado por armador:"):texto.index("Recepciones por jugador:")]
        self.assertIn("Jugador 99: 3 armados", seccion)
        self.assertIn("Jugador 7: 1 armados", seccion)
        self.assertNotIn("Otros jugadores", seccion)

    def test_el_reporte_muestra_las_secciones_por_set(self):
        texto = formatear_estadisticas(self._puntos(), {"A": "Local", "B": "Rival"}, {"A": {99}})
        self.assertIn("Armado por zona (set 1):", texto)
        self.assertIn("Armado por zona (set 2):", texto)
        self.assertIn("Armado por armador (set 1):", texto)
        self.assertIn("Armado por armador (set 2):", texto)

    def test_con_un_solo_set_no_repite_el_desglose(self):
        puntos = [self._punto_armado(1, "3_5_4_X/1_3/99_4/9_1_P")]
        texto = formatear_estadisticas(puntos, {"A": "Local", "B": "Rival"}, {"A": {99}})
        self.assertNotIn("(set 1)", texto)

    def test_junta_los_armadores_de_la_rotacion_y_de_los_cambios(self):
        # 28 arranca de armador; despues entra 7 con _S y lo desplaza. Los dos
        # tienen que quedar en el desglose, y el armado suelto del 88 en "Otros".
        entradas = [
            "Local", "Rival",
            "28_S 5 13 88 3 40", "",          # UVC sin rotacion
            "B",                              # saca Rival, asi que arma Local
            "9_1_5_X/3_3/28_4/13_1_P",        # arma el 28
            "C_7_S_5",                        # cambio de armador: entra 7, sale 5
            "1_5_E",                          # error de saque de Local
            "9_1_5_X/3_3/7_4/13_1_P",         # arma el 7
            "1_5_E",
            "9_1_5_X/3_3/88_4/13_1_P",        # armado de emergencia del 88
            "salir", "n",
        ]
        with patch("builtins.input", side_effect=entradas), \
                patch("getpass.getpass", return_value=CONTRASENA_CARGA), \
                patch("sys.stdout", new_callable=io.StringIO) as salida:
            cargar_jugadas()
        texto = salida.getvalue()

        archivo = re.search(r"Archivo generado: (.+)", texto).group(1)
        try:
            seccion = texto[texto.index("Armado por armador:"):texto.index("Recepciones por jugador:")]
            self.assertIn("Jugador 7: 1 armados", seccion)
            self.assertIn("Jugador 28: 1 armados", seccion)
            self.assertNotIn("Jugador 88:", seccion)
            self.assertIn("Otros jugadores: 1 armados", seccion)
        finally:
            os.remove(archivo)


class TestMatrizCalidadZonaPorArmador(unittest.TestCase):

    def _punto_rec(self, jugada, equipo="A"):
        bloque = {**parsear_bloque_saque(jugada), "equipo_receptor": equipo, "equipo_set": equipo}
        return _punto("B", "A", [bloque])

    def _puntos(self):
        return [
            self._punto_rec("3_5_4_X/1_3/99_4/9_1_P"),
            self._punto_rec("3_5_4_X/1_3/99_2/9_1_P"),
            self._punto_rec("3_5_4_X/1_0/99_1/9_1_P"),
            self._punto_rec("3_5_4_X/1_3/7_4/9_1_P"),     # armado de emergencia
            self._punto_rec("3_5_4_X/1_2/2_4/9_1_P", equipo="B"),
        ]

    def test_abre_la_matriz_por_jugador(self):
        datos = calcular_armado_por_calidad_recepcion_por_armador(self._puntos())["A"]
        self.assertEqual(datos[99][3], {"4": 1, "2": 1})
        self.assertEqual(datos[99][0], {"1": 1})
        self.assertEqual(datos[99][2], {})
        self.assertEqual(datos[7][3], {"4": 1})

    def test_no_mezcla_equipos(self):
        datos = calcular_armado_por_calidad_recepcion_por_armador(self._puntos())
        self.assertNotIn(2, datos["A"])
        self.assertEqual(datos["B"][2][2], {"4": 1})

    def test_los_armadores_nunca_superan_la_matriz_global(self):
        puntos = self._puntos()
        glob = calcular_armado_por_calidad_recepcion(puntos)
        por_armador = calcular_armado_por_calidad_recepcion_por_armador(puntos)
        for equipo in ("A", "B"):
            for calidad in (-1, 0, 1, 2, 3):
                for zona, cantidad in glob[equipo][calidad].items():
                    suma = sum(datos[calidad].get(zona, 0) for datos in por_armador[equipo].values())
                    self.assertEqual(suma, cantidad, f"{equipo} cal {calidad} zona {zona}")

    def test_el_armado_marcado_con_x_no_cuenta(self):
        puntos = [self._punto_rec("3_5_4_X/1_3/99_4_X/9_1_P")]
        self.assertEqual(calcular_armado_por_calidad_recepcion_por_armador(puntos)["A"], {})

    def test_el_reporte_solo_lista_a_los_marcados_con_s(self):
        texto = formatear_estadisticas(self._puntos(), {"A": "Local", "B": "Rival"}, {"A": {99}})
        self.assertIn("Armado segun calidad de recepcion (armador 99):", texto)
        self.assertNotIn("Armado segun calidad de recepcion (armador 7):", texto)

    def test_sin_armadores_conocidos_lista_a_todos(self):
        texto = formatear_estadisticas(self._puntos(), {"A": "Local", "B": "Rival"})
        self.assertIn("Armado segun calidad de recepcion (armador 99):", texto)
        self.assertIn("Armado segun calidad de recepcion (armador 7):", texto)

    def test_la_matriz_global_sigue_estando(self):
        texto = formatear_estadisticas(self._puntos(), {"A": "Local", "B": "Rival"}, {"A": {99}})
        seccion = texto[texto.index("Armado segun calidad de recepcion (global del equipo):"):]
        seccion = seccion[:seccion.index("(armador")]
        self.assertIn("Hacia zona 4: 2", seccion)   # 99 y 7 armaron a zona 4 con calidad 3


class TestPuntosPorFase(unittest.TestCase):

    def _jugar(self, entradas, equipo_saca="A"):
        with patch("builtins.input", side_effect=entradas), \
                patch("sys.stdout", new_callable=io.StringIO):
            ganador, secuencia, _ = jugar_punto(equipo_saca)
        return {"set": 1, "equipo_saca": equipo_saca, "equipo_gana": ganador,
                "jugadas": secuencia}

    def test_k1_es_el_ataque_del_que_recibe(self):
        punto = self._jugar(["5_1_6_X/3_3/2_4/4_1_P"])
        self.assertEqual(fase_del_punto(punto), "K1")

    def test_k1_tambien_cuando_el_punto_lo_gana_el_que_defiende(self):
        # el ataque de K1 se va afuera: el punto es del otro, pero la fase es K1
        punto = self._jugar(["5_1_6_X/3_3/2_4/4_1_O"])
        self.assertEqual(fase_del_punto(punto), "K1")
        self.assertEqual(punto["equipo_gana"], "A")

    def test_k2_es_la_primera_defensa_del_rally(self):
        punto = self._jugar(["5_1_6_X/3_3/2_4/4_1_D", "8_2/1_4/9_5_P"])
        self.assertEqual(fase_del_punto(punto), "K2")

    def test_k3_es_de_la_segunda_defensa_en_adelante(self):
        punto = self._jugar(["5_1_6_X/3_3/2_4/4_1_D", "8_2/1_4/9_5_D", "7_2/1_4/4_1_P"])
        self.assertEqual(fase_del_punto(punto), "K3")
        punto = self._jugar(["5_1_6_X/3_3/2_4/4_1_D", "8_2/1_4/9_5_D",
                             "7_2/1_4/4_1_D", "3_1/2_4/8_6_P"])
        self.assertEqual(fase_del_punto(punto), "K3")

    def test_el_saque_directo_no_es_k1(self):
        # no hubo recepcion ni armado ni ataque, asi que va aparte
        self.assertEqual(fase_del_punto(self._jugar(["5_1_6_A"])), "Saque")
        self.assertEqual(fase_del_punto(self._jugar(["5_1_6_E"])), "Saque")

    def test_error_en_juego_sin_nada_cargado_no_tiene_fase(self):
        self.assertEqual(fase_del_punto(self._jugar(["f"])), "Sin fase")

    def test_error_en_juego_cuenta_en_la_fase_que_venia(self):
        # el ataque de K1 lo defienden y ahi hay un error: se define en K2
        punto = self._jugar(["5_1_6_X/3_3/2_4/4_1_D", "f"])
        self.assertEqual(fase_del_punto(punto), "K2")

    def test_los_recibidos_de_uno_son_los_hechos_del_otro(self):
        puntos = [
            self._jugar(["5_1_6_X/3_3/2_4/4_1_P"]),      # gana B en K1
            self._jugar(["5_1_6_A"]),                     # gana A en el saque
            self._jugar(["5_1_6_X/3_3/2_4/4_1_D", "8_2/1_4/9_5_P"]),   # gana A en K2
        ]
        fases = calcular_puntos_por_fase(puntos)
        for fase in FASES_RALLY:
            self.assertEqual(fases["A"]["hechos"][fase], fases["B"]["recibidos"][fase], fase)
            self.assertEqual(fases["B"]["hechos"][fase], fases["A"]["recibidos"][fase], fase)

    def test_los_totales_cierran_con_el_marcador(self):
        puntos = [
            self._jugar(["5_1_6_X/3_3/2_4/4_1_P"]),
            self._jugar(["5_1_6_A"]),
            self._jugar(["5_1_6_E"]),
            self._jugar(["f"]),
            self._jugar(["5_1_6_X/3_3/2_4/4_1_D", "8_2/1_4/9_5_P"]),
        ]
        fases = calcular_puntos_por_fase(puntos)
        for equipo in ("A", "B"):
            ganados = sum(1 for p in puntos if p["equipo_gana"] == equipo)
            perdidos = len(puntos) - ganados
            for clase, esperado in (("hechos", ganados), ("recibidos", perdidos)):
                datos = fases[equipo][clase].values()
                self.assertEqual(sum(d["total"] for d in datos), esperado, (equipo, clase))
                # cada punto entra en exactamente una de las dos columnas
                self.assertEqual(
                    sum(d["ganados"] + d["error"] for d in datos), esperado, (equipo, clase)
                )

    def test_separa_el_merito_propio_del_error_del_rival(self):
        ganados = {
            "5_1_6_X/3_3/2_4/4_1_P": "ataque punto",
            "5_1_6_X/3_3/2_4/4_1_U_7": "usa el bloqueo",
            "5_1_6_X/3_3/2_4/4_1_B_7_P": "bloqueo punto",
            "5_1_6_A": "as",
        }
        errores = {
            "5_1_6_X/3_3/2_4/4_1_O": "afuera",
            "5_1_6_X/3_3/2_4/4_1_M": "a la malla",
            "5_1_6_E": "error de saque",
            "5_1_6_X/3_3/2_-2": "armado malo",
        }
        for jugada, que_es in ganados.items():
            punto = self._jugar([jugada])
            fases = calcular_puntos_por_fase([punto])
            fase = fase_del_punto(punto)
            self.assertEqual(fases[punto["equipo_gana"]]["hechos"][fase]["ganados"], 1, que_es)
        for jugada, que_es in errores.items():
            punto = self._jugar([jugada])
            fases = calcular_puntos_por_fase([punto])
            fase = fase_del_punto(punto)
            self.assertEqual(fases[punto["equipo_gana"]]["hechos"][fase]["error"], 1, que_es)

    def test_el_ejemplo_del_pedido(self):
        # ganar en K2 porque el rival mando el balon fuera
        punto = self._jugar(["5_1_6_X/3_3/2_4/4_1_D", "8_2/1_4/9_5_O"])
        self.assertEqual(fase_del_punto(punto), "K2")
        self.assertEqual(causa_del_punto(punto), "O")
        fases = calcular_puntos_por_fase([punto])
        gana = punto["equipo_gana"]
        self.assertEqual(fases[gana]["hechos"]["K2"], {"total": 1, "ganados": 0, "error": 1})
        # y del otro lado el mismo punto es un error propio
        self.assertEqual(
            fases[otro_equipo(gana)]["recibidos"]["K2"], {"total": 1, "ganados": 0, "error": 1}
        )

    def test_el_error_en_juego_es_error_del_rival(self):
        punto = self._jugar(["5_1_6_X/3_3/2_4/4_1_D", "f"])
        self.assertEqual(causa_del_punto(punto), "EJ")
        fases = calcular_puntos_por_fase([punto])
        self.assertEqual(fases[punto["equipo_gana"]]["hechos"]["K2"]["error"], 1)

    def test_las_causas_cierran_con_las_fases(self):
        puntos = [
            self._jugar(["5_1_6_X/3_3/2_4/4_1_P"]),
            self._jugar(["5_1_6_X/3_3/2_4/4_1_O"]),
            self._jugar(["5_1_6_A"]),
            self._jugar(["5_1_6_E"]),
            self._jugar(["f"]),
        ]
        fases = calcular_puntos_por_fase(puntos)
        causas = calcular_puntos_por_causa(puntos)
        for equipo in ("A", "B"):
            for clase in ("hechos", "recibidos"):
                por_fase = fases[equipo][clase].values()
                por_causa = causas[equipo][clase]
                self.assertEqual(
                    sum(d["ganados"] for d in por_fase),
                    sum(por_causa[c] for c in CAUSAS_GANADAS), (equipo, clase),
                )
                self.assertEqual(
                    sum(d["error"] for d in por_fase),
                    sum(por_causa[c] for c in CAUSAS_ERROR), (equipo, clase),
                )

    def test_el_reporte_muestra_el_apartado(self):
        puntos = [self._jugar(["5_1_6_X/3_3/2_4/4_1_P"]), self._jugar(["5_1_6_A"])]
        texto = formatear_estadisticas(puntos, {"A": "Local", "B": "Rival"})
        self.assertIn("Puntos por fase del rally:", texto)
        self.assertIn("Hechos: 1", texto)
        self.assertIn("K1: 1 (100.0%)", texto)


class TestRotacionEnCancha(unittest.TestCase):

    INICIAL = [28, 15, 99, 88, 10, 13]
    ROTACIONES = {"A": {"jugadores": list(INICIAL), "armador": 99}}

    def _punto(self, equipo_saca, equipo_gana, numero_set=1):
        return {"set": numero_set, "equipo_saca": equipo_saca,
                "equipo_gana": equipo_gana, "jugadas": []}

    def test_el_de_la_zona_2_pasa_a_la_1_y_el_de_la_1_a_la_6(self):
        girada = rotar(self.INICIAL)
        self.assertEqual(girada, [15, 99, 88, 10, 13, 28])
        self.assertEqual(girada[0], self.INICIAL[1])   # zona 2 -> zona 1
        self.assertEqual(girada[5], self.INICIAL[0])   # zona 1 -> zona 6

    def test_seis_giros_vuelven_a_la_formacion_inicial(self):
        self.assertEqual(rotar(self.INICIAL, 6), self.INICIAL)
        self.assertEqual(rotar(self.INICIAL, 7), rotar(self.INICIAL, 1))

    def test_no_rota_hasta_recuperar_el_saque(self):
        # gana puntos pero sacando el: la formacion no se mueve
        puntos = [self._punto("A", "A"), self._punto("A", "A")]
        self.assertEqual(rotacion_en_cancha(self.ROTACIONES, puntos, 1, "A"), self.INICIAL)

    def test_rota_al_recuperar_el_saque(self):
        puntos = [self._punto("B", "A")]     # side-out
        self.assertEqual(rotacion_en_cancha(self.ROTACIONES, puntos, 1, "A"),
                         [15, 99, 88, 10, 13, 28])

    def test_la_zona_1_es_siempre_el_que_saca(self):
        puntos = []
        for esperado in [15, 99, 88, 10, 13, 28]:
            puntos.append(self._punto("B", "A"))
            en_cancha = rotacion_en_cancha(self.ROTACIONES, puntos, 1, "A")
            self.assertEqual(en_cancha[0], esperado)
            self.assertEqual(
                jugador_que_saca(self.ROTACIONES, puntos, 1, "A"), en_cancha[0]
            )
            puntos.append(self._punto("A", "B"))   # devuelve el saque

    def test_cada_set_arranca_de_la_formacion_inicial(self):
        puntos = [self._punto("B", "A", numero_set=1) for _ in range(3)]
        self.assertEqual(rotacion_en_cancha(self.ROTACIONES, puntos, 2, "A"), self.INICIAL)

    def test_sin_rotacion_cargada_no_hay_formacion(self):
        self.assertIsNone(rotacion_en_cancha({}, [], 1, "A"))

    def test_el_cambio_informa_la_zona_donde_esta_parado_ahora(self):
        rotaciones = {"A": {"jugadores": list(self.INICIAL), "armador": 99}}
        puntos = [self._punto("B", "A")]          # ya roto una vez
        # el 28 arranco en zona 1, pero despues del giro esta en la 6
        with patch("sys.stdout", new_callable=io.StringIO) as salida:
            cambio = aplicar_cambio(rotaciones, "C_7_28", {"A": "Local", "B": "Rival"},
                                    puntos, 1)
        self.assertEqual(cambio["zona"], 6)
        self.assertIn("(zona 6)", salida.getvalue())
        # y el que entra hereda el turno de saque del que sale
        self.assertEqual(rotacion_en_cancha(rotaciones, puntos, 1, "A")[5], 7)


class TestPuntosPorZonaArmador(unittest.TestCase):

    def _rot(self):
        return {"A": {"jugadores": [28, 15, 99, 88, 10, 13], "armador": 99}}

    def _punto(self, equipo_saca, equipo_gana, numero_set=1):
        return {"set": numero_set, "equipo_saca": equipo_saca,
                "equipo_gana": equipo_gana, "jugadas": []}

    def test_la_zona_del_armador_sale_de_la_formacion_actual(self):
        rot = self._rot()
        # el 99 arranca en zona 3
        self.assertEqual(zona_del_armador(rot, [], 1, "A"), 3)
        # tras un side-out la formacion gira y el armador pasa a la zona 2
        self.assertEqual(zona_del_armador(rot, [self._punto("B", "A")], 1, "A"), 2)

    def test_da_la_vuelta_completa(self):
        rot, puntos = self._rot(), []
        zonas = []
        for _ in range(6):
            zonas.append(zona_del_armador(rot, puntos, 1, "A"))
            puntos.append(self._punto("B", "A"))
            puntos.append(self._punto("A", "B"))
        self.assertEqual(zonas, [3, 2, 1, 6, 5, 4])

    def test_sigue_al_armador_despues_de_un_cambio(self):
        rot = self._rot()
        with patch("sys.stdout", new_callable=io.StringIO):
            aplicar_cambio(rot, "C_7_S_13", {"A": "Local", "B": "Rival"})
        # ahora el armador es el 7, que entro en la zona 6
        self.assertEqual(zona_del_armador(rot, [], 1, "A"), 6)

    def test_sin_rotacion_no_hay_zona(self):
        self.assertIsNone(zona_del_armador({}, [], 1, "A"))

    def _punto_con_zona(self, gana, zona, jugadas=None):
        return {"set": 1, "equipo_saca": "A", "equipo_gana": gana,
                "zona_armador": {"A": zona, "B": zona}, "jugadas": jugadas or []}

    def test_reparte_los_puntos_entre_hechas_y_recibidas(self):
        puntos = [self._punto_con_zona("A", 3), self._punto_con_zona("B", 3),
                  self._punto_con_zona("A", 5)]
        datos = calcular_puntos_por_zona_armador(puntos)
        self.assertEqual(sum(datos["A"][3]["hechos"].values()), 1)
        self.assertEqual(sum(datos["A"][3]["recibidos"].values()), 1)
        self.assertEqual(sum(datos["A"][5]["hechos"].values()), 1)
        # lo que uno hace es lo que el otro recibe
        self.assertEqual(sum(datos["B"][3]["recibidos"].values()), 1)
        self.assertEqual(sum(datos["B"][5]["recibidos"].values()), 1)

    def test_los_totales_cierran_con_el_marcador(self):
        with patch("builtins.input", side_effect=["5_1_6_X/3_3/2_4/4_1_P"]), \
                patch("sys.stdout", new_callable=io.StringIO):
            _, secuencia, _ = jugar_punto("A")
        puntos = [self._punto_con_zona("B", 2, secuencia), self._punto_con_zona("A", 4)]
        datos = calcular_puntos_por_zona_armador(puntos)
        for equipo in ("A", "B"):
            hechos = sum(sum(z["hechos"].values()) for z in datos[equipo].values())
            recib = sum(sum(z["recibidos"].values()) for z in datos[equipo].values())
            self.assertEqual(hechos, sum(1 for p in puntos if p["equipo_gana"] == equipo))
            self.assertEqual(hechos + recib, len(puntos))

    def test_los_puntos_sin_zona_no_entran(self):
        puntos = [{"set": 1, "equipo_saca": "A", "equipo_gana": "A", "jugadas": []}]
        datos = calcular_puntos_por_zona_armador(puntos)
        self.assertEqual(sum(sum(z["hechos"].values()) for z in datos["A"].values()), 0)

    def test_el_reporte_muestra_la_seccion(self):
        puntos = [self._punto_con_zona("A", 3)]
        texto = formatear_estadisticas(puntos, {"A": "Local", "B": "Rival"})
        self.assertIn("Puntos con el armador en cada zona:", texto)
        self.assertIn("Zona 3:", texto)
        self.assertIn("Hechos: 1 - K1 0, K2 0, K3 0, As 0, Error de saque 0, Sin fase 1", texto)

    def test_el_saque_se_separa_en_as_y_error_del_rival(self):
        # un as y un error de saque del rival caen los dos en la fase "Saque",
        # pero no son la misma situacion: en uno sacamos nosotros
        with patch("builtins.input", side_effect=["5_1_6_A"]), \
                patch("sys.stdout", new_callable=io.StringIO):
            _, as_directo, _ = jugar_punto("A")
        with patch("builtins.input", side_effect=["5_1_6_E"]), \
                patch("sys.stdout", new_callable=io.StringIO):
            _, error_saque, _ = jugar_punto("B")

        puntos = [
            {"set": 1, "equipo_saca": "A", "equipo_gana": "A",
             "zona_armador": {"A": 1, "B": 1}, "jugadas": as_directo},
            {"set": 1, "equipo_saca": "B", "equipo_gana": "A",
             "zona_armador": {"A": 1, "B": 1}, "jugadas": error_saque},
        ]
        hechos = calcular_puntos_por_zona_armador(puntos)["A"][1]["hechos"]
        self.assertEqual(hechos["As"], 1)
        self.assertEqual(hechos["Error de saque"], 1)
        self.assertNotIn("Saque", hechos)

    def test_sin_rotacion_el_reporte_lo_dice(self):
        puntos = [{"set": 1, "equipo_saca": "A", "equipo_gana": "A", "jugadas": []}]
        texto = formatear_estadisticas(puntos, {"A": "Local", "B": "Rival"})
        self.assertIn("(sin rotacion cargada)", texto)


class TestSesionWeb(unittest.TestCase):
    """La interfaz web no reimplementa el motor: reproduce las mismas lineas.
    Estos tests son la garantia de que no puede divergir de la consola."""

    SETUP = ["Local", "Rival", "28_S 5 13 88 3 40", "", "B"]
    # secuencia valida: Rival saca, Local hace side-out y pasa a sacar el de
    # zona 2 (el 5), gana un as, falla el siguiente y devuelve el saque.
    PUNTOS = ["9_1_5_X/3_3/28_4/13_1_P", "1_5_A", "1_5_E", "9_1_5_X/3_0/88_4/13_1_O"]

    def _sesion(self, lineas):
        sesion = sesion_web.SesionPartido()
        for linea in lineas:
            sesion.enviar(linea)
        return sesion

    def test_da_lo_mismo_que_la_consola(self):
        lineas = self.SETUP + self.PUNTOS
        sesion = self._sesion(lineas)

        with patch("builtins.input", side_effect=lineas + ["salir", "n"]), \
                patch("getpass.getpass", return_value=CONTRASENA_CARGA), \
                patch("sys.stdout", new_callable=io.StringIO) as salida:
            cargar_jugadas()
        archivo = re.search(r"Archivo generado: (.+)", salida.getvalue()).group(1)
        try:
            desde_consola = salida.getvalue()
            self.assertIn(sesion.estadisticas(), desde_consola)
        finally:
            os.remove(archivo)

    def test_no_guarda_la_linea_que_el_motor_rechaza(self):
        sesion = self._sesion(self.SETUP)
        antes = list(sesion.lineas)
        resultado = sesion.enviar("no_es_una_jugada")
        self.assertFalse(resultado["ok"])
        self.assertEqual(sesion.lineas, antes)
        self.assertIn("Formato invalido", resultado["mensaje"])

    def test_devuelve_el_mensaje_del_motor_tal_cual(self):
        sesion = self._sesion(self.SETUP + self.PUNTOS[:1])
        # le toca sacar al 5 (zona 2, porque Local recibio primero y rota)
        resultado = sesion.enviar("99_1_5_A")
        self.assertFalse(resultado["ok"])
        self.assertIn("Saca el jugador 5, no el 99", resultado["mensaje"])

    def test_borrar_la_ultima_linea_vuelve_al_estado_anterior(self):
        sesion = self._sesion(self.SETUP + self.PUNTOS)
        marcador = dict(sesion.instantanea()["marcador"])
        # saca Rival, que no tiene rotacion: el numero va explicito
        self.assertTrue(sesion.enviar("9_1_5_A")["ok"])
        self.assertNotEqual(sesion.instantanea()["marcador"], marcador)
        sesion.deshacer_linea()
        self.assertEqual(sesion.instantanea()["marcador"], marcador)

    def test_las_etapas_guian_la_carga(self):
        sesion = sesion_web.SesionPartido()
        etapas = [sesion.instantanea()["etapa"]]
        for linea in self.SETUP:
            etapas.append(sesion.enviar(linea)["estado"]["etapa"])
        self.assertEqual(
            etapas,
            ["nombres", "nombres", "rotacion", "rotacion", "saque_inicial", "jugadas"],
        )

    def test_la_instantanea_dice_quien_saca(self):
        sesion = self._sesion(self.SETUP)
        instantanea = sesion.instantanea()
        self.assertEqual(instantanea["equipo_saca"], "B")
        self.assertIsNone(instantanea["jugador_saca"])   # Rival no tiene rotacion
        sesion.enviar(self.PUNTOS[0])                    # gana Local el side-out
        self.assertEqual(sesion.instantanea()["jugador_saca"], 5)

    def test_cargar_un_partido_entero_de_una(self):
        sesion = sesion_web.SesionPartido()
        resultado = sesion.cargar_lineas("\n".join(self.SETUP + self.PUNTOS))
        self.assertTrue(resultado["ok"])
        self.assertEqual(len(sesion.lineas), len(self.SETUP) + len(self.PUNTOS))

    def test_la_carga_masiva_se_corta_en_la_linea_mala(self):
        sesion = sesion_web.SesionPartido()
        resultado = sesion.cargar_lineas("\n".join(self.SETUP + ["basura"] + self.PUNTOS))
        self.assertFalse(resultado["ok"])
        self.assertIn("Se corto en", resultado["mensaje"])
        self.assertEqual(len(sesion.lineas), len(self.SETUP))

    def test_reiniciar_deja_el_partido_en_cero(self):
        sesion = self._sesion(self.SETUP + self.PUNTOS)
        sesion.reiniciar()
        self.assertEqual(sesion.lineas, [])
        self.assertEqual(sesion.instantanea()["etapa"], "nombres")
        self.assertEqual(sesion.instantanea()["puntos_cargados"], 0)


class TestDeshacerUnSet(unittest.TestCase):
    """Cerrar un set toca media docena de cosas a la vez. Si se cerro de mas, o
    la rotacion del siguiente se cargo mal, la unica salida es volver atras el
    cambio de set entero."""

    SETUP = ["Local", "Rival", "1_S 2 3 4 5 6", "7_S 8 9 10 11 12", "A"]
    SET1 = ["1_5_A", "1_5_A", "1_5_E"]              # Local 2 - Rival 1
    CIERRE = ["w", "n", "99_S 2 3 4 5 6", "7_S 8 9 10 11 12", "A"]

    def _sesion(self, extra=()):
        sesion = sesion_web.SesionPartido()
        for linea in list(self.SETUP) + list(self.SET1) + list(self.CIERRE) + list(extra):
            sesion.enviar(linea)
        return sesion

    def test_la_x_reabre_el_set_anterior(self):
        sesion = self._sesion()
        resultado = sesion.enviar("x")
        self.assertTrue(resultado["ok"])
        self.assertIn("se reabrio el set 1", resultado["mensaje"])

    def test_vuelve_el_marcador_del_set_reabierto(self):
        sesion = self._sesion(["x"])
        instantanea = sesion.instantanea()
        self.assertEqual(instantanea["set"], 1)
        self.assertEqual(instantanea["marcador"], {"A": 2, "B": 1})
        self.assertEqual(instantanea["sets_ganados"], {"A": 0, "B": 0})
        self.assertEqual(instantanea["historial_sets"], [])

    def test_vuelve_la_rotacion_que_tenia_ese_set(self):
        # la del set 2 se habia cargado mal (el 99); tiene que desaparecer
        sesion = self._sesion(["x"])
        self.assertEqual(sesion.instantanea()["rotaciones"]["A"]["jugadores"],
                         [1, 2, 3, 4, 5, 6])

    def test_vuelve_quien_tenia_el_saque(self):
        sesion = self._sesion(["x"])
        # el ultimo punto del set 1 lo gano Rival (error de saque de Local)
        self.assertEqual(sesion.instantanea()["equipo_saca"], "B")

    def test_los_puntos_del_set_reabierto_siguen_estando(self):
        sesion = self._sesion(["x"])
        self.assertEqual(sesion.instantanea()["puntos_cargados"], len(self.SET1))

    def test_otra_x_deshace_el_ultimo_punto_de_ese_set(self):
        sesion = self._sesion(["x", "x"])
        instantanea = sesion.instantanea()
        self.assertEqual(instantanea["marcador"], {"A": 2, "B": 0})
        self.assertEqual(instantanea["puntos_cargados"], len(self.SET1) - 1)

    def test_se_puede_volver_a_cerrar_el_set_con_la_rotacion_bien(self):
        sesion = self._sesion(["x", "w", "n", "13_S 2 3 4 5 6",
                               "7_S 8 9 10 11 12", "B"])
        instantanea = sesion.instantanea()
        self.assertEqual(instantanea["set"], 2)
        self.assertEqual(instantanea["sets_ganados"], {"A": 1, "B": 0})
        self.assertEqual(instantanea["historial_sets"], [{"A": 2, "B": 1}])
        self.assertEqual(instantanea["rotaciones"]["A"]["jugadores"],
                         [13, 2, 3, 4, 5, 6])

    def test_el_volcado_no_guarda_lo_deshecho(self):
        sesion = self._sesion(["x", "w", "n", "13_S 2 3 4 5 6",
                               "7_S 8 9 10 11 12", "B"])
        entradas = sesion.estado["entradas_totales"]
        self.assertNotIn("99_S 2 3 4 5 6", entradas)   # la rotacion mal
        self.assertNotIn("x", entradas)                # el comando de deshacer
        self.assertEqual(entradas.count("w"), 1)       # un solo cierre de set

    def test_la_rotacion_por_set_queda_con_la_buena(self):
        sesion = self._sesion(["x", "w", "n", "13_S 2 3 4 5 6",
                               "7_S 8 9 10 11 12", "B"])
        por_set = sesion.estado["rotaciones_por_set"]
        self.assertEqual(sorted(por_set), [1, 2])
        self.assertEqual(por_set[2]["A"]["jugadores"], [13, 2, 3, 4, 5, 6])

    def test_sin_sets_cerrados_no_hay_nada_que_deshacer(self):
        sesion = sesion_web.SesionPartido()
        for linea in self.SETUP:
            sesion.enviar(linea)
        resultado = sesion.enviar("x")
        self.assertIn("recien empieza", resultado["mensaje"])
        self.assertEqual(sesion.instantanea()["marcador"], {"A": 0, "B": 0})

    def test_se_pueden_reabrir_dos_sets_seguidos(self):
        sesion = self._sesion(["1_5_A", "w", "s", "A"])   # se juega y cierra el set 2
        self.assertEqual(sesion.instantanea()["set"], 3)
        sesion.enviar("x")                                 # reabre el 2
        self.assertEqual(sesion.instantanea()["set"], 2)
        sesion.enviar("x")                                 # deshace su unico punto
        sesion.enviar("x")                                 # reabre el 1
        instantanea = sesion.instantanea()
        self.assertEqual(instantanea["set"], 1)
        self.assertEqual(instantanea["marcador"], {"A": 2, "B": 1})
        self.assertEqual(instantanea["sets_ganados"], {"A": 0, "B": 0})


class TestLiberos(unittest.TestCase):
    """El libero no rota: entra por el jugador al que cubre mientras ese esta
    atras, y sale cuando pasa adelante. La excepcion es el saque, porque el
    libero no saca."""

    ROTACION = "3_S 88 15 13 16 10"          # zonas 1..6
    CON_LIBERO = ROTACION + "  9_L_15_10"

    def _rotaciones(self, entrada=None):
        entrada = entrada or self.CON_LIBERO
        jugadores, armador = parsear_rotacion(entrada)
        return {"A": {"jugadores": jugadores, "armador": armador,
                      "liberos": av.parsear_liberos(entrada, jugadores)}}

    # ---------------- como se declara ----------------
    def test_la_rotacion_se_lee_igual_con_libero_al_final(self):
        self.assertEqual(parsear_rotacion(self.CON_LIBERO),
                         parsear_rotacion(self.ROTACION))

    def test_el_libero_dice_a_quien_cubre(self):
        jugadores, _ = parsear_rotacion(self.CON_LIBERO)
        self.assertEqual(av.parsear_liberos(self.CON_LIBERO, jugadores),
                         [{"jugador": 9, "cubre": [15, 10]}])

    def test_una_rotacion_sin_libero_no_declara_ninguno(self):
        jugadores, _ = parsear_rotacion(self.ROTACION)
        self.assertEqual(av.parsear_liberos(self.ROTACION, jugadores), [])

    def test_se_pueden_declarar_dos_liberos(self):
        entrada = self.ROTACION + "  9_L_15  2_L_10"
        jugadores, _ = parsear_rotacion(entrada)
        self.assertEqual(av.parsear_liberos(entrada, jugadores),
                         [{"jugador": 9, "cubre": [15]},
                          {"jugador": 2, "cubre": [10]}])

    def test_el_libero_no_puede_estar_en_la_rotacion(self):
        entrada = self.ROTACION + "  15_L_10"
        jugadores, _ = parsear_rotacion(entrada)
        with self.assertRaisesRegex(ValueError, "no puede estar tambien"):
            av.parsear_liberos(entrada, jugadores)

    def test_no_puede_cubrir_a_alguien_que_no_juega(self):
        entrada = self.ROTACION + "  9_L_77"
        jugadores, _ = parsear_rotacion(entrada)
        with self.assertRaisesRegex(ValueError, "no esta en la rotacion"):
            av.parsear_liberos(entrada, jugadores)

    # ---------------- quien esta en la cancha ----------------
    def _formacion(self, rotaciones, giros=0, equipo_saca=None):
        # cada punto que gana el que recibia hace rotar al que recupera el saque
        puntos = [{"set": 1, "equipo_saca": "B", "equipo_gana": "A", "jugadas": []}
                  for _ in range(giros)]
        return av.formacion_en_cancha(rotaciones, puntos, 1, "A", equipo_saca)

    def test_el_libero_entra_por_el_cubierto_que_esta_atras(self):
        # el 10 arranca en zona 6 (atras) y el 15 en zona 3 (adelante)
        formacion = self._formacion(self._rotaciones(), equipo_saca="B")
        self.assertEqual(formacion, [3, 88, 15, 13, 16, 9])

    def test_sin_libero_declarado_la_formacion_es_la_rotacion(self):
        rotaciones = self._rotaciones(self.ROTACION)
        self.assertEqual(self._formacion(rotaciones, equipo_saca="B"),
                         av.rotacion_en_cancha(rotaciones, [], 1, "A"))

    def test_el_cubierto_que_esta_adelante_sigue_jugando(self):
        # el 15 esta en zona 3: adelante, asi que no lo cubre nadie
        self.assertIn(15, self._formacion(self._rotaciones(), equipo_saca="B"))

    def test_el_cubierto_entra_a_sacar_y_el_libero_sale(self):
        # con dos giros el 15 queda en zona 1; si saca su equipo, saca el
        rotaciones = self._rotaciones()
        sacando = self._formacion(rotaciones, giros=2, equipo_saca="A")
        self.assertEqual(sacando[0], 15)
        self.assertNotIn(9, sacando)

    def test_si_recibe_el_libero_ocupa_tambien_la_zona_1(self):
        rotaciones = self._rotaciones()
        recibiendo = self._formacion(rotaciones, giros=2, equipo_saca="B")
        self.assertEqual(recibiendo[0], 9)
        self.assertNotIn(15, recibiendo)

    def test_un_libero_ocupa_una_sola_zona(self):
        # cubre a dos que estan los dos atras: entra por uno solo
        rotaciones = {"A": {"jugadores": [3, 88, 15, 13, 16, 10], "armador": 3,
                            "liberos": [{"jugador": 9, "cubre": [16, 10]}]}}
        formacion = self._formacion(rotaciones, equipo_saca="B")
        self.assertEqual(formacion.count(9), 1)

    def test_dos_liberos_ocupan_las_dos_zonas(self):
        rotaciones = {"A": {"jugadores": [3, 88, 15, 13, 16, 10], "armador": 3,
                            "liberos": [{"jugador": 9, "cubre": [16]},
                                        {"jugador": 2, "cubre": [10]}]}}
        self.assertEqual(self._formacion(rotaciones, equipo_saca="B"),
                         [3, 88, 15, 13, 9, 2])

    def test_la_rotacion_nominal_no_se_toca(self):
        # de ella salen el turno de saque y la zona del armador, asi que las
        # estadisticas no se pueden enterar del libero
        rotaciones = self._rotaciones()
        self.assertEqual(av.rotacion_en_cancha(rotaciones, [], 1, "A"),
                         [3, 88, 15, 13, 16, 10])
        self.assertEqual(av.jugador_que_saca(rotaciones, [], 1, "A"), 3)

    # ---------------- cambiar un libero por otro ----------------
    def test_entra_un_libero_por_el_otro_cubriendo_a_los_mismos(self):
        rotaciones = self._rotaciones()
        cambio = av.aplicar_cambio(rotaciones, "C_2_9", {"A": "Local", "B": "Rival"})
        self.assertEqual(cambio["entra"], 2)
        self.assertTrue(cambio["libero"])
        self.assertEqual(rotaciones["A"]["liberos"], [{"jugador": 2, "cubre": [15, 10]}])
        self.assertEqual(self._formacion(rotaciones, equipo_saca="B")[5], 2)

    def test_el_cambio_de_un_jugador_de_campo_sigue_igual(self):
        rotaciones = self._rotaciones()
        cambio = av.aplicar_cambio(rotaciones, "C_7_88", {"A": "Local", "B": "Rival"})
        self.assertEqual((cambio["entra"], cambio["sale"]), (7, 88))
        self.assertNotIn("libero", cambio)

    def test_la_sesion_web_publica_la_formacion_y_los_liberos(self):
        sesion = sesion_web.SesionPartido()
        for linea in ["Local", "Rival", self.CON_LIBERO, "", "B"]:
            sesion.enviar(linea)
        rotacion = sesion.instantanea()["rotaciones"]["A"]
        self.assertEqual(rotacion["jugadores"], [3, 88, 15, 13, 16, 10])
        self.assertEqual(rotacion["formacion"], [3, 88, 15, 13, 16, 9])
        self.assertEqual(rotacion["liberos"], [{"jugador": 9, "cubre": [15, 10]}])


class TestPuntoPendiente(unittest.TestCase):
    """Una jugada que deja el punto abierto obliga a cargar otra linea, y hasta
    que llegue el punto queda a medias. El que tipea se acuerda de que estaba
    haciendo; una pantalla que arma la jugada tocando, no. Estos tests son la
    garantia de que el motor lo cuenta en vez de que el navegador lo adivine."""

    SETUP = ["Local", "Rival", "28_S 5 13 88 3 40", "", "B"]
    # Saca Rival (B, sin rotacion), asi que ataca Local (A). El 13 remata.
    SAQUE = "9_1_5_X/3_3/28_4/13_1_"

    def _pendiente(self, *jugadas):
        sesion = sesion_web.SesionPartido()
        for linea in list(self.SETUP) + list(jugadas):
            sesion.enviar(linea)
        return sesion.instantanea()["pendiente"]

    def test_sin_nada_cargado_espera_un_saque(self):
        pendiente = self._pendiente()
        self.assertFalse(pendiente["hay"])
        self.assertEqual(pendiente["espera"], "saque")
        self.assertEqual(pendiente["equipo_con_la_pelota"], "B")
        self.assertEqual(pendiente["jugadas"], [])

    def test_un_punto_cerrado_no_deja_nada_pendiente(self):
        pendiente = self._pendiente(self.SAQUE + "P")
        self.assertFalse(pendiente["hay"])
        self.assertEqual(pendiente["espera"], "saque")

    def test_el_as_y_el_error_de_saque_cierran_el_punto(self):
        for linea in ("9_1_5_A", "9_1_5_E"):
            with self.subTest(linea=linea):
                self.assertEqual(self._pendiente(linea)["espera"], "saque")

    def test_defendido_sigue_del_lado_del_que_defendio(self):
        pendiente = self._pendiente(self.SAQUE + "D")
        self.assertTrue(pendiente["hay"])
        self.assertEqual(pendiente["espera"], "continuacion")
        self.assertEqual(pendiente["equipo_con_la_pelota"], "B")

    def test_el_bloqueo_rejugable_sigue_del_mismo_lado_que_ataco(self):
        # el unico resultado que NO cambia de lado: por eso va aparte
        pendiente = self._pendiente(self.SAQUE + "R_5")
        self.assertEqual(pendiente["equipo_con_la_pelota"], "A")

    def test_el_bloqueo_punto_y_el_toque_de_bloqueo_cierran(self):
        for linea in (self.SAQUE + "B_5_P", self.SAQUE + "U_5"):
            with self.subTest(linea=linea):
                self.assertEqual(self._pendiente(linea)["espera"], "saque")

    def test_afuera_y_malla_cierran(self):
        for linea in (self.SAQUE + "O", self.SAQUE + "M"):
            with self.subTest(linea=linea):
                self.assertEqual(self._pendiente(linea)["espera"], "saque")

    def test_el_overpass_de_recepcion_pasa_la_pelota_al_otro_lado(self):
        pendiente = self._pendiente("9_1_5_X/3_-1")
        self.assertEqual(pendiente["espera"], "continuacion")
        self.assertEqual(pendiente["equipo_con_la_pelota"], "B")

    def test_el_overpass_de_armado_pasa_la_pelota_al_otro_lado(self):
        pendiente = self._pendiente("9_1_5_X/3_3/28_-1")
        self.assertEqual(pendiente["espera"], "continuacion")
        self.assertEqual(pendiente["equipo_con_la_pelota"], "B")

    def test_la_armada_mala_cierra_el_punto(self):
        self.assertEqual(self._pendiente("9_1_5_X/3_3/28_-2")["espera"], "saque")

    def test_el_libre_y_el_toque_siguen_del_otro_lado(self):
        for linea in ("9_1_5_X/3_3/28_4/13_F_8", "9_1_5_X/3_3/28_4/13_T_8"):
            with self.subTest(linea=linea):
                pendiente = self._pendiente(linea)
                self.assertEqual(pendiente["espera"], "continuacion")
                self.assertEqual(pendiente["equipo_con_la_pelota"], "B")

    def test_la_defensa_perdida_cierra_el_punto(self):
        pendiente = self._pendiente(self.SAQUE + "D", "5_-2")
        self.assertEqual(pendiente["espera"], "saque")

    def test_cinco_intercambios_van_alternando_el_lado(self):
        jugadas = [self.SAQUE + "D"]
        lados = ["B"]
        for _ in range(4):
            jugadas.append("5_2/88_4/13_1_D" if lados[-1] == "B" else "3_2/28_4/13_1_D")
            lados.append("A" if lados[-1] == "B" else "B")
        pendiente = self._pendiente(*jugadas)
        self.assertEqual(pendiente["equipo_con_la_pelota"], lados[-1])
        self.assertEqual(len(pendiente["jugadas"]), 5)

    def test_la_f_en_medio_de_un_punto_lo_cierra_y_no_deja_nada_colgando(self):
        pendiente = self._pendiente(self.SAQUE + "D", "f")
        self.assertFalse(pendiente["hay"])
        self.assertEqual(pendiente["espera"], "saque")
        self.assertEqual(pendiente["jugadas"], [])

    def test_la_x_del_motor_borra_la_jugada_y_devuelve_la_pelota(self):
        # el punto queda abierto del lado de B; "x" saca esa continuacion y la
        # pelota vuelve a donde estaba antes de cargarla
        pendiente = self._pendiente(self.SAQUE + "D", "5_2/88_4/13_1_D", "x")
        self.assertEqual(len(pendiente["jugadas"]), 1)
        self.assertEqual(pendiente["equipo_con_la_pelota"], "B")

    def test_el_punto_a_medias_no_cuenta_para_el_marcador(self):
        sesion = sesion_web.SesionPartido()
        for linea in self.SETUP + [self.SAQUE + "D"]:
            sesion.enviar(linea)
        instantanea = sesion.instantanea()
        self.assertEqual(instantanea["marcador"], {"A": 0, "B": 0})
        self.assertEqual(instantanea["puntos_cargados"], 0)
        self.assertTrue(instantanea["pendiente"]["hay"])

    def test_la_secuencia_del_punto_en_curso_viene_descrita(self):
        pendiente = self._pendiente(self.SAQUE + "D")
        self.assertIn("Saque: jugador 9", pendiente["secuencia"])
        self.assertIn("defendido", pendiente["secuencia"])
        self.assertEqual(pendiente["lineas"], [self.SAQUE + "D"])


class TestQuePreguntaElMotor(unittest.TestCase):
    """La etapa sale de lo que el motor esta esperando, no de contar lineas.
    Contando, las cuatro respuestas de un cambio de set caian en "jugadas" y la
    pantalla ofrecia la cancha mientras el motor pedia una "n"."""

    SETUP = ["Local", "Rival", "28_S 5 13 88 3 40", "5_S 9 12 2 7 4", "B"]

    def _etapas(self, lineas):
        sesion = sesion_web.SesionPartido()
        etapas = [sesion.instantanea()["etapa"]]
        for linea in lineas:
            etapas.append(sesion.enviar(linea)["estado"]["etapa"])
        return etapas

    def test_la_preparacion_va_pidiendo_una_cosa_por_vez(self):
        self.assertEqual(
            self._etapas(self.SETUP),
            ["nombres", "nombres", "rotacion", "rotacion", "saque_inicial", "jugadas"],
        )

    def test_el_cambio_de_set_ya_no_finge_que_toca_una_jugada(self):
        etapas = self._etapas(self.SETUP + ["1_5_A", "w", "n", "", "", "A"])
        self.assertEqual(
            etapas[-6:],
            # despues del "w": mantener?, rotacion A, rotacion B, quien saca, y recien ahi jugadas
            ["jugadas", "mantener_rotacion", "rotacion", "rotacion", "saque_inicial", "jugadas"],
        )

    def test_el_prompt_dice_a_quien_le_toca_en_medio_de_un_punto(self):
        sesion = sesion_web.SesionPartido()
        for linea in self.SETUP + ["1_5_X/3_3/28_4/13_1_D"]:
            sesion.enviar(linea)
        # saco Rival, ataco Local y Rival defendio: ahora juega Rival
        self.assertEqual(sesion.instantanea()["prompt"], "Juega Rival")

    def test_la_instantanea_dice_que_equipo_le_toca_contestar(self):
        sesion = sesion_web.SesionPartido()
        self.assertEqual(sesion.instantanea()["esperando"],
                         {"que": "nombre_equipo", "equipo": "A"})
        sesion.enviar("Local")
        self.assertEqual(sesion.instantanea()["esperando"],
                         {"que": "nombre_equipo", "equipo": "B"})

    def test_un_cambio_ambiguo_pregunta_de_que_equipo_sale(self):
        # el 13 esta en los dos equipos: el motor no puede adivinar de cual
        # sale y pregunta. Sin avisarlo, la web se comia la jugada siguiente
        # como respuesta y la sesion quedaba trabada.
        sesion = sesion_web.SesionPartido()
        for linea in ["Local", "Rival", "1_S 2 3 13 5 6", "7_S 8 9 13 11 12", "A"]:
            sesion.enviar(linea)
        resultado = sesion.enviar("C_77_13")
        self.assertTrue(resultado["ok"])
        instantanea = sesion.instantanea()
        self.assertEqual(instantanea["etapa"], "equipo_del_cambio")
        self.assertEqual(instantanea["esperando"],
                         {"que": "equipo_del_cambio", "jugador": 13, "entra": 77})
        self.assertIn("los dos equipos", instantanea["prompt"])

    def test_contestado_el_cambio_ambiguo_la_carga_sigue(self):
        sesion = sesion_web.SesionPartido()
        for linea in ["Local", "Rival", "1_S 2 3 13 5 6", "7_S 8 9 13 11 12", "A"]:
            sesion.enviar(linea)
        sesion.enviar("C_77_13")
        sesion.enviar("B")                       # sale el 13 de Rival
        instantanea = sesion.instantanea()
        self.assertEqual(instantanea["etapa"], "jugadas")
        self.assertEqual(instantanea["rotaciones"]["B"]["jugadores"][3], 77)
        self.assertEqual(instantanea["rotaciones"]["A"]["jugadores"][3], 13)
        self.assertTrue(sesion.enviar("1_5_A")["ok"])

    def test_un_cambio_sin_ambiguedad_no_pregunta_nada(self):
        sesion = sesion_web.SesionPartido()
        for linea in ["Local", "Rival", "13_S 2 3 4 5 6", "7_S 8 9 10 11 12", "A"]:
            sesion.enviar(linea)
        sesion.enviar("C_77_2")
        self.assertEqual(sesion.instantanea()["etapa"], "jugadas")

    def test_el_punto_a_medias_no_viaja_dos_veces(self):
        # "esperando" es el resumen y "pendiente" el detalle: los bloques van
        # en uno solo, que es lo que se manda en cada jugada
        sesion = sesion_web.SesionPartido()
        for linea in self.SETUP + ["1_5_X/3_3/28_4/13_1_D"]:
            sesion.enviar(linea)
        esperando = sesion.instantanea()["esperando"]
        self.assertNotIn("jugadas", esperando)
        self.assertNotIn("entradas", esperando)
        self.assertEqual(esperando["que"], "continuacion")


try:
    import openpyxl
    import valores_excel
    HAY_OPENPYXL = True
except ImportError:
    HAY_OPENPYXL = False


class TestNombresEnElInforme(unittest.TestCase):
    """El .txt guarda numeros; el nombre se le agrega al armar el informe.

    Se renombra la CLAVE de cada jugador y no solo la etiqueta que se ve: en
    el informe el mismo texto es lo que dice la celda y el criterio con el que
    las formulas suman desde Datos_Base. Cambiar una sola de las dos dejaria
    las tablas en cero."""

    VOLCADO = Path(__file__).resolve().parent / "Datos" / "partido_20260908_141454.txt"

    def _datos(self):
        import generar_informe_volley as gi
        return gi.parse_volcado(self.VOLCADO.read_text(encoding="utf-8"))["teams"]["Palestino"]

    def test_le_pone_el_nombre_al_jugador(self):
        import generar_informe_volley as gi
        datos = gi.poner_nombres(self._datos(), {"13": "Sofia"})
        self.assertIn("Jugador 13 · Sofia", datos["ataques_jugador"])
        self.assertNotIn("Jugador 13", datos["ataques_jugador"])

    def test_el_que_no_tiene_nombre_queda_igual(self):
        import generar_informe_volley as gi
        datos = gi.poner_nombres(self._datos(), {"13": "Sofia"})
        self.assertIn("Jugador 88", datos["ataques_jugador"])

    def test_sin_nombres_no_cambia_nada(self):
        import generar_informe_volley as gi
        antes = self._datos()["ataques_jugador"]
        self.assertEqual(gi.poner_nombres(self._datos(), {})["ataques_jugador"], antes)

    def test_los_numeros_no_se_mueven(self):
        import generar_informe_volley as gi
        antes = self._datos()["ataques_jugador"]["Jugador 13"]
        despues = gi.poner_nombres(self._datos(), {"13": "Sofia"})["ataques_jugador"]
        self.assertEqual(despues["Jugador 13 · Sofia"], antes)

    def test_renombra_tambien_el_detalle_por_zona(self):
        # es el que alimenta la tabla de ataques por zona de origen: si la
        # clave no coincide con la de la fila, esa tabla sale toda en cero
        import generar_informe_volley as gi
        datos = gi.poner_nombres(self._datos(), {"13": "Sofia"})
        jugadores = {fila[0] for fila in datos["ataques_detalle"]}
        self.assertIn("Jugador 13 · Sofia", jugadores)
        self.assertNotIn("Jugador 13", jugadores)

    def test_renombra_en_todas_las_tablas_por_jugador(self):
        import generar_informe_volley as gi
        datos = gi.poner_nombres(self._datos(), {"13": "Sofia", "3": "Camila"})
        for campo in ("recepciones", "armado_armador", "armado_calidad_armador",
                      "ataques_jugador", "bloqueos_jugador"):
            for clave in datos[campo]:
                if clave.startswith("Jugador 13"):
                    self.assertEqual(clave, "Jugador 13 · Sofia", campo)
                if clave.startswith("Jugador 3 "):
                    self.assertEqual(clave, "Jugador 3 · Camila", campo)


@unittest.skipUnless(HAY_OPENPYXL, "necesita openpyxl")
class TestElInformeConNombres(unittest.TestCase):
    """Ponerle nombres al informe no puede cambiar ningun numero."""

    VOLCADO = str(Path(__file__).resolve().parent / "Datos" / "partido_20260908_141454.txt")

    def _celdas(self, nombres):
        import generar_informe_volley as gi
        import valores_excel
        parsed = gi.parse_volcado(Path(self.VOLCADO).read_text(encoding="utf-8"))
        libro, _ = gi.build_workbook("Palestino", "O'sommer", parsed, nombres)
        valores_excel.convertir_a_valores(libro)
        hoja = libro["Ataque jugador"]
        return [[hoja.cell(row=f, column=c).value for c in range(1, 10)]
                for f in range(1, hoja.max_row + 1)]

    def test_los_numeros_son_los_mismos_con_y_sin_nombres(self):
        sin = self._celdas(None)
        con = self._celdas({"13": "Sofia", "88": "Valentina"})
        self.assertEqual(len(sin), len(con))
        for fila_sin, fila_con in zip(sin, con):
            # la primera celda es la etiqueta; el resto son los numeros
            self.assertEqual(fila_sin[1:], fila_con[1:])

    def test_el_nombre_aparece_en_la_planilla(self):
        con = self._celdas({"13": "Sofia"})
        etiquetas = [str(f[0]) for f in con]
        self.assertTrue(any("Jugador 13 · Sofia" in e for e in etiquetas))

    def test_la_tabla_por_zona_no_queda_en_cero(self):
        con = self._celdas({"13": "Sofia"})
        fila = next(f for f in con if str(f[0]) == "Jugador 13 · Sofia" and f[8] is not None)
        self.assertGreater(num_o_cero(fila[8]), 0)   # ataques totales


def num_o_cero(valor):
    try:
        return float(valor)
    except (TypeError, ValueError):
        return 0


@unittest.skipUnless(HAY_OPENPYXL, "necesita openpyxl")
class TestValoresExcel(unittest.TestCase):
    """El informe se guarda tambien sin formulas para poder mirarlo desde el
    celular. Un numero mal aca es peor que no tener el archivo."""

    def _libro(self, celdas, datos=None):
        libro = openpyxl.Workbook()
        hoja = libro.active
        hoja.title = "Hoja"
        base = libro.create_sheet("Datos_Base")
        for coord, valor in (datos or {}).items():
            base[coord] = valor
        for coord, valor in celdas.items():
            hoja[coord] = valor
        return libro

    def _valor(self, formula, celdas=None, datos=None):
        libro = self._libro({**(celdas or {}), "Z1": formula}, datos)
        convertidas, fallidas = valores_excel.convertir_a_valores(libro)
        self.assertEqual(fallidas, [], formula)
        return libro["Hoja"]["Z1"].value

    def test_suma_un_rango(self):
        self.assertEqual(self._valor("=SUM(A1:A3)", {"A1": 2, "A2": 3, "A3": 5}), 10)

    def test_las_celdas_vacias_valen_cero(self):
        self.assertEqual(self._valor("=SUM(A1:A3)", {"A1": 2}), 2)

    def test_sumif_con_criterio_de_texto(self):
        datos = {"H2": "4", "H3": "2", "H4": "4", "I2": 10, "I3": 5, "I4": 7}
        self.assertEqual(
            self._valor('=SUMIF(Datos_Base!$H$2:$H$4,"4",Datos_Base!$I$2:$I$4)', datos=datos), 17
        )

    def test_sumifs_con_dos_criterios(self):
        datos = {"A2": "hechos", "A3": "hechos", "A4": "recibidos",
                 "B2": "K1", "B3": "K2", "B4": "K1",
                 "C2": 3, "C3": 4, "C4": 9}
        formula = ('=SUMIFS(Datos_Base!$C$2:$C$4,Datos_Base!$A$2:$A$4,"hechos",'
                   'Datos_Base!$B$2:$B$4,"K1")')
        self.assertEqual(self._valor(formula, datos=datos), 3)

    def test_un_criterio_que_parece_celda_no_se_confunde(self):
        # "K1" es un criterio de texto, no la celda K1: si se tomara como
        # referencia el resultado daria 0 y nadie lo notaria
        datos = {"A2": "K1", "A3": "K3", "B2": 7, "B3": 1}
        formula = '=SUMIF(Datos_Base!$A$2:$A$3,"K1",Datos_Base!$B$2:$B$3)'
        self.assertEqual(self._valor(formula, {"K1": 999}, datos=datos), 7)

    def test_criterio_numerico(self):
        datos = {"A2": 1, "A3": 2, "B2": 4, "B3": 8}
        self.assertEqual(
            self._valor("=SUMIF(Datos_Base!$A$2:$A$3,2,Datos_Base!$B$2:$B$3)", datos=datos), 8
        )

    def test_el_criterio_puede_ser_una_celda_de_texto(self):
        # La tabla de ataques por zona filtra por jugador con SUMIFS(...,$A26),
        # y esa celda dice "Jugador 13". Leyendola solo como numero el criterio
        # quedaba en 0 y la tabla entera salia en cero.
        datos = {"A2": "Jugador 13", "A3": "Jugador 8", "B2": 7, "B3": 1}
        formula = '=SUMIF(Datos_Base!$A$2:$A$3,$A$5,Datos_Base!$B$2:$B$3)'
        self.assertEqual(self._valor(formula, {"A5": "Jugador 13"}, datos=datos), 7)

    def test_el_criterio_de_celda_tambien_sirve_con_dos_condiciones(self):
        datos = {"U2": "Jugador 13", "U3": "Jugador 13", "U4": "Jugador 8",
                 "V2": "4", "V3": "2", "V4": "4",
                 "X2": 8, "X3": 11, "X4": 1}
        formula = ('=SUMIFS(Datos_Base!$X$2:$X$4,Datos_Base!$U$2:$U$4,$A$5,'
                   'Datos_Base!$V$2:$V$4,"4")')
        self.assertEqual(self._valor(formula, {"A5": "Jugador 13"}, datos=datos), 8)

    def test_una_celda_de_texto_sigue_sumando_como_cero(self):
        # envolver la celda no puede romper la aritmetica: un texto en una
        # suma vale 0, como en Excel
        self.assertEqual(self._valor("=A1+A2", {"A1": "Jugador 13", "A2": 5}), 5)

    def test_una_formula_que_es_solo_una_referencia_conserva_el_texto(self):
        self.assertEqual(self._valor("=A1", {"A1": "Jugador 13"}), "Jugador 13")

    def test_division_y_porcentaje(self):
        self.assertAlmostEqual(
            self._valor('=IFERROR(A1/A2,"")', {"A1": 3, "A2": 4}), 0.75
        )

    def test_iferror_atrapa_la_division_por_cero(self):
        self.assertEqual(self._valor('=IFERROR(A1/A2,"")', {"A1": 3, "A2": 0}), "")

    def test_sumproduct_de_la_hoja_partido(self):
        # cuenta sets ganados: filas donde B > C y el set se jugo
        celdas = {"B1": 25, "C1": 18, "B2": 22, "C2": 25, "B3": 0, "C3": 0}
        formula = "=SUMPRODUCT((B1:B3>C1:C3)*((B1:B3+C1:C3)>0))"
        self.assertEqual(self._valor(formula, celdas), 1)

    def test_resuelve_formulas_encadenadas(self):
        celdas = {"A1": 5, "A2": "=A1*2", "A3": "=SUM(A1:A2)"}
        libro = self._libro({**celdas, "Z1": "=A3+1"})
        valores_excel.convertir_a_valores(libro)
        self.assertEqual(libro["Hoja"]["A2"].value, 10)
        self.assertEqual(libro["Hoja"]["A3"].value, 15)
        self.assertEqual(libro["Hoja"]["Z1"].value, 16)

    def test_una_tabla_sin_filas_no_se_suma_a_si_misma(self):
        # el TOTAL de una tabla vacia queda como SUM(B5:B4), que Excel lee al
        # reves e incluye la propia celda: ahi el total correcto es 0
        import generar_informe_volley as gi
        self.assertEqual(gi.sin_rangos_invertidos("=SUM(B5:B4)"), "=0")
        self.assertEqual(gi.sin_rangos_invertidos("=SUM(B5:B9)"), "=SUM(B5:B9)")
        self.assertEqual(gi.sin_rangos_invertidos("=SUM(B5:B5)"), "=SUM(B5:B5)")
        # los rangos a la hoja oculta nunca se tocan
        formula = "=SUM(Datos_Base!$I$2:$I$6)"
        self.assertEqual(gi.sin_rangos_invertidos(formula), formula)

    def test_un_partido_sin_datos_igual_genera_el_informe(self):
        import generar_informe_volley as gi, tempfile, os
        volcado = gi.parse_volcado("\n".join([
            "=== Resultado final ===",
            "Marcador final: Local 1 - 0 Rival",
            "Total de puntos cargados: 1",
            "",
            "=== Estadisticas por equipo ===",
            "--- Local ---",
            "Armado por zona:",
            "--- Rival ---",
            "Armado por zona:",
        ]))
        # el marcador final alcanza cuando no hay lineas "Set N:"
        self.assertEqual(volcado["sets_rows"], [(1, 1, 0)])
        libro, _ = gi.build_workbook("Local", "Rival", volcado)
        with tempfile.TemporaryDirectory() as carpeta:
            ruta, avisos = gi.guardar_informe(libro, os.path.join(carpeta, "x.xlsx"))
            self.assertEqual(avisos, [])
            guardado = openpyxl.load_workbook(ruta)
            self.assertEqual(
                [c.coordinate for h in guardado for f in h.iter_rows() for c in f
                 if isinstance(c.value, str) and c.value.startswith("=")],
                [],
            )

    def test_una_formula_que_no_entiende_queda_intacta(self):
        libro = self._libro({"Z1": "=VLOOKUP(A1,B:C,2,FALSE)"})
        convertidas, fallidas = valores_excel.convertir_a_valores(libro)
        self.assertEqual(convertidas, 0)
        self.assertEqual(len(fallidas), 1)
        # sigue siendo una formula: se ve que no es un numero calculado
        self.assertTrue(str(libro["Hoja"]["Z1"].value).startswith("="))

    def test_el_informe_que_se_entrega_no_tiene_ni_una_formula(self):
        # el archivo se abre en el celular, que no calcula nada
        import generar_informe_volley as gi, tempfile, os
        # un volcado de verdad de Datos/, para probar contra un partido entero
        # y no contra un caso armado a mano
        with open("Datos/partido_20260908_141454.txt", encoding="utf-8") as archivo:
            volcado = gi.parse_volcado(archivo.read())
        libro, _ = gi.build_workbook("Palestino", "O'sommer", volcado)

        with tempfile.TemporaryDirectory() as carpeta:
            destino = os.path.join(carpeta, "informe.xlsx")
            ruta, avisos = gi.guardar_informe(libro, destino)
            self.assertEqual(avisos, [])
            guardado = openpyxl.load_workbook(ruta)
            self.assertEqual(
                [f"{h.title}!{c.coordinate}" for h in guardado for f in h.iter_rows()
                 for c in f if isinstance(c.value, str) and c.value.startswith("=")],
                [],
            )
            # se entrega un solo archivo, no una copia al lado
            self.assertEqual(os.listdir(carpeta), ["informe.xlsx"])

            # y los numeros del archivo guardado son los del volcado
            fila = next(f for f in guardado["Fases y armador"].iter_rows()
                        if str(f[0].value).strip() == "K1")
            esperado = volcado["teams"]["Palestino"]["fases"]["hechos"]["K1"]
            self.assertEqual((fila[1].value, fila[3].value, fila[4].value),
                             (esperado["total"], esperado["ganados"], esperado["error"]))

            partido = guardado["Partido"]
            totales = next(f for f in partido.iter_rows()
                           if str(f[0].value).strip() == "Ataques registrados")
            self.assertEqual(
                totales[1].value,
                sum(v["totales"] for v in volcado["teams"]["Palestino"]["ataques_jugador"].values()),
            )


if __name__ == "__main__":
    unittest.main()
