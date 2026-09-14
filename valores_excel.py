"""
Convierte un informe con formulas en uno con los valores ya calculados.

openpyxl escribe la formula pero no su resultado, y Excel lo calcula recien al
abrir el archivo. Los visores del celular no calculan nada: leen el valor
guardado, que esta vacio, y las tablas se ven en blanco. Esta copia guarda los
numeros para poder mirar el informe desde el telefono.

Se evalua solo la gramatica que genera generar_informe_volley.py (SUM, SUMIF,
SUMIFS, IFERROR, SUMPRODUCT y aritmetica entre celdas). Si aparece algo que no
entiende, deja la formula tal cual en vez de inventar un numero: es preferible
una celda que se ve rara a una cifra equivocada en un informe.
"""
import re

from openpyxl.utils import get_column_letter, column_index_from_string

# Un solo patron y una sola pasada, por dos motivos:
#   - si se sustituyera en dos pasadas, la segunda volveria a encontrar las
#     referencias que la primera ya metio dentro de comillas;
#   - los textos entrecomillados van en la alternativa de adelante para que
#     criterios como "K1" o "6-5" no se confundan con una celda.
_HOJA = r"(?:(?:'([^']+)'|([A-Za-z_][\w.]*))!)?"
RE_REFERENCIA = re.compile(
    r'"[^"]*"'
    r"|" + _HOJA + r"\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?"
)


class Arreglo(list):
    """Lista de numeros con operaciones elemento a elemento, lo justo para el
    unico SUMPRODUCT que genera el informe."""

    def _par(self, otro):
        return otro if isinstance(otro, list) else [otro] * len(self)

    def __gt__(self, otro):
        return Arreglo(int(a > b) for a, b in zip(self, self._par(otro)))

    def __add__(self, otro):
        return Arreglo(a + b for a, b in zip(self, self._par(otro)))

    def __mul__(self, otro):
        return Arreglo(a * b for a, b in zip(self, self._par(otro)))


def _numero(valor):
    if valor is None or valor == "":
        return 0
    if isinstance(valor, bool):
        return int(valor)
    if isinstance(valor, (int, float)):
        return valor
    try:
        return float(valor)
    except (TypeError, ValueError):
        return 0


class Valor(float):
    """El numero de una celda, sin perder lo que la celda decia.

    Una celda puede aparecer en una formula como numero (para sumarla) o como
    criterio de SUMIF/SUMIFS, y ahi lo que importa es su texto: el informe
    filtra por jugador con SUMIFS(...,$A26) y $A26 dice "Jugador 13".
    Convirtiendola solo a numero el criterio quedaba en 0 y la tabla de
    ataques por zona salia entera en cero."""

    def __new__(cls, crudo):
        propio = super().__new__(cls, _numero(crudo))
        propio.crudo = crudo
        return propio


def _sin_envolver(valor):
    return valor.crudo if isinstance(valor, Valor) else valor


def _coincide(valor, criterio) -> bool:
    """Comparacion laxa como la de Excel: "4" y 4 son lo mismo."""
    valor, criterio = _sin_envolver(valor), _sin_envolver(criterio)
    try:
        return float(valor) == float(criterio)
    except (TypeError, ValueError):
        return str(valor).strip().lower() == str(criterio).strip().lower()


class Evaluador:

    def __init__(self, libro):
        self.libro = libro
        self.cache = {}
        self.en_curso = set()

    # ------------------------------------------------------------------
    def valor(self, hoja: str, coord: str):
        clave = (hoja, coord)
        if clave in self.cache:
            return self.cache[clave]
        if clave in self.en_curso:
            raise ValueError(f"referencia circular en {hoja}!{coord}")

        crudo = self.libro[hoja][coord].value
        if not (isinstance(crudo, str) and crudo.startswith("=")):
            self.cache[clave] = crudo
            return crudo

        self.en_curso.add(clave)
        try:
            resultado = self._evaluar(crudo[1:], hoja)
        finally:
            self.en_curso.discard(clave)
        self.cache[clave] = resultado
        return resultado

    def rango(self, hoja: str, desde: str, hasta: str) -> Arreglo:
        c1, f1 = column_index_from_string(re.match(r"[A-Z]+", desde).group()), int(re.search(r"\d+", desde).group())
        c2, f2 = column_index_from_string(re.match(r"[A-Z]+", hasta).group()), int(re.search(r"\d+", hasta).group())
        salida = Arreglo()
        for fila in range(min(f1, f2), max(f1, f2) + 1):
            for col in range(min(c1, c2), max(c1, c2) + 1):
                salida.append(self.valor(hoja, f"{get_column_letter(col)}{fila}"))
        return salida

    # ------------------------------------------------------------------
    def _evaluar(self, formula: str, hoja_actual: str):
        expresion = self._a_python(formula, hoja_actual)
        entorno = {
            "SUM": lambda *args: sum(_numero(v) for a in args
                                     for v in (a if isinstance(a, list) else [a])),
            "SUMIF": self._sumif,
            "SUMIFS": self._sumifs,
            "SUMPRODUCT": lambda *args: sum(_numero(v) for v in args[0]),
            "IFERROR": None,      # se resuelve aparte, necesita evaluacion perezosa
            "_r": lambda h, a, b: self.rango(h, a, b),
            "_c": lambda h, c: Valor(self.valor(h, c)),
        }
        if expresion.startswith("IFERROR("):
            return self._iferror(expresion, entorno)
        return eval(expresion, {"__builtins__": {}}, entorno)   # gramatica cerrada y propia

    def _iferror(self, expresion: str, entorno: dict):
        interior = expresion[len("IFERROR("):-1]
        alternativa = interior.rsplit(",", 1)[1].strip()
        try:
            return eval(interior.rsplit(",", 1)[0], {"__builtins__": {}}, entorno)
        except (ZeroDivisionError, TypeError, ValueError):
            return "" if alternativa in ('""', "''") else eval(alternativa, {"__builtins__": {}}, entorno)

    @staticmethod
    def _sumif(rango_criterio, criterio, rango_suma=None):
        origen = rango_suma if rango_suma is not None else rango_criterio
        return sum(_numero(v) for v, c in zip(origen, rango_criterio) if _coincide(c, criterio))

    @staticmethod
    def _sumifs(rango_suma, *pares):
        condiciones = list(zip(pares[0::2], pares[1::2]))
        total = 0
        for i, valor in enumerate(rango_suma):
            if all(_coincide(rango[i], criterio) for rango, criterio in condiciones):
                total += _numero(valor)
        return total

    # ------------------------------------------------------------------
    def _a_python(self, formula: str, hoja_actual: str) -> str:
        def reemplazo(m):
            if m.group(0).startswith('"'):
                return m.group(0)          # es un criterio de texto, no una celda
            hoja = m.group(1) or m.group(2) or hoja_actual
            desde = m.group(3) + m.group(4)
            if m.group(5):
                return f"_r({hoja!r},{desde!r},{m.group(5) + m.group(6)!r})"
            return f"_c({hoja!r},{desde!r})"

        return RE_REFERENCIA.sub(reemplazo, formula).replace("<>", "!=")


def convertir_a_valores(libro) -> tuple[int, list[str]]:
    """Reemplaza en el libro cada formula por su resultado.

    Devuelve (cuantas se convirtieron, las que no se pudieron). Las que fallan
    quedan con la formula intacta, para que se note que no son un numero."""
    evaluador = Evaluador(libro)
    pendientes = [(hoja.title, celda.coordinate)
                  for hoja in libro
                  for fila in hoja.iter_rows()
                  for celda in fila
                  if isinstance(celda.value, str) and celda.value.startswith("=")]

    convertidas, fallidas = 0, []
    calculados = {}
    for hoja, coord in pendientes:
        try:
            calculados[(hoja, coord)] = evaluador.valor(hoja, coord)
        except Exception as error:
            fallidas.append(f"{hoja}!{coord}: {type(error).__name__}: {error}")

    # se escribe al final para no alterar lo que otras formulas van leyendo
    for (hoja, coord), valor in calculados.items():
        # una formula que es solo una referencia devuelve el Valor envuelto:
        # en la celda va lo que decia la celda original, no su lectura numerica
        libro[hoja][coord].value = _sin_envolver(valor)
        convertidas += 1
    return convertidas, fallidas
